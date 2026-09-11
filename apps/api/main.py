from datetime import datetime, timedelta, timezone
from copy import deepcopy
from contextlib import contextmanager
from secrets import token_urlsafe
from uuid import uuid4
import re
import logging
import os
from time import perf_counter
from typing import Annotated

from fastapi import Body, Cookie, Depends, FastAPI, HTTPException, Query, Request, Response, UploadFile, File, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.exceptions import RequestValidationError
from fastapi.responses import JSONResponse
from pydantic import BaseModel, Field
from pwdlib import PasswordHash
from .storage import mode
from .db_auth import AuthDatabase, StorageError, utcnow
from .db_customers import CustomerDatabase
from .db_assignment import AssignmentDatabase
from .db_activity import ActivityDatabase
from sqlalchemy import inspect

app = FastAPI(title="Call Center API", version="0.1.0")
logger = logging.getLogger("call-center.api")
ALLOWED_ORIGINS = [os.environ["FRONTEND_ORIGIN"]] if os.environ.get("FRONTEND_ORIGIN") else ["http://localhost:3000", "http://127.0.0.1:3000", "http://localhost:4174", "http://127.0.0.1:4174"]
app.add_middleware(CORSMiddleware, allow_origins=ALLOWED_ORIGINS, allow_credentials=True, allow_methods=["GET", "POST", "PATCH", "PUT", "DELETE", "OPTIONS"], allow_headers=["*"])
password_hash = PasswordHash.recommended()
SESSION_SECONDS = 8 * 60 * 60
ALEMBIC_HEAD = "009_auth_constraints"


@app.exception_handler(StorageError)
async def storage_error(_: Request, __: StorageError) -> JSONResponse:
    return JSONResponse({"detail": "Storage operation failed."}, status_code=503)


@app.exception_handler(RequestValidationError)
async def validation_error(_: Request, __: RequestValidationError) -> JSONResponse:
    return JSONResponse({"detail": "Invalid request."}, status_code=422)


class User(BaseModel):
    id: str
    name: str
    email: str
    role: str
    active: bool = True


class Customer(BaseModel):
    bcn: str
    name: str
    ownerId: str | None
    ownerName: str | None
    status: str
    phones: list[str] = Field(default_factory=list)
    source: dict[str, str | int | float | bool | None] = Field(default_factory=dict)
    version: int = 0
    histories: list[dict] = Field(default_factory=list)
    followUps: list[dict] = Field(default_factory=list)


class CustomerPage(BaseModel):
    items: list[Customer]
    page: int
    page_size: int
    total: int


class AssignmentRequest(BaseModel):
    ownerId: str | None
    submissionId: str = Field(min_length=1)
    expectedVersion: int | None = None


class Login(BaseModel):
    email: str = Field(min_length=3, max_length=254)
    password: str = Field(min_length=12, max_length=128)

class InteractionCreate(BaseModel):
    outcome: str
    note: str | None = Field(default=None, max_length=4000)
    submissionId: str = Field(min_length=1)

class NoteCreate(BaseModel):
    text: str = Field(min_length=1, max_length=4000)
    submissionId: str = Field(min_length=1)

class FollowUpCreate(BaseModel):
    type: str
    due: str | None = None
    note: str = Field(min_length=1, max_length=4000)
    submissionId: str = Field(min_length=1)

class LifecycleRequest(BaseModel):
    reasonId: str | None = None
    submissionId: str = Field(min_length=1)


class MemoryRepo:
    def __init__(self) -> None:
        self.users: dict[str, dict] = {}
        self.sessions: dict[str, tuple[str, datetime]] = {}
        self.customers: dict[str, dict] = {}
        self.submissions: dict[tuple[str, str], dict] = {}
        self.reasons: dict[str, dict] = {}
        self.interactions: dict[str, dict] = {}
        self.notes: dict[str, dict] = {}
        self.followups: dict[str, dict] = {}
        self.rules: list[dict] = []
        self.assignment_runs: dict[tuple[str, str], dict] = {}
        self.fallback_sales: list[str] = []
        self.assignment_version: int = 1
        self.login_failures: dict[tuple[str, str], list[datetime]] = {}
        self.login_failures_by_ip: dict[str, list[datetime]] = {}
        self.imports: dict[str, dict] = {}
        self.reset()

    def reset(self) -> None:
        self.users.clear(); self.sessions.clear(); self.submissions.clear(); self.interactions.clear(); self.notes.clear(); self.followups.clear(); self.imports.clear(); self.rules.clear(); self.assignment_runs.clear(); self.fallback_sales.clear(); self.assignment_version = 1; self.login_failures.clear(); self.login_failures_by_ip.clear(); self.reasons = {"closure-1": {"id": "closure-1", "label": "Won", "active": True}}
        self.customers = {
            "000123": {"bcn": "000123", "name": "Acme North", "ownerId": "sales-river", "ownerName": "River Sales", "status": "Open", "phones": ["(555) 010-0101"], "source": {"propensity_score": 0.98}, "version": 0, "histories": []},
            "000124": {"bcn": "000124", "name": "Acme North", "ownerId": "sales-sky", "ownerName": "Sky Sales", "status": "Closed", "phones": ["555 010 0103"], "source": {"propensity_score": 0.7}, "version": 0, "histories": []},
            "000125": {"bcn": "000125", "name": "Beta Works", "ownerId": None, "ownerName": None, "status": "Open", "phones": [], "source": {}, "version": 0, "histories": []},
        }

    @contextmanager
    def transaction(self):
        snapshot = (deepcopy(self.users), deepcopy(self.sessions), deepcopy(self.customers), deepcopy(self.submissions), deepcopy(self.interactions), deepcopy(self.notes), deepcopy(self.followups), deepcopy(self.imports))
        try:
            yield self
        except Exception:
            self.users, self.sessions, self.customers, self.submissions, self.interactions, self.notes, self.followups, self.imports = snapshot
            raise


repo = MemoryRepo()
storage_mode = mode()
auth_db = AuthDatabase(__import__("os").environ["DATABASE_URL"], create_schema=False) if storage_mode == "postgres" else None
customer_db = CustomerDatabase(__import__("os").environ["DATABASE_URL"], create_schema=False) if storage_mode == "postgres" else None
assignment_db = AssignmentDatabase(__import__("os").environ["DATABASE_URL"], create_schema=False) if storage_mode == "postgres" else None
activity_db = ActivityDatabase(__import__("os").environ["DATABASE_URL"], create_schema=False) if storage_mode == "postgres" else None
if storage_mode == "postgres":
    # Persistent mode must never let the demonstration fixture shadow database state after restart.
    repo.customers.clear(); repo.followups.clear(); repo.interactions.clear(); repo.notes.clear(); repo.imports.clear(); repo.rules.clear(); repo.assignment_runs.clear()

def append_audit(actor_id: str | None, action: str, target: str, details: dict) -> None:
    if assignment_db is not None: assignment_db.append_audit(actor_id=actor_id, action=action, target=target, details=details)

def persist_activity(record: dict) -> None:
    if activity_db is not None: activity_db.save_activity(record_id=record["id"], bcn=record["bcn"], actor_id=record["actorId"], kind=record["kind"], outcome=record.get("outcome"), text=record.get("text") or record.get("note"))

def persist_followup(record: dict) -> None:
    if activity_db is not None: activity_db.save_followup(record)

def persist_reason(reason: dict) -> None:
    if activity_db is not None: activity_db.save_reason(reason)

def readable_followups() -> list[dict]:
    if activity_db is None: return list(repo.followups.values())
    return [{"id": item.id, "bcn": item.bcn, "type": item.type, "due": item.due.isoformat() if item.due else None, "note": item.note, "status": item.status, "actorId": item.actor_id, "createdAt": item.created_at.isoformat()} for item in activity_db.all_followups()]

@app.get("/health/live")
def health_live() -> dict:
    return {"status": "ok"}

@app.get("/health/ready")
def health_ready() -> dict:
    if auth_db is None: return {"status": "ok", "storage": "memory"}
    try:
        with auth_db.engine.connect() as connection:
            connection.exec_driver_sql("SELECT 1")
            if not inspect(connection).has_table("users"): raise RuntimeError("migrations incomplete")
            if not inspect(connection).has_table("alembic_version") or connection.exec_driver_sql("SELECT version_num FROM alembic_version").scalar() != ALEMBIC_HEAD: raise RuntimeError("migrations incomplete")
        return {"status": "ok", "storage": "postgres"}
    except Exception as exc:
        raise HTTPException(status.HTTP_503_SERVICE_UNAVAILABLE, "Storage is not ready.") from exc

def readable_rows() -> list[dict]:
    if customer_db is None: return list(repo.customers.values())
    pairs = customer_db.all_with_phones(); histories = activity_db.history_map([item.bcn for item, _ in pairs]) if activity_db is not None else {}
    return [{"bcn": item.bcn, "name": item.name, "ownerId": item.owner_id, "ownerName": repo.users.get(item.owner_id or "", {}).get("name"), "status": item.status, "phones": phones, "source": item.source, "version": item.version, "histories": histories.get(item.bcn, [])} for item, phones in pairs]

def import_value(value: object) -> str | int | float | bool | None:
    if isinstance(value, datetime):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


@app.middleware("http")
async def origin_guard(request: Request, call_next):
    started = perf_counter()
    candidate = request.headers.get("x-request-id", "")
    request_id = candidate if re.fullmatch(r"[A-Za-z0-9._-]{1,128}", candidate) else str(uuid4())
    if request.url.path == "/admin/imports":
        try: content_length = int(request.headers.get("content-length", "0"))
        except ValueError: content_length = 0
        if content_length > 11 * 1024 * 1024:
            logger.warning("request id=%s method=%s route=%s status=413 duration_ms=%.3f error=upload_limit", request_id, request.method, request.url.path, (perf_counter() - started) * 1000)
            return Response("Upload is too large.", status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, headers={"x-request-id": request_id}, media_type="application/json")
    if request.method in {"POST", "PATCH", "PUT", "DELETE"} and request.url.path != "/session/login":
        origin = request.headers.get("origin")
        referer = request.headers.get("referer", "")
        if origin not in set(ALLOWED_ORIGINS) and not any(referer.startswith(value + "/") for value in ALLOWED_ORIGINS):
            logger.warning("request id=%s method=%s route=%s status=403 duration_ms=%.3f error=origin", request_id, request.method, request.url.path, (perf_counter() - started) * 1000)
            return Response("Origin not allowed.", status_code=403, headers={"x-request-id": request_id}, media_type="application/json")
    response = await call_next(request)
    response.headers["x-request-id"] = request_id
    route = getattr(request.scope.get("route"), "path", request.url.path)
    logger.info("request id=%s method=%s route=%s status=%s duration_ms=%.3f error=%s", request_id, request.method, route, response.status_code, (perf_counter() - started) * 1000, "none" if response.status_code < 400 else "http_error")
    return response


def safe_email(value: str) -> str:
    return value.strip().casefold()


def current_user(session: Annotated[str | None, Cookie(alias="call_center_session")] = None) -> User:
    if auth_db is not None:
        row = auth_db.user_for_session(session or "")
        if not row or not row.active: raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Authentication required.")
        for item in auth_db.all_users(): repo.users[item.id] = {"id": item.id, "name": item.name, "email": item.email, "role": item.role, "active": item.active, "password": ""}
        if activity_db is not None:
            for item in activity_db.reasons(): repo.reasons[item.id] = {"id": item.id, "label": item.label, "active": item.active}
        if assignment_db is not None:
            repo.rules[:] = [{"id": item.id, "name": item.name, "ownerId": item.owner_id, "active": item.active, "order": item.position} for item in assignment_db.ordered_rules()]
            fallback = assignment_db.get_setting("fallback_sales")
            if fallback is not None: repo.fallback_sales = list(fallback.get("ids", []))
            version = assignment_db.get_setting("assignment_version")
            if version is not None: repo.assignment_version = int(version.get("value", repo.assignment_version))
        return User(id=row.id, name=row.name, email=row.email, role=row.role, active=row.active)
    if not session or session not in repo.sessions:
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Authentication required.")
    user_id, expires = repo.sessions[session]
    if utcnow() >= expires or not repo.users.get(user_id, {}).get("active", False):
        repo.sessions.pop(session, None)
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Authentication required.")
    return User.model_validate(repo.users[user_id])


@app.post("/session/login", response_model=User)
def login(body: Login, request: Request, response: Response) -> User:
    now = utcnow(); ip = request.client.host if request.client else "unknown"; key = (safe_email(body.email), ip)
    recent = [stamp for stamp in repo.login_failures.get(key, []) if now - stamp < timedelta(minutes=15)]
    ip_recent = [stamp for stamp in repo.login_failures_by_ip.get(ip, []) if now - stamp < timedelta(minutes=15)]
    if len(ip_recent) >= 50 or len(recent) >= 5:
        raise HTTPException(status.HTTP_429_TOO_MANY_REQUESTS, "Too many sign-in attempts.", headers={"Retry-After": "900"})
    if auth_db is not None:
        record = auth_db.user_by_email(safe_email(body.email))
        if not record or not record.active or not password_hash.verify(body.password, record.password_hash):
            repo.login_failures[key] = recent + [now]; repo.login_failures_by_ip[ip] = ip_recent + [now]
            raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Unable to sign in.")
        repo.login_failures.pop(key, None)
        token, expires = auth_db.issue(record.id, SESSION_SECONDS); response.set_cookie("call_center_session", token, httponly=True, samesite="lax", secure=request.url.hostname not in {"localhost", "127.0.0.1"}, path="/", max_age=SESSION_SECONDS); return User(id=record.id, name=record.name, email=record.email, role=record.role, active=record.active)
    record = next((u for u in repo.users.values() if u["email"] == safe_email(body.email)), None)
    if not record or not record["active"] or not password_hash.verify(body.password, record["password"]):
        repo.login_failures[key] = recent + [now]
        repo.login_failures_by_ip[ip] = ip_recent + [now]
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Unable to sign in.")
    repo.login_failures.pop(key, None)
    token = token_urlsafe(32); repo.sessions[token] = (record["id"], utcnow() + timedelta(seconds=SESSION_SECONDS))
    response.set_cookie("call_center_session", token, httponly=True, samesite="lax", secure=request.url.hostname not in {"localhost", "127.0.0.1"}, path="/", max_age=SESSION_SECONDS)
    return User.model_validate(record)


@app.post("/session/logout", status_code=204)
def logout(response: Response, session: Annotated[str | None, Cookie(alias="call_center_session")] = None) -> None:
    if session:
        if auth_db is not None: auth_db.revoke(session)
        else: repo.sessions.pop(session, None)
    response.delete_cookie("call_center_session", path="/")


@app.get("/session/me", response_model=User)
def me(user: Annotated[User, Depends(current_user)]) -> User:
    return user

@app.get("/sample-records")
def sample_records(_: Annotated[User, Depends(current_user)]) -> list[dict]:
    return []


@app.get("/customers", response_model=CustomerPage)
def list_customers(user: Annotated[User, Depends(current_user)], page: int = Query(1, ge=1), page_size: int = Query(25, ge=1, le=100), mine: bool = False, q: str = "", status_filter: str | None = Query(None, alias="status"), owner: str | None = None) -> CustomerPage:
    if customer_db is not None:
        owner_id = user.id if mine else (None if not owner or owner == "unassigned" else owner)
        if owner == "unassigned":
            rows = customer_db.search(page=page, page_size=page_size, unassigned=True, status=status_filter, query=q)
        else: rows = customer_db.search(page=page, page_size=page_size, owner_id=owner_id, status=status_filter, query=q)
        items = [{"bcn": row.bcn, "name": row.name, "ownerId": row.owner_id, "ownerName": repo.users.get(row.owner_id or "", {}).get("name"), "status": row.status, "phones": phones, "source": row.source, "version": row.version, "histories": []} for row, phones in rows[0]]
        return CustomerPage(items=[Customer.model_validate(row) for row in items], page=page, page_size=page_size, total=rows[1])
    rows = readable_rows()
    if mine: rows = [row for row in rows if row["ownerId"] == user.id]
    if q: rows = [row for row in rows if q.casefold() in f"{row['bcn']} {row['name']} {' '.join(row.get('phones', []))}".casefold()]
    if status_filter: rows = [row for row in rows if row["status"] == status_filter]
    if owner: rows = [row for row in rows if (row["ownerId"] or "unassigned") == owner]
    rows.sort(key=lambda row: row["bcn"]); total = len(rows); start = (page - 1) * page_size
    return CustomerPage(items=[Customer.model_validate(row) for row in rows[start:start + page_size]], page=page, page_size=page_size, total=total)


@app.get("/customers/{bcn}", response_model=Customer)
def get_customer(bcn: str, user: Annotated[User, Depends(current_user)]) -> Customer:
    db_row = customer_db.get(bcn) if customer_db is not None else None
    if db_row:
        histories = []
        followups = []
        if activity_db is not None:
            histories = [{"id": item.id, "bcn": item.bcn, "kind": item.kind, "outcome": item.outcome, "text": None if item.deleted_at else item.text, "actorId": item.actor_id, "timestamp": item.created_at.isoformat(), "deleted": item.deleted_at is not None} for item in activity_db.history(bcn, 1, 10000)[0]]
            followups = [{"id": item.id, "bcn": item.bcn, "type": item.type, "due": item.due.isoformat() if item.due else None, "note": item.note, "status": item.status, "actorId": item.actor_id, "createdAt": item.created_at.isoformat(), "updatedAt": item.updated_at.isoformat()} for item in activity_db.followups(bcn)]
        row = {"bcn": db_row.bcn, "name": db_row.name, "ownerId": db_row.owner_id, "ownerName": repo.users.get(db_row.owner_id or "", {}).get("name"), "status": db_row.status, "phones": customer_db.phones(bcn) if customer_db is not None else [], "source": db_row.source, "version": db_row.version, "histories": histories, "followUps": followups}
    else: row = repo.customers.get(bcn)
    if not row: raise HTTPException(status.HTTP_404_NOT_FOUND, "Customer not found.")
    return Customer.model_validate(row)


@app.get("/customers/{bcn}/history")
def customer_history(bcn: str, user: Annotated[User, Depends(current_user)], page: int = Query(1, ge=1), page_size: int = Query(25, ge=1, le=100)) -> dict:
    if activity_db is not None:
        rows, total = activity_db.history(bcn, page, page_size)
        return {"items": [{"id": item.id, "bcn": item.bcn, "kind": item.kind, "outcome": item.outcome, "text": None if item.deleted_at else item.text, "actorId": item.actor_id, "timestamp": item.created_at.isoformat(), "deleted": item.deleted_at is not None} for item in rows], "page": page, "page_size": page_size, "total": total}
    row = repo.customers.get(bcn)
    if row is None and customer_db is not None:
        stored = customer_db.get(bcn)
        if stored:
            row = {"bcn": stored.bcn, "name": stored.name, "ownerId": stored.owner_id, "ownerName": repo.users.get(stored.owner_id or "", {}).get("name"), "status": stored.status, "phones": [], "source": stored.source, "version": stored.version, "histories": []}; repo.customers[bcn] = row
    if not row: raise HTTPException(status.HTTP_404_NOT_FOUND, "Customer not found.")
    events = row["histories"]; start = (page - 1) * page_size
    return {"items": events[start:start + page_size], "page": page, "page_size": page_size, "total": len(events)}

def writable_customer(bcn: str, user: User) -> dict:
    row = repo.customers.get(bcn)
    if row is None and customer_db is not None:
        stored = customer_db.get(bcn)
        if stored:
            history = []
            if activity_db is not None: history = [{"id": item.id, "bcn": item.bcn, "kind": item.kind, "outcome": item.outcome, "note": item.text, "actorId": item.actor_id, "timestamp": item.created_at.isoformat(), "deleted": item.deleted_at is not None} for item in activity_db.history(bcn, 1, 10000)[0]]
            row = {"bcn": stored.bcn, "name": stored.name, "ownerId": stored.owner_id, "ownerName": repo.users.get(stored.owner_id or "", {}).get("name"), "status": stored.status, "phones": customer_db.phones(bcn), "source": stored.source, "version": stored.version, "histories": history}; repo.customers[bcn] = row
    if not row: raise HTTPException(status.HTTP_404_NOT_FOUND, "Customer not found.")
    if row["status"] == "Closed": raise HTTPException(status.HTTP_409_CONFLICT, "Customer is closed.")
    if user.role != "Admin" and row["ownerId"] != user.id: raise HTTPException(status.HTTP_403_FORBIDDEN, "Customer access denied.")
    return row

@app.post("/customers/{bcn}/interactions")
def create_interaction(bcn: str, body: InteractionCreate, user: Annotated[User, Depends(current_user)], *, persist: bool = True) -> dict:
    row = writable_customer(bcn, user); key = f"{user.id}:{bcn}:interaction:{body.submissionId}"
    payload = f"{bcn}|{body.outcome}|{body.note or ''}"
    if persist and activity_db is not None:
        try: persisted = activity_db.get_idempotent(actor_id=user.id, operation="interaction", submission_id=body.submissionId, payload=payload)
        except ValueError as exc: raise HTTPException(status.HTTP_409_CONFLICT, str(exc)) from exc
        if persisted: return persisted
    prior = repo.interactions.get(key)
    if prior:
        if prior["outcome"] != body.outcome or prior.get("note") != body.note: raise HTTPException(status.HTTP_409_CONFLICT, "Submission already used.")
        return prior
    if body.outcome not in {"Attempt", "Contact"}: raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Invalid outcome.")
    record = {"id": f"interaction-{len(repo.interactions)+1}", "bcn": bcn, "kind": "Interaction", "outcome": body.outcome, "note": body.note, "actor": user.name, "actorId": user.id, "timestamp": datetime.now(timezone.utc).isoformat(), "deleted": False}
    repo.interactions[key] = record; row["histories"].append(record)
    if persist:
        persist_activity(record)
        if activity_db is not None: activity_db.save_idempotent(actor_id=user.id, operation="interaction", submission_id=body.submissionId, payload=payload, result=record)
    append_audit(user.id, "Interaction created", bcn, {"outcome": body.outcome}); return record

@app.post("/customers/{bcn}/notes")
def create_note(bcn: str, body: NoteCreate, user: Annotated[User, Depends(current_user)]) -> dict:
    row = writable_customer(bcn, user); key = f"{user.id}:{bcn}:note:{body.submissionId}"
    payload = f"{bcn}|{body.text}"
    if activity_db is not None:
        try: persisted = activity_db.get_idempotent(actor_id=user.id, operation="note", submission_id=body.submissionId, payload=payload)
        except ValueError as exc: raise HTTPException(status.HTTP_409_CONFLICT, str(exc)) from exc
        if persisted: return persisted
    prior = repo.notes.get(key)
    if prior:
        if prior["text"] != body.text: raise HTTPException(status.HTTP_409_CONFLICT, "Submission already used.")
        return prior
    record = {"id": f"note-{len(repo.notes)+1}", "bcn": bcn, "kind": "Standalone note", "text": body.text, "actor": user.name, "actorId": user.id, "timestamp": datetime.now(timezone.utc).isoformat(), "deleted": False}
    repo.notes[key] = record; row["histories"].append(record); persist_activity(record)
    if activity_db is not None: activity_db.save_idempotent(actor_id=user.id, operation="note", submission_id=body.submissionId, payload=payload, result=record)
    append_audit(user.id, "Note created", bcn, {}); return record

@app.delete("/customers/{bcn}/history/{record_id}")
def delete_history(bcn: str, record_id: str, user: Annotated[User, Depends(current_user)]) -> dict:
    row = repo.customers.get(bcn)
    if row is None and customer_db is not None:
        stored = customer_db.get(bcn)
        if stored:
            events = [{"id": item.id, "bcn": item.bcn, "kind": item.kind, "outcome": item.outcome, "text": item.text, "actorId": item.actor_id, "timestamp": item.created_at.isoformat(), "deleted": item.deleted_at is not None} for item in (activity_db.history(bcn, 1, 10000)[0] if activity_db is not None else [])]
            row = {"bcn": stored.bcn, "name": stored.name, "ownerId": stored.owner_id, "ownerName": repo.users.get(stored.owner_id or "", {}).get("name"), "status": stored.status, "phones": customer_db.phones(bcn), "source": stored.source, "version": stored.version, "histories": events}; repo.customers[bcn] = row
    if not row: raise HTTPException(status.HTTP_404_NOT_FOUND, "Customer not found.")
    record = next((item for item in row["histories"] if item.get("id") == record_id), None)
    if not record: raise HTTPException(status.HTTP_404_NOT_FOUND, "Record not found.")
    if user.role != "Admin" and row.get("ownerId") != user.id: raise HTTPException(status.HTTP_403_FORBIDDEN, "Record access denied.")
    if record.get("deleted"): return {"id": record_id, "deleted": True, "deletedBy": record.get("deletedBy"), "deletedAt": record.get("deletedAt")}
    record.update(deleted=True, deletedBy=user.id, deletedAt=datetime.now(timezone.utc).isoformat())
    if activity_db is not None: activity_db.soft_delete(record_id, user.id)
    append_audit(user.id, "History deleted", bcn, {"recordId": record_id})
    return {"id": record_id, "deleted": True, "deletedBy": user.id, "deletedAt": record["deletedAt"]}

@app.post("/customers/{bcn}/follow-ups")
def create_followup(bcn: str, body: FollowUpCreate, user: Annotated[User, Depends(current_user)]) -> dict:
    row = writable_customer(bcn, user); key = f"{user.id}:{bcn}:followup:{body.submissionId}"
    payload = f"{bcn}|{body.type}|{body.due or ''}|{body.note}"
    if activity_db is not None:
        try: persisted = activity_db.get_idempotent(actor_id=user.id, operation="followup", submission_id=body.submissionId, payload=payload)
        except ValueError as exc: raise HTTPException(status.HTTP_409_CONFLICT, str(exc)) from exc
        if persisted: return persisted
    if key in repo.followups:
        prior = repo.followups[key]
        if (prior.get("type"), prior.get("due"), prior.get("note")) != (body.type, body.due, body.note):
            raise HTTPException(status.HTTP_409_CONFLICT, "Submission already used.")
        return prior
    if body.type not in {"Appointment", "Reminder"}: raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Invalid follow-up type.")
    record = {"id": f"followup-{len(repo.followups)+1}", "bcn": bcn, "type": body.type, "due": body.due, "note": body.note, "status": "Open", "actor": user.name, "actorId": user.id, "createdAt": datetime.now(timezone.utc).isoformat()}
    repo.followups[key] = record; row["histories"].append({**record, "kind": "Follow-up"}); persist_followup(record)
    if activity_db is not None: activity_db.save_idempotent(actor_id=user.id, operation="followup", submission_id=body.submissionId, payload=payload, result=record)
    return record

@app.post("/customers/{bcn}/close")
def close_customer(bcn: str, body: LifecycleRequest, user: Annotated[User, Depends(current_user)]) -> Customer:
    local_key = (user.id, bcn, "close", body.submissionId)
    payload = f"{bcn}|close|{body.reasonId}"
    prior = repo.submissions.get(local_key)
    if prior:
        if prior["payload"] != payload: raise HTTPException(status.HTTP_409_CONFLICT, "Submission already used.")
        return Customer.model_validate(prior["result"])
    if activity_db is not None:
        try: persisted = activity_db.get_idempotent(actor_id=user.id, operation="close", submission_id=body.submissionId, payload=payload)
        except ValueError as exc: raise HTTPException(status.HTTP_409_CONFLICT, str(exc)) from exc
        if persisted: return get_customer(bcn, user)
    row = writable_customer(bcn, user); reason = repo.reasons.get(body.reasonId or "")
    if not reason or not reason["active"]: raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Choose an active closure reason.")
    timestamp = datetime.now(timezone.utc).isoformat(); row["status"] = "Closed"; row["version"] += 1; event = {"id": f"closure-{bcn}-{row['version']}", "bcn": bcn, "kind": "Closure", "reasonId": reason["id"], "reason": reason["label"], "actor": user.name, "actorId": user.id, "timestamp": timestamp}; row["histories"].append(event); persist_activity({**event, "text": reason["label"]}); append_audit(user.id, "Customer closed", bcn, {"reasonId": reason["id"], "reason": reason["label"]})
    if activity_db is not None:
        for stored in activity_db.followups(bcn):
            key = f"{user.id}:{bcn}:followup:{stored.id}"
            repo.followups.setdefault(key, {"id": stored.id, "bcn": stored.bcn, "type": stored.type, "due": stored.due.isoformat() if stored.due else None, "note": stored.note, "status": stored.status, "actorId": stored.actor_id, "createdAt": stored.created_at.isoformat()})
    for item in repo.followups.values():
        if item["bcn"] == bcn and item["status"] == "Open":
            item["status"] = "Cancelled"; item["updatedAt"] = datetime.now(timezone.utc).isoformat(); persist_followup(item)
            persist_activity({"id": f"activity-{uuid4()}", "bcn": bcn, "kind": "Follow-up cancel", "actorId": user.id, "text": "Customer closed"})
    if customer_db is not None: customer_db.save_operational(bcn=bcn, owner_id=row["ownerId"], status=row["status"], version=row["version"])
    result = Customer.model_validate(row)
    repo.submissions[local_key] = {"payload": payload, "result": result.model_dump()}
    if activity_db is not None: activity_db.save_idempotent(actor_id=user.id, operation="close", submission_id=body.submissionId, payload=payload, result=result.model_dump())
    return result

@app.post("/customers/{bcn}/reopen")
def reopen_customer(bcn: str, body: LifecycleRequest, user: Annotated[User, Depends(current_user)]) -> Customer:
    row = repo.customers.get(bcn)
    if row is None and customer_db is not None:
        stored = customer_db.get(bcn)
        if stored: row = {"bcn": stored.bcn, "name": stored.name, "ownerId": stored.owner_id, "ownerName": repo.users.get(stored.owner_id or "", {}).get("name"), "status": stored.status, "phones": customer_db.phones(bcn), "source": stored.source, "version": stored.version, "histories": []}; repo.customers[bcn] = row
    if not row: raise HTTPException(status.HTTP_404_NOT_FOUND, "Customer not found.")
    if user.role != "Admin" and row["ownerId"] != user.id: raise HTTPException(status.HTTP_403_FORBIDDEN, "Customer access denied.")
    payload = f"{bcn}|reopen"
    local_key = (user.id, bcn, "reopen", body.submissionId)
    prior = repo.submissions.get(local_key)
    if prior:
        if prior["payload"] != payload: raise HTTPException(status.HTTP_409_CONFLICT, "Submission already used.")
        return Customer.model_validate(prior["result"])
    if activity_db is not None:
        try: persisted = activity_db.get_idempotent(actor_id=user.id, operation="reopen", submission_id=body.submissionId, payload=payload)
        except ValueError as exc: raise HTTPException(status.HTTP_409_CONFLICT, str(exc)) from exc
        if persisted: return Customer.model_validate(row)
    timestamp = datetime.now(timezone.utc).isoformat(); row["status"] = "Open"; row["version"] += 1; event = {"id": f"reopen-{bcn}-{row['version']}", "bcn": bcn, "kind": "Reopen", "actor": user.name, "actorId": user.id, "timestamp": timestamp}; row["histories"].append(event); persist_activity({**event, "text": None}); append_audit(user.id, "Customer reopened", bcn, {})
    if customer_db is not None: customer_db.save_operational(bcn=bcn, owner_id=row["ownerId"], status=row["status"], version=row["version"])
    result = Customer.model_validate(row)
    repo.submissions[local_key] = {"payload": payload, "result": result.model_dump()}
    if activity_db is not None: activity_db.save_idempotent(actor_id=user.id, operation="reopen", submission_id=body.submissionId, payload=payload, result=result.model_dump())
    return result

def find_followup(bcn: str, followup_id: str, user: User) -> dict:
    row = writable_customer(bcn, user)
    item = next((value for value in repo.followups.values() if value["bcn"] == bcn and value["id"] == followup_id), None)
    if item is None and activity_db is not None:
        stored = next((value for value in activity_db.followups(bcn) if value.id == followup_id), None)
        if stored:
            item = {"id": stored.id, "bcn": stored.bcn, "type": stored.type, "due": stored.due.isoformat() if stored.due else None, "note": stored.note, "status": stored.status, "actorId": stored.actor_id, "createdAt": stored.created_at.isoformat()}; repo.followups[f"{user.id}:{bcn}:followup:{followup_id}"] = item
    if not item: raise HTTPException(status.HTTP_404_NOT_FOUND, "Follow-up not found.")
    return item

@app.patch("/customers/{bcn}/follow-ups/{followup_id}")
def update_followup(bcn: str, followup_id: str, body: FollowUpCreate, user: Annotated[User, Depends(current_user)]) -> dict:
    item = find_followup(bcn, followup_id, user)
    if item["status"] != "Open": raise HTTPException(status.HTTP_409_CONFLICT, "Follow-up is no longer open.")
    payload = f"{bcn}|{followup_id}|{body.type}|{body.due or ''}|{body.note}"
    if activity_db is not None:
        try: persisted = activity_db.get_idempotent(actor_id=user.id, operation="followup-edit", submission_id=body.submissionId, payload=payload)
        except ValueError as exc: raise HTTPException(status.HTTP_409_CONFLICT, str(exc)) from exc
        if persisted: return persisted
    item.update(type=body.type, due=body.due, note=body.note, version=item.get("version", 0) + 1, updatedAt=datetime.now(timezone.utc).isoformat()); persist_followup(item); persist_activity({"id": f"activity-{uuid4()}", "bcn": bcn, "kind": "Follow-up edit", "actorId": user.id, "text": item.get("note")})
    if activity_db is not None: activity_db.save_idempotent(actor_id=user.id, operation="followup-edit", submission_id=body.submissionId, payload=payload, result=item)
    return item

@app.post("/customers/{bcn}/follow-ups/{followup_id}/cancel")
def cancel_followup(bcn: str, followup_id: str, user: Annotated[User, Depends(current_user)]) -> dict:
    item = find_followup(bcn, followup_id, user)
    if item["status"] == "Cancelled": return item
    if item["status"] == "Completed": raise HTTPException(status.HTTP_409_CONFLICT, "Follow-up is completed.")
    item["status"] = "Cancelled"; item["updatedAt"] = datetime.now(timezone.utc).isoformat(); persist_followup(item); persist_activity({"id": f"activity-{uuid4()}", "bcn": bcn, "kind": "Follow-up cancel", "actorId": user.id, "text": None}); return item

@app.delete("/customers/{bcn}/follow-ups/{followup_id}")
def cancel_followup_contract(bcn: str, followup_id: str, user: Annotated[User, Depends(current_user)]) -> dict:
    return cancel_followup(bcn, followup_id, user)

@app.post("/customers/{bcn}/follow-ups/{followup_id}/complete")
def complete_followup(bcn: str, followup_id: str, body: InteractionCreate, user: Annotated[User, Depends(current_user)]) -> dict:
    item = find_followup(bcn, followup_id, user)
    if item["status"] == "Completed": return item
    if item["status"] != "Open": raise HTTPException(status.HTTP_409_CONFLICT, "Follow-up is not open.")
    interaction = create_interaction(bcn, body, user, persist=activity_db is None); item["status"] = "Completed"; item["interactionId"] = interaction["id"]; item["updatedAt"] = datetime.now(timezone.utc).isoformat()
    if activity_db is not None: activity_db.complete_followup(item, interaction)
    else: persist_followup(item)
    return item


@app.post("/admin/assignments/manual/{bcn}", response_model=Customer)
def assign_customer(bcn: str, body: AssignmentRequest, user: Annotated[User, Depends(current_user)]) -> Customer:
    if user.role != "Admin": raise HTTPException(status.HTTP_403_FORBIDDEN, "Admin access required.")
    row = repo.customers.get(bcn)
    if row is None and customer_db is not None:
        stored = customer_db.get(bcn)
        if stored: row = {"bcn": stored.bcn, "name": stored.name, "ownerId": stored.owner_id, "ownerName": repo.users.get(stored.owner_id or "", {}).get("name"), "status": stored.status, "phones": customer_db.phones(bcn), "source": stored.source, "version": stored.version, "histories": []}; repo.customers[bcn] = row
    if not row: raise HTTPException(status.HTTP_404_NOT_FOUND, "Customer not found.")
    if body.ownerId and not any(item["id"] == body.ownerId and item["role"] == "Sales" and item["active"] for item in repo.users.values()): raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Owner must be an active Sales user.")
    payload = f"{bcn}|{body.ownerId}|{body.expectedVersion}"
    if activity_db is not None:
        try: persisted = activity_db.get_idempotent(actor_id=user.id, operation="assignment", submission_id=body.submissionId, payload=payload)
        except ValueError as exc: raise HTTPException(status.HTTP_409_CONFLICT, str(exc)) from exc
        if persisted: return Customer.model_validate(row)
    key = (bcn, body.submissionId)
    prior = repo.submissions.get(key)
    if prior:
        if prior["ownerId"] != body.ownerId or prior["expectedVersion"] != body.expectedVersion:
            raise HTTPException(status.HTTP_409_CONFLICT, "Submission already used.")
        return Customer.model_validate(row)
    if body.expectedVersion is not None and body.expectedVersion != row["version"]:
        raise HTTPException(status.HTTP_409_CONFLICT, "Customer version is stale.")
    if row["ownerId"] == body.ownerId:
        repo.submissions[key] = {"ownerId": body.ownerId, "expectedVersion": body.expectedVersion}
        if activity_db is not None: activity_db.save_idempotent(actor_id=user.id, operation="assignment", submission_id=body.submissionId, payload=payload, result=Customer.model_validate(row).model_dump())
        return Customer.model_validate(row)
    with repo.transaction():
        old = row["ownerId"]; owner = repo.users.get(body.ownerId) if body.ownerId else None
        row["ownerId"] = body.ownerId; row["ownerName"] = owner["name"] if owner else None
        row["version"] += 1
        row["histories"].append({"kind": "Assignment", "actor": user.name, "actorId": user.id, "oldOwner": old, "newOwner": body.ownerId, "reason": "Manual assignment", "timestamp": datetime.now(timezone.utc).isoformat()})
        if assignment_db is not None: assignment_db.append_assignment(bcn=bcn, actor_id=user.id, old_owner_id=old, new_owner_id=body.ownerId, reason="Manual assignment")
        if customer_db is not None: customer_db.save_operational(bcn=bcn, owner_id=row["ownerId"], status=row["status"], version=row["version"])
        append_audit(user.id, "Customer assigned", bcn, {"oldOwner": old, "newOwner": body.ownerId})
        repo.submissions[key] = {"ownerId": body.ownerId, "expectedVersion": body.expectedVersion}
    result = Customer.model_validate(row)
    if activity_db is not None: activity_db.save_idempotent(actor_id=user.id, operation="assignment", submission_id=body.submissionId, payload=payload, result=result.model_dump())
    return result


class Provision(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    email: str = Field(min_length=3, max_length=254)
    role: str = "Admin"
    password: str = Field(min_length=12, max_length=128)


class ResetPassword(BaseModel):
    password: str = Field(min_length=12, max_length=128)

class UserDraft(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    email: str = Field(min_length=3, max_length=254)
    role: str = "Sales"
    password: str = Field(min_length=12, max_length=128)

class UserPatch(BaseModel):
    name: str | None = Field(default=None, min_length=1, max_length=120)
    role: str | None = None
    active: bool | None = None

class ClosureReason(BaseModel):
    id: str
    label: str
    active: bool = True

class ClosureReasonDraft(BaseModel):
    label: str = Field(min_length=1, max_length=120)

class ClosureReasonPatch(BaseModel):
    label: str | None = Field(default=None, min_length=1, max_length=120)
    active: bool | None = None

class AssignmentRuleDraft(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    ownerId: str
    active: bool = True

class AssignmentRunRequest(BaseModel):
    scope: str = "unassigned"
    submissionId: str = Field(min_length=1)


def provision_user(data: Provision) -> User:
    data = Provision.model_validate({key: getattr(data, key) for key in ("name", "email", "role", "password")})
    email = safe_email(data.email)
    if auth_db is not None:
        if auth_db.user_by_email(email): raise ValueError("normalized identity already exists")
        if data.role not in {"Admin", "Sales"}: raise ValueError("invalid role")
        row = auth_db.create_user(user_id=f"user-{uuid4()}", name=data.name, email=email, role=data.role, password_hash=password_hash.hash(data.password))
        return User(id=row.id, name=row.name, email=row.email, role=row.role, active=row.active)
    if any(item["email"] == email for item in repo.users.values()):
        raise ValueError("normalized identity already exists")
    if data.role not in {"Admin", "Sales"}:
        raise ValueError("invalid role")
    record = {"id": f"user-{len(repo.users) + 1}", "name": data.name, "email": email, "role": data.role, "active": True, "password": password_hash.hash(data.password)}
    repo.users[record["id"]] = record
    return User.model_validate(record)

def admin_user(user: Annotated[User, Depends(current_user)]) -> User:
    if user.role != "Admin": raise HTTPException(status.HTTP_403_FORBIDDEN, "Admin access required.")
    return user

@app.get("/admin/users", response_model=list[User])
def list_users(_: Annotated[User, Depends(admin_user)]) -> list[User]:
    return [User.model_validate(item) for item in repo.users.values()]

@app.post("/admin/users", response_model=User, status_code=201)
def create_user(data: UserDraft, _: Annotated[User, Depends(admin_user)]) -> User:
    try: return provision_user(Provision(**data.model_dump()))
    except ValueError as exc: raise HTTPException(status.HTTP_409_CONFLICT, str(exc)) from exc

@app.patch("/admin/users/{user_id}", response_model=User)
def update_user(user_id: str, patch: UserPatch, actor: Annotated[User, Depends(admin_user)]) -> User:
    record = repo.users.get(user_id)
    if not record: raise HTTPException(status.HTTP_404_NOT_FOUND, "User not found.")
    changes = patch.model_dump(exclude_none=True)
    if user_id == actor.id and (changes.get("active") is False or changes.get("role") == "Sales"): raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Cannot disable or demote yourself.")
    if changes.get("role") not in {None, "Admin", "Sales"}: raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Invalid role.")
    if changes.get("email"): changes["email"] = safe_email(changes["email"])
    prior_active = record["active"]; prior_role = record["role"]; record.update(changes)
    if auth_db is not None:
        auth_db.update_user(user_id, changes)
    append_audit(actor.id, "User changed", user_id, {key: value for key, value in changes.items() if key in {"name", "role", "active"}})
    if prior_role != record["role"] or prior_active != record["active"]:
        for token, (owner, _) in list(repo.sessions.items()):
            if owner == user_id: repo.sessions.pop(token, None)
    if prior_active and (record["active"] is False or record["role"] != "Sales"):
        for row in repo.customers.values():
            if row["ownerId"] == user_id and row["status"] == "Open": row.update(ownerId=None, ownerName=None, version=row["version"] + 1)
        if customer_db is not None: customer_db.release_open_owner(user_id)
    return User.model_validate(record)

@app.get("/admin/closure-reasons", response_model=list[ClosureReason])
def list_reasons(_: Annotated[User, Depends(admin_user)]) -> list[ClosureReason]:
    return [ClosureReason.model_validate(item) for item in repo.reasons.values()]

@app.post("/admin/closure-reasons", response_model=ClosureReason, status_code=201)
def create_reason(data: ClosureReasonDraft, _: Annotated[User, Depends(admin_user)]) -> ClosureReason:
    if any(item["label"].casefold() == data.label.strip().casefold() for item in repo.reasons.values()): raise HTTPException(status.HTTP_409_CONFLICT, "Reason already exists.")
    reason = {"id": f"closure-{len(repo.reasons) + 1}", "label": data.label.strip(), "active": True}; repo.reasons[reason["id"]] = reason; persist_reason(reason); append_audit(_.id, "Closure reason created", reason["id"], {"label": reason["label"]}); return ClosureReason.model_validate(reason)

@app.patch("/admin/closure-reasons/{reason_id}", response_model=ClosureReason)
def update_reason(reason_id: str, patch: ClosureReasonPatch, _: Annotated[User, Depends(admin_user)]) -> ClosureReason:
    reason = repo.reasons.get(reason_id)
    if not reason: raise HTTPException(status.HTTP_404_NOT_FOUND, "Reason not found.")
    label = patch.label.strip() if patch.label else reason["label"]
    if any(item["id"] != reason_id and item["label"].casefold() == label.casefold() for item in repo.reasons.values()): raise HTTPException(status.HTTP_409_CONFLICT, "Reason already exists.")
    reason.update(label=label, **({"active": patch.active} if patch.active is not None else {})); persist_reason(reason); append_audit(_.id, "Closure reason changed", reason_id, {"label": reason["label"], "active": reason["active"]}); return ClosureReason.model_validate(reason)

@app.post("/admin/imports", status_code=201)
async def import_customers(file: UploadFile = File(...), submission_id: str = Query(..., min_length=1), user: Annotated[User, Depends(admin_user)] = None) -> dict:
    if not file.filename or not file.filename.casefold().endswith(".xlsx"):
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Upload an .xlsx workbook.")
    if submission_id in repo.imports: return repo.imports[submission_id]
    if customer_db is not None:
        persisted_job = customer_db.import_job(user.id, submission_id)
        if persisted_job:
            repo.imports[submission_id] = persisted_job
            return persisted_job
    payload = await file.read()
    if len(payload) > 10 * 1024 * 1024: raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, "Workbook is too large.")
    if activity_db is not None:
        try:
            persisted = activity_db.get_idempotent(actor_id=user.id, operation="import", submission_id=submission_id, payload=payload)
        except ValueError as exc:
            raise HTTPException(status.HTTP_409_CONFLICT, str(exc)) from exc
        if persisted: return persisted
    try:
        from io import BytesIO
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True, keep_links=False)
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows, ())]
        if not headers or "bcn" not in {item.casefold() for item in headers}: raise ValueError("Missing bcn header")
        bcn_index = next(index for index, value in enumerate(headers) if value.casefold() == "bcn")
        name_index = next((index for index, value in enumerate(headers) if value.casefold() in {"customer_name", "name"}), None)
        phone_index = next((index for index, value in enumerate(headers) if value.casefold() == "phone"), None)
        errors = []; created = updated = 0; nonblank_rows = 0; expanded_bytes = 0; source_rows = []; seen_bcns: set[str] = set()
        with repo.transaction():
            for row_number, values in enumerate(rows, 2):
                expanded_bytes += sum(len(str(value).encode("utf-8")) for value in values if value is not None)
                if expanded_bytes > 100 * 1024 * 1024: raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, "Workbook expands beyond the processing limit.")
                if not any(value not in (None, "") for value in values): continue
                nonblank_rows += 1
                if nonblank_rows > 10_000: raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, "Workbook has too many rows.")
                raw = values[bcn_index] if bcn_index < len(values) else None
                bcn = str(raw).strip() if raw is not None else ""
                if not bcn or not bcn.isdigit(): errors.append({"row": row_number, "field": "bcn", "reason": "Invalid bcn"}); continue
                bcn = bcn.zfill(6); name = str(values[name_index]).strip() if name_index is not None and name_index < len(values) and values[name_index] is not None else ""
                if len(bcn) > 128: errors.append({"row": row_number, "field": "bcn", "reason": "BCN exceeds 128 characters"}); continue
                if not name: errors.append({"row": row_number, "field": "customer_name", "reason": "Customer name is required"}); continue
                phone = str(values[phone_index]).strip() if phone_index is not None and phone_index < len(values) and values[phone_index] is not None else None
                if bcn in seen_bcns: errors.append({"row": row_number, "field": "bcn", "reason": "Duplicate bcn"}); continue
                seen_bcns.add(bcn)
                source = {header: import_value(values[index] if index < len(values) else None) for index, header in enumerate(headers) if header}
                source["customer_name"] = name or bcn
                source_rows.append({"bcn": bcn, "name": name, "source": source, "primary_phone": phone})
                if bcn in repo.customers:
                    record = repo.customers[bcn]
                    record["name"] = name or record["name"]
                    phones = record.setdefault("phones", [])
                    if phone:
                        if phones: phones[0] = phone
                        else: phones.append(phone)
                    elif phones:
                        phones.pop(0)
                    record["source"] = source
                    updated += 1
                else:
                    repo.customers[bcn] = {"bcn": bcn, "name": name or bcn, "ownerId": None, "ownerName": None, "status": "Open", "phones": [phone] if phone else [], "source": source, "version": 0, "histories": []}; created += 1
        result = {"jobId": f"import-{len(repo.imports)+1}", "submissionId": submission_id, "filename": file.filename, "completedAt": datetime.now(timezone.utc).isoformat(), "status": "Partial" if errors else "Completed", "created": created, "updated": updated, "errors": errors, "errorRows": len(errors), "processed": created + updated + len(errors), "actorId": user.id}
        if customer_db is not None:
            customer_db.ingest_sources(source_rows, result)
        repo.imports[submission_id] = result
        if activity_db is not None: activity_db.save_idempotent(actor_id=user.id, operation="import", submission_id=submission_id, payload=payload, result=result)
        append_audit(user.id, "Import completed", result["jobId"], {"created": created, "updated": updated, "errors": len(errors)}); return result
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Workbook could not be processed.") from exc

@app.get("/admin/assignment-rules")
def list_assignment_rules(_: Annotated[User, Depends(admin_user)]) -> list[dict]:
    return repo.rules

@app.post("/admin/assignment-rules", status_code=201)
def create_assignment_rule(body: AssignmentRuleDraft, _: Annotated[User, Depends(admin_user)]) -> dict:
    if any(item["name"].casefold() == body.name.strip().casefold() for item in repo.rules): raise HTTPException(status.HTTP_409_CONFLICT, "Rule already exists.")
    owner = repo.users.get(body.ownerId)
    if not owner or owner["role"] != "Sales" or not owner["active"]: raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Owner must be active Sales.")
    rule = {"id": f"rule-{len(repo.rules)+1}", "name": body.name.strip(), "ownerId": body.ownerId, "active": body.active, "order": len(repo.rules)+1}; repo.rules.append(rule); repo.assignment_version += 1
    if assignment_db is not None:
        assignment_db.create_rule(rule_id=rule["id"], name=rule["name"], position=rule["order"], actor_id=_.id, owner_id=rule["ownerId"])
        assignment_db.set_setting("assignment_version", {"value": repo.assignment_version})
    return rule

@app.patch("/admin/assignment-rules/{rule_id}")
def update_assignment_rule(rule_id: str, patch: dict, _: Annotated[User, Depends(admin_user)]) -> dict:
    rule = next((item for item in repo.rules if item["id"] == rule_id), None)
    if not rule: raise HTTPException(status.HTTP_404_NOT_FOUND, "Rule not found.")
    if "version" in patch and patch["version"] != repo.assignment_version: raise HTTPException(status.HTTP_409_CONFLICT, "Assignment configuration is stale.")
    if "name" in patch and any(item["id"] != rule_id and item["name"].casefold() == str(patch["name"]).strip().casefold() for item in repo.rules): raise HTTPException(status.HTTP_409_CONFLICT, "Rule already exists.")
    for key in ("name", "active", "order"):
        if key in patch: rule[key] = patch[key]
    repo.rules.sort(key=lambda item: item["order"]); repo.assignment_version += 1
    if assignment_db is not None:
        assignment_db.update_rule(rule_id, {"name": rule["name"], "active": rule["active"], "position": rule["order"], "owner_id": rule["ownerId"]}, _.id)
        assignment_db.set_setting("assignment_version", {"value": repo.assignment_version})
    return rule

@app.get("/admin/assignment-fallback")
def assignment_fallback(_: Annotated[User, Depends(admin_user)]) -> list[str]:
    if assignment_db is not None:
        stored = assignment_db.get_setting("fallback_sales")
        if stored is not None: repo.fallback_sales = list(stored.get("ids", []))
    return repo.fallback_sales

@app.put("/admin/assignment-fallback")
def set_assignment_fallback(ids: Annotated[list[str], Body()], _: Annotated[User, Depends(admin_user)]) -> list[str]:
    valid = {item["id"] for item in repo.users.values() if item["role"] == "Sales" and item["active"]}
    if any(item not in valid for item in ids): raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Fallback members must be active Sales users.")
    repo.fallback_sales = list(dict.fromkeys(ids)); repo.assignment_version += 1
    if assignment_db is not None:
        assignment_db.set_setting("fallback_sales", {"ids": repo.fallback_sales})
        assignment_db.set_setting("assignment_version", {"value": repo.assignment_version})
    return repo.fallback_sales

@app.get("/admin/assignment-version")
def get_assignment_version(_: Annotated[User, Depends(admin_user)]) -> int:
    return repo.assignment_version

@app.post("/admin/assignment-runs")
def run_assignment(body: AssignmentRunRequest, _: Annotated[User, Depends(admin_user)]) -> dict:
    if assignment_db is not None:
        persisted = assignment_db.get_run(_.id, body.submissionId)
        if persisted: return persisted
    run_key = (_.id, body.submissionId)
    if run_key in repo.assignment_runs: return repo.assignment_runs[run_key]
    if body.scope not in {"unassigned", "all-open"}: raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Invalid assignment scope.")
    source_rows = readable_rows()
    candidates = [row for row in source_rows if row["status"] == "Open" and (body.scope == "all-open" or row["ownerId"] is None)]
    counts = {item["id"]: sum(1 for row in source_rows if row["ownerId"] == item["id"] and row["status"] == "Open") for item in repo.users.values() if item["role"] == "Sales" and item["active"]}
    assigned = 0
    for row in sorted(candidates, key=lambda item: item["bcn"]):
        eligible = [rule for rule in sorted(repo.rules, key=lambda item: item["order"]) if rule["active"] and rule["ownerId"] in counts]
        if not eligible: continue
        owner = min((rule["ownerId"] for rule in eligible), key=lambda user_id: (counts[user_id], user_id))
        old_owner = row["ownerId"]; row.update(ownerId=owner, ownerName=repo.users[owner]["name"], version=row["version"] + 1); counts[owner] += 1; assigned += 1
        if customer_db is not None:
            customer_db.save_operational(bcn=row["bcn"], owner_id=owner, status=row["status"], version=row["version"])
            if assignment_db is not None: assignment_db.append_assignment(bcn=row["bcn"], actor_id=_.id, old_owner_id=old_owner, new_owner_id=owner, reason="Bulk assignment")
            append_audit(_.id, "Customer assigned", row["bcn"], {"oldOwner": old_owner, "newOwner": owner, "source": "bulk"})
    result = {"submissionId": body.submissionId, "scope": body.scope, "candidates": len(candidates), "assigned": assigned, "skipped": len(candidates) - assigned}
    repo.assignment_runs[run_key] = result
    if assignment_db is not None: assignment_db.save_run(actor_id=_.id, submission_id=body.submissionId, scope=body.scope, result=result)
    return result

@app.get("/admin/assignments")
def list_assignment_rules_contract(user: Annotated[User, Depends(admin_user)]) -> list[dict]:
    return list_assignment_rules(user)

@app.post("/admin/assignments/run")
def run_assignment_contract(body: AssignmentRunRequest, user: Annotated[User, Depends(admin_user)]) -> dict:
    return run_assignment(body, user)

@app.get("/sales/workload")
def sales_workload(user: Annotated[User, Depends(current_user)]) -> dict:
    if user.role != "Sales": raise HTTPException(status.HTTP_403_FORBIDDEN, "Sales access required.")
    if customer_db is not None and activity_db is not None:
        owned = customer_db.open_owned(user.id); bcns = [row.bcn for row, _ in owned]; followups = activity_db.followup_summary(bcns, datetime.now(timezone.utc).date().isoformat()); interactions = activity_db.interaction_counts(bcns)
        return {"ownerId": user.id, "totalOpen": len(owned), "neverContacted": sum(1 for bcn in bcns if not interactions.get(bcn)), "pendingFollowUps": sum(item["open"] for item in followups.values())}
    owned = [row for row in readable_rows() if row["ownerId"] == user.id and row["status"] == "Open"]
    followups = readable_followups(); return {"ownerId": user.id, "totalOpen": len(owned), "neverContacted": sum(1 for row in owned if not any(item.get("kind") == "Interaction" and not item.get("deleted") for item in row["histories"])), "pendingFollowUps": sum(1 for item in followups if item["bcn"] in {row["bcn"] for row in owned} and item["status"] == "Open")}

@app.get("/workload")
def workload_contract(user: Annotated[User, Depends(current_user)]) -> dict:
    if user.role != "Sales": raise HTTPException(status.HTTP_403_FORBIDDEN, "Sales access required.")
    if customer_db is not None and activity_db is not None:
        today = datetime.now(timezone.utc).date().isoformat(); owned = customer_db.open_owned(user.id); bcns = [row.bcn for row, _ in owned]; followups = activity_db.followup_summary(bcns, today); interactions = activity_db.interaction_counts(bcns)
        counts = {bucket: 0 for bucket in ("overdue", "today", "undated", "never-contacted", "other")}; customers = []
        for row, phones in owned:
            summary = followups[row.bcn]; bucket = "overdue" if summary["overdue"] else "today" if summary["today"] else "undated" if summary["undated"] else "never-contacted" if not interactions.get(row.bcn) else "other"; counts[bucket] += 1; customers.append(Customer.model_validate({"bcn": row.bcn, "name": row.name, "ownerId": row.owner_id, "ownerName": repo.users.get(row.owner_id or "", {}).get("name"), "status": row.status, "phones": phones, "source": row.source, "version": row.version, "histories": []}).model_dump() | {"workloadBucket": bucket, "relevantDue": None})
        return {"asOf": datetime.now(timezone.utc).isoformat(), "today": today, "counts": counts, "customers": customers}
    owned = [row for row in readable_rows() if row["ownerId"] == user.id and row["status"] == "Open"]
    return {"asOf": datetime.now(timezone.utc).isoformat(), "today": datetime.now(timezone.utc).date().isoformat(), "counts": {"overdue": 0, "today": 0, "undated": 0, "never-contacted": sum(1 for row in owned if not any(event.get("kind") == "Interaction" and not event.get("deleted") for event in row["histories"])), "other": 0}, "customers": [Customer.model_validate(row).model_dump() | {"workloadBucket": None, "relevantDue": None} for row in owned]}

@app.get("/admin/reports")
def admin_reports(start: str | None = None, end: str | None = None, _: Annotated[User, Depends(admin_user)] = None) -> dict:
    if start and end and start > end: raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Start date must not be after end date.")
    if customer_db is not None and activity_db is not None:
        metrics, daily = activity_db.report_interactions(start, end); owners: dict[str, dict] = {}; customer_rows = customer_db.all()
        for owner_id, status_value, count in customer_db.owner_counts():
            key = owner_id or "unassigned"; row = owners.setdefault(key, {"ownerId": owner_id, "owner": repo.users.get(key, {}).get("name", "Unassigned"), "open": 0, "closed": 0, "neverContacted": 0, "attempts": 0, "contacts": 0, "pendingFollowUps": 0}); row["open" if status_value == "Open" else "closed"] += count
        all_counts, _ = activity_db.report_interactions(); pending = {item.bcn: item for item in activity_db.all_followups() if item.status == "Open"}
        for owner_id, row in ((key, value) for key, value in owners.items()):
            bcns = [item.bcn for item in customer_rows if (item.owner_id or "unassigned") == owner_id]
            row["neverContacted"] = sum(1 for bcn in bcns if bcn not in all_counts); row["attempts"] = sum(all_counts.get(bcn, {}).get("attempts", 0) for bcn in bcns); row["contacts"] = sum(all_counts.get(bcn, {}).get("contacts", 0) for bcn in bcns); row["pendingFollowUps"] = sum(1 for item in pending.values() if item.bcn in bcns); row["contactRate"] = row["contacts"] / row["attempts"] * 100 if row["attempts"] else None
        all_followups = activity_db.all_followups(); followups = {"overdue": 0, "today": 0, "undated": 0, "completed": sum(1 for item in all_followups if item.status == "Completed")}; today = datetime.now(timezone.utc).date()
        for item in all_followups:
            if item.status != "Open": continue
            if item.due is None: followups["undated"] += 1
            elif item.due.date() < today: followups["overdue"] += 1
            elif item.due.date() == today: followups["today"] += 1
        return {"owners": list(owners.values()), "daily": sorted(daily.values(), key=lambda item: item["date"]), "closureReasons": list(repo.reasons.values()), "followUps": followups}
    owners: dict[str, dict] = {}; daily: dict[str, dict] = {}; all_followups = readable_followups()
    for row in readable_rows():
        key = row["ownerId"] or "unassigned"; bucket = owners.setdefault(key, {"ownerId": row["ownerId"], "owner": repo.users.get(key, {}).get("name", "Unassigned"), "open": 0, "closed": 0, "neverContacted": 0, "attempts": 0, "contacts": 0, "pendingFollowUps": 0})
        bucket["open" if row["status"] == "Open" else "closed"] += 1
        if row["status"] == "Open" and not any(event.get("kind") == "Interaction" and not event.get("deleted") for event in row["histories"]): bucket["neverContacted"] += 1
        bucket["pendingFollowUps"] += sum(1 for item in all_followups if item["bcn"] == row["bcn"] and item["status"] == "Open")
        for event in row["histories"]:
            if event.get("deleted") or event.get("kind") != "Interaction": continue
            date = event.get("timestamp", "")[:10]
            if start and date < start or end and date > end: continue
            bucket["attempts" if event.get("outcome") == "Attempt" else "contacts"] += 1
            daily.setdefault(date, {"date": date, "attempts": 0, "contacts": 0})["attempts" if event.get("outcome") == "Attempt" else "contacts"] += 1
    for bucket in owners.values(): bucket["contactRate"] = bucket["contacts"] / bucket["attempts"] * 100 if bucket["attempts"] else None
    followups = {"overdue": 0, "today": 0, "undated": 0, "completed": sum(1 for item in all_followups if item["status"] == "Completed")}
    for item in all_followups:
        if item["status"] != "Open": continue
        if not item.get("due"): followups["undated"] += 1
    return {"owners": list(owners.values()), "daily": sorted(daily.values(), key=lambda item: item["date"]), "closureReasons": list(repo.reasons.values()), "followUps": followups}

@app.get("/admin/audit")
def admin_audit(actor: str | None = None, action: str | None = None, bcn: str | None = None, start: str | None = None, end: str | None = None, page: int = Query(1, ge=1), page_size: int = Query(25, ge=1, le=100), _: Annotated[User, Depends(admin_user)] = None) -> dict:
    if assignment_db is not None:
        rows, total = assignment_db.audit(page, page_size, actor=actor, action=action, target=bcn, start=start, end=end)
        return {"items": [{"id": str(item.id), "actor": item.actor_id or "", "actorId": item.actor_id or "", "action": item.action, "target": item.target, "timestamp": item.created_at.isoformat(), "details": item.details} for item in rows], "page": page, "page_size": page_size, "total": total}
    events = []
    for row in repo.customers.values():
        for index, event in enumerate(row["histories"]):
            item = {"id": event.get("id", f"{row['bcn']}-{index}"), "actor": event.get("actor", ""), "actorId": event.get("actorId", ""), "action": event.get("kind", ""), "target": row["bcn"], "timestamp": event.get("timestamp", ""), "details": {key: value for key, value in event.items() if key not in {"text", "note", "outcome"}}}
            date = item["timestamp"][:10]
            if (bcn and bcn != row["bcn"]) or (actor and actor.casefold() not in item["actor"].casefold()) or (action and action.casefold() not in item["action"].casefold()) or (start and date < start) or (end and date > end): continue
            events.append(item)
    events.sort(key=lambda item: item.get("timestamp", ""), reverse=True); start = (page - 1) * page_size
    return {"items": events[start:start + page_size], "page": page, "page_size": page_size, "total": len(events)}


def operator_provision(data: Provision) -> User:
    if repo.users or (auth_db is not None and auth_db.all_users()):
        raise HTTPException(status.HTTP_409_CONFLICT, "Initial Admin already provisioned.")
    try:
        return provision_user(data)
    except ValueError as exc:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, str(exc)) from exc


def operator_reset_password(user_id: str, data: ResetPassword) -> User:
    if auth_db is not None:
        row = auth_db.reset_password(user_id, password_hash.hash(data.password))
        if not row: raise HTTPException(status.HTTP_404_NOT_FOUND, "User not found.")
        return User(id=row.id, name=row.name, email=row.email, role=row.role, active=row.active)
    record = repo.users.get(user_id)
    if not record:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "User not found.")
    record["password"] = password_hash.hash(data.password)
    for token, (owner, _) in list(repo.sessions.items()):
        if owner == user_id: repo.sessions.pop(token, None)
    return User.model_validate(record)


if storage_mode == "memory":
    app.post("/operator/provision", response_model=User, include_in_schema=False)(operator_provision)
    app.post("/operator/reset-password/{user_id}", response_model=User, include_in_schema=False)(operator_reset_password)
