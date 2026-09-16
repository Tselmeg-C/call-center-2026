from datetime import datetime, timedelta, timezone
from io import BytesIO
from time import perf_counter
import logging
import os
import subprocess
from openpyxl import Workbook

from fastapi.testclient import TestClient
from sqlalchemy.orm import Session
from sqlalchemy import text

from ..main import app, password_hash, repo, provision_user
from ..storage import mode
from ..db_auth import AuthDatabase, UserRow, SessionRow, digest
from ..db_customers import CustomerDatabase
from ..db_assignment import AssignmentDatabase
from ..db_activity import ActivityDatabase


def test_login_logout_and_generic_failure() -> None:
    repo.reset()
    provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    client = TestClient(app, base_url="http://localhost")
    assert client.post("/session/login", json={"email": "unknown@example.test", "password": "wrong password"}).status_code == 401
    login = client.post("/session/login", json={"email": " ADMIN@example.test ", "password": "correct horse battery staple"})
    assert login.status_code == 200 and login.json()["role"] == "Admin"
    assert client.get("/session/me").status_code == 200
    assert client.post("/session/logout", headers={"origin": "http://localhost:3000"}).status_code == 204
    assert client.get("/session/me").status_code == 401

def test_storage_selection_is_explicit(monkeypatch) -> None:
    monkeypatch.setenv("CALL_CENTER_STORAGE", "postgres"); monkeypatch.delenv("DATABASE_URL", raising=False)
    try: mode(); assert False
    except RuntimeError as exc: assert "DATABASE_URL" in str(exc)

def test_database_url_pins_psycopg_dialect(monkeypatch) -> None:
    from ..storage import database_url
    # #27: Railway's managed Postgres plugin (and most standard providers) hand back a bare
    # postgresql:// URL. SQLAlchemy's default dialect for that scheme is psycopg2, which is not
    # installed (apps/api/requirements.txt only installs psycopg v3) -- every boot's `alembic
    # upgrade head` failed against a real deployment until this normalization existed.
    monkeypatch.setenv("DATABASE_URL", "postgresql://user:pass@host:5432/db")
    assert database_url() == "postgresql+psycopg://user:pass@host:5432/db"
    monkeypatch.setenv("DATABASE_URL", "postgres://user:pass@host:5432/db")
    assert database_url() == "postgresql+psycopg://user:pass@host:5432/db"
    monkeypatch.setenv("DATABASE_URL", "postgresql+psycopg://user:pass@host:5432/db")
    assert database_url() == "postgresql+psycopg://user:pass@host:5432/db"
    monkeypatch.setenv("DATABASE_URL", "sqlite+pysqlite:///:memory:")
    assert database_url() == "sqlite+pysqlite:///:memory:"

def test_login_failure_throttle_is_generic() -> None:
    repo.reset(); client = TestClient(app, base_url="http://localhost")
    for _ in range(5): assert client.post("/session/login", json={"email": "unknown@example.test", "password": "wrong password"}).status_code == 401
    limited = client.post("/session/login", json={"email": "unknown@example.test", "password": "wrong password"})
    assert limited.status_code == 429 and limited.headers.get("retry-after") == "900" and "unknown@example.test" not in limited.text

def test_login_failure_ip_throttle_limits_identity_spraying() -> None:
    repo.reset(); client = TestClient(app, base_url="http://localhost")
    for index in range(50): assert client.post("/session/login", json={"email": f"unknown-{index}@example.test", "password": "wrong password"}).status_code == 401
    limited = client.post("/session/login", json={"email": "unknown-final@example.test", "password": "wrong password"})
    assert limited.status_code == 429 and limited.headers.get("retry-after") == "900"

def test_start_rejects_multi_worker_rate_limit_topology() -> None:
    env = os.environ.copy(); env["WEB_CONCURRENCY"] = "2"; env["CALL_CENTER_STORAGE"] = "memory"
    result = subprocess.run(["sh", "apps/api/start.sh"], env=env, capture_output=True, text=True)
    assert result.returncode == 78 and "WEB_CONCURRENCY must be 1" in result.stderr

def test_request_id_rejects_malformed_client_value() -> None:
    response = TestClient(app, base_url="http://localhost").get("/health/live", headers={"x-request-id": "bad\nvalue"})
    assert response.status_code == 200 and "\n" not in response.headers["x-request-id"] and len(response.headers["x-request-id"]) > 10

def test_request_id_injection_cannot_add_log_lines(caplog) -> None:
    # A client-supplied id containing newlines/format-string-shaped content must not: (1) reach
    # the log line verbatim (it fails the id regex, so a server UUID is used instead), or
    # (2) otherwise cause more than the one structured log line the middleware always emits.
    with caplog.at_level("INFO", logger="call-center.api"):
        response = TestClient(app, base_url="http://localhost").get(
            "/health/live", headers={"x-request-id": "id\nrequest id=fake method=GET route=/admin status=200 duration_ms=0.0 error=none"}
        )
    assert response.status_code == 200
    request_records = [record for record in caplog.records if record.message.startswith("request id=")]
    assert len(request_records) == 1
    assert "\n" not in request_records[0].message
    assert "fake" not in request_records[0].message

def test_login_never_logs_credentials_or_cookie(caplog) -> None:
    repo.reset()
    provision_user(type("P", (), {"name": "Admin", "email": "creds@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    with caplog.at_level("INFO"):
        client = TestClient(app, base_url="http://localhost")
        login = client.post("/session/login", json={"email": "creds@example.test", "password": "correct horse battery staple"})
        token = client.cookies.get("call_center_session")
    assert login.status_code == 200
    assert "correct horse battery staple" not in caplog.text
    assert token is not None and token not in caplog.text

def test_client_ip_prefers_leftmost_forwarded_for_over_tcp_peer() -> None:
    from ..main import client_ip
    from unittest.mock import Mock
    request = Mock()
    request.headers = {"x-forwarded-for": "203.0.113.7, 10.0.0.5, 10.0.0.1"}
    assert client_ip(request) == "203.0.113.7"
    request.headers = {}
    request.client = Mock(host="10.0.0.9")
    assert client_ip(request) == "10.0.0.9"
    request.client = None
    assert client_ip(request) == "unknown"

def test_request_log_reaches_stdout_without_otel() -> None:
    # #27: observability/otel_setup.py only wires this logger's handler/level when OTel is
    # actually enabled, and this deployment runs with it deliberately unset -- the request log
    # must still reach a real stream handler (what `railway logs` captures) on its own, not just
    # propagate to pytest's own log capture. Verified for real against a running container while
    # implementing #27 (`docker logs` showed the line only after this fix).
    from .. import main as main_module
    assert main_module.logger.getEffectiveLevel() <= logging.INFO
    assert any(isinstance(handler, logging.StreamHandler) and handler.level <= logging.INFO for handler in main_module.logger.handlers)

def test_cors_allows_configured_frontend_origin() -> None:
    response = TestClient(app, base_url="http://localhost").options("/health/live", headers={"origin": "http://localhost:3000", "access-control-request-method": "GET"})
    assert response.status_code == 200 and response.headers.get("access-control-allow-origin") == "http://localhost:3000"
    dev = TestClient(app, base_url="http://localhost").options("/health/live", headers={"origin": "http://localhost:4174", "access-control-request-method": "POST"})
    assert dev.status_code == 200 and dev.headers.get("access-control-allow-origin") == "http://localhost:4174"

def test_import_rejects_oversized_multipart_envelope() -> None:
    response = TestClient(app, base_url="http://localhost").post("/admin/imports?submission_id=large", headers={"origin": "http://localhost:3000", "content-length": str(11 * 1024 * 1024 + 1)})
    assert response.status_code == 413 and response.headers.get("x-request-id")

def test_health_endpoints_are_minimal_and_safe() -> None:
    client = TestClient(app, base_url="http://localhost")
    live = client.get("/health/live"); ready = client.get("/health/ready")
    assert live.status_code == 200 and live.json() == {"status": "ok"}
    assert ready.status_code == 200 and ready.json()["storage"] == "memory"
    from ..main import ALEMBIC_HEAD
    assert ALEMBIC_HEAD == "029_assignment_conditions"

def test_version_is_visible_without_shell_access(monkeypatch) -> None:
    # #27: the deployed commit must be identifiable from an HTTP response alone. GIT_SHA is
    # baked into the image at build time (apps/api/Dockerfile) into main.APP_VERSION; nothing
    # DB/credential-shaped rides along with it.
    from .. import main as main_module
    monkeypatch.setattr(main_module, "APP_VERSION", "abc1234")
    client = TestClient(app, base_url="http://localhost")
    live = client.get("/health/live")
    assert live.headers["x-app-version"] == "abc1234"
    ready = client.get("/health/ready")
    assert ready.json()["version"] == "abc1234"
    assert "DATABASE_URL" not in str(ready.json()) and "password" not in str(ready.json()).lower()

def test_postgres_readiness_rejects_stale_migration(monkeypatch) -> None:
    from ..main import health_ready
    database = AuthDatabase("sqlite+pysqlite:///:memory:")
    with database.engine.begin() as connection:
        connection.execute(text("CREATE TABLE alembic_version (version_num VARCHAR(32) NOT NULL)"))
        connection.execute(text("INSERT INTO alembic_version VALUES ('001_auth')"))
    monkeypatch.setitem(health_ready.__globals__, "auth_db", database)
    response = TestClient(app, base_url="http://localhost").get("/health/ready")
    assert response.status_code == 503

def test_postgres_readiness_times_out_instead_of_hanging(monkeypatch) -> None:
    # #27 failure drill: a DB outage that black-holes an already-open connection (observed
    # locally via `docker pause` on the Postgres container) must not hang this endpoint forever.
    from .. import main as main_module
    import time

    monkeypatch.setattr(main_module, "HEALTH_READY_TIMEOUT_SECONDS", 0.2)

    def hangs_forever() -> str:
        time.sleep(5)
        return main_module.ALEMBIC_HEAD

    monkeypatch.setitem(main_module.health_ready.__globals__, "auth_db", object())
    monkeypatch.setattr(main_module, "_check_postgres_ready", hangs_forever)
    started = perf_counter()
    response = TestClient(app, base_url="http://localhost").get("/health/ready")
    elapsed = perf_counter() - started
    assert response.status_code == 503
    assert elapsed < 2

def test_database_session_lookup_uses_digest_and_revocation() -> None:
    database = AuthDatabase("sqlite+pysqlite:///:memory:")
    from datetime import datetime, timedelta, timezone
    with Session(database.engine) as session:
        session.add(UserRow(id="u1", name="Admin", email="admin@example.test", role="Admin", active=True, password_hash="hash"))
        session.add(SessionRow(digest=digest("opaque-token"), user_id="u1", issued_at=datetime.now(timezone.utc), expires_at=datetime.now(timezone.utc) + timedelta(hours=1))); session.commit()
    assert database.user_for_session("opaque-token").id == "u1"
    assert database.user_for_session("opaque-token").id == "u1"
    database.revoke("opaque-token"); assert database.user_for_session("opaque-token") is None
    assert "opaque-token" not in {row.digest for row in Session(database.engine).query(SessionRow).all()}

def test_database_duplicate_identity_is_safe_conflict() -> None:
    database = AuthDatabase("sqlite+pysqlite:///:memory:")
    database.create_user(user_id="u1", name="One", email="same@example.test", role="Admin", password_hash="hash")
    try: database.create_user(user_id="u2", name="Two", email="same@example.test", role="Admin", password_hash="hash")
    except ValueError as exc: assert "already exists" in str(exc)
    else: assert False

def test_database_customer_upsert_preserves_operational_owner() -> None:
    database = CustomerDatabase("sqlite+pysqlite:///:memory:")
    first = database.upsert_source(bcn="000123", name="Original", source={"score": 1}, primary_phone="555")
    with Session(database.engine) as session:
        stored = session.get(type(first), "000123"); stored.owner_id = "sales-river"; stored.status = "Closed"; session.commit()
    second = database.upsert_source(bcn="000123", name="Imported", source={"score": 2}, primary_phone="777")
    assert second.name == "Imported" and second.owner_id == "sales-river" and second.status == "Closed"
    database.upsert_sources([{"bcn": "000123", "name": "Imported", "source": {}, "primary_phone": None}])
    assert database.phones("000123") == []

def test_database_release_open_owner_returns_released_customers() -> None:
    database = CustomerDatabase("sqlite+pysqlite:///:memory:")
    database.upsert_source(bcn="000123", name="Open", source={})
    database.save_operational(bcn="000123", owner_id="sales", status="Open", version=1)
    assert database.release_open_owner("sales") == ["000123"]
    assert database.get("000123").owner_id is None

def test_database_customer_ingest_stores_typed_source_fields() -> None:
    from datetime import date
    from decimal import Decimal
    database = CustomerDatabase("sqlite+pysqlite:///:memory:")
    database.ingest_sources([{"bcn": "000123", "name": "Typed", "source": {}, "typed": {"propensity_score": Decimal("0.875000"), "last_purchase_date": date(2025, 1, 2), "previously_contacted": True}, "primary_phone": None}], {"jobId": "job-typed", "submissionId": "typed", "filename": "x.xlsx", "processed": 1, "created": 1, "updated": 0, "errorRows": 0, "status": "Completed", "errors": [], "actorId": "admin"})
    row = database.get("000123")
    assert row.propensity_score == Decimal("0.875000") and row.last_purchase_date == date(2025, 1, 2) and row.previously_contacted is True

def test_database_customer_ingest_stores_extensible_collections() -> None:
    from decimal import Decimal
    from ..db_customers import CustomerCollectionRow
    database = CustomerDatabase("sqlite+pysqlite:///:memory:")
    result = {"jobId": "job-collections", "submissionId": "collections", "filename": "x.xlsx", "processed": 1, "created": 1, "updated": 0, "errorRows": 0, "status": "Completed", "errors": [], "actorId": "admin"}
    record = {"bcn": "000123", "name": "Collections", "source": {}, "collections": [{"kind": "vendor", "slot": 4, "name": "New vendor", "revenue": Decimal("12.50")}], "primary_phone": None}
    database.ingest_sources([record], result)
    with Session(database.engine) as session:
        row = session.query(CustomerCollectionRow).one()
        assert row.kind == "vendor" and row.slot == 4 and row.revenue == Decimal("12.500000")

def test_assignment_configuration_and_run_are_admin_only() -> None:
    repo.reset(); admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})()); repo.users["sales-river"] = {"id": "sales-river", "name": "River Sales", "email": "river@example.test", "role": "Sales", "active": True, "password": "unused"}
    client = TestClient(app, base_url="http://localhost"); client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"})
    created = client.post("/admin/assignment-rules", json={"name": "River", "conditions": [], "memberIds": ["sales-river"]}, headers={"origin": "http://localhost:3000"})
    assert created.status_code == 201
    assert client.put("/admin/assignment-fallback", json=["sales-river"], headers={"origin": "http://localhost:3000"}).json() == ["sales-river"]
    result = client.post("/admin/assignment-runs", json={"scope": "unassigned", "submissionId": "run-1"}, headers={"origin": "http://localhost:3000"})
    assert result.status_code == 200 and result.json()["assigned"] == 1
    assert client.post("/admin/assignment-runs", json={"scope": "all-open", "submissionId": "run-1"}, headers={"origin": "http://localhost:3000"}).status_code == 409
    assert client.patch("/admin/assignment-rules/rule-1", json={"version": 0}, headers={"origin": "http://localhost:3000"}).status_code == 409

def test_assignment_database_keeps_ordered_rules_and_audit() -> None:
    database = AssignmentDatabase("sqlite+pysqlite:///:memory:")
    database.create_rule(rule_id="r2", name="Second", position=2, actor_id="admin")
    database.create_rule(rule_id="r1", name="First", position=1, actor_id="admin")
    assert [row.id for row in database.ordered_rules()] == ["r1", "r2"]
    events, total = database.audit(); assert total == 2 and len(events) == 2
    filtered, filtered_total = database.audit(actor="admin", action="created", page=1, page_size=1)
    assert filtered_total == 2 and len(filtered) == 1
    result = {"submissionId": "s1", "scope": "unassigned", "assigned": 1}
    assert database.save_run(actor_id="admin", submission_id="s1", scope="unassigned", result=result) == result
    assert database.get_run("admin", "s1") == result
    database.set_setting("fallback_sales", {"ids": ["u1"]})
    assert database.get_setting("fallback_sales") == {"ids": ["u1"]}

def test_assignment_run_retry_returns_persisted_result() -> None:
    database = AssignmentDatabase("sqlite+pysqlite:///:memory:")
    assert database.save_run(actor_id="admin", submission_id="same", scope="all", result={"assigned": 1}) == {"assigned": 1}
    assert database.save_run(actor_id="admin", submission_id="same", scope="all", result={"assigned": 99}) == {"assigned": 1}
    database.save_run(actor_id="admin", submission_id="scoped", scope="all", result={"assigned": 1}, payload="all")
    try: database.get_run("admin", "scoped", "unassigned")
    except ValueError as exc: assert "already used" in str(exc)
    else: assert False

def test_assignment_run_concurrent_winner_replays_or_conflicts(tmp_path, monkeypatch) -> None:
    from hashlib import sha256
    from sqlalchemy import event
    from .. import main
    from ..db_assignment import AssignmentRunRow

    database = AssignmentDatabase(f"sqlite+pysqlite:///{tmp_path / 'assignment.db'}")
    monkeypatch.setattr(main, "assignment_db", database)
    monkeypatch.setattr(main, "customer_db", None)
    actor = main.User(id="admin", name="Admin", email="admin@example.test", role="Admin")
    monkeypatch.setitem(app.dependency_overrides, main.admin_user, lambda: actor)
    try:
        for scope, expected_status in (("unassigned", 200), ("all-open", 409)):
            repo.reset()
            winner = {"submissionId": scope, "scope": "unassigned", "assigned": 7}
            def commit_winner(session, *_):
                if session.bind is database.engine:
                    with database.engine.begin() as connection:
                        connection.execute(AssignmentRunRow.__table__.insert().values(actor_id=actor.id, submission_id=scope, scope="unassigned", fingerprint=sha256(b"unassigned").hexdigest(), result=winner, created_at=datetime.now(timezone.utc)))
            # Commit another request after save_run's read, before its INSERT.
            event.listen(Session, "before_flush", commit_winner)
            try:
                with TestClient(app, base_url="http://localhost") as client:
                    response = client.post("/admin/assignments/run", json={"scope": scope, "submissionId": scope}, headers={"origin": "http://localhost:3000"})
            finally:
                event.remove(Session, "before_flush", commit_winner)
            assert response.status_code == expected_status
            assert database.get_run(actor.id, scope) == winner
            if expected_status == 200:
                assert response.json() == winner and repo.assignment_runs[(actor.id, scope)] == winner
            else:
                assert response.json() == {"detail": "submission already used"}
                assert (actor.id, scope) not in repo.assignment_runs
    finally:
        database.engine.dispose()
        repo.reset()

def test_assignment_order_fallback_and_unchanged_owner(monkeypatch, caplog) -> None:
    """First-position-match precedence over a later also-matching rule; a matched rule with an empty
    eligible set (a member who is no longer active Sales) falls through to the next rule and logs a
    warning; a workload tie is broken by ascending user id, and the pick shifts mid-run once the first
    candidate's in-run count is bumped."""
    from .. import main
    monkeypatch.setattr(main, "assignment_db", None)
    monkeypatch.setattr(main, "customer_db", None)
    repo.reset()
    try:
        for user_id in ("first", "second", "third", "inactive"):
            repo.users[user_id] = {"id": user_id, "name": user_id, "role": "Sales", "active": user_id != "inactive"}
        repo.customers = {"000001": {"bcn": "000001", "ownerId": "first", "ownerName": "first", "status": "Open", "version": 4}, "000002": {"bcn": "000002", "ownerId": None, "status": "Open", "version": 0}}
        repo.rules = [
            {"id": "r0", "name": "Empty-eligible", "conditions": [], "memberIds": ["inactive"], "active": True, "order": 1},
            {"id": "r1", "name": "Primary", "conditions": [], "memberIds": ["first", "second"], "active": True, "order": 2},
            {"id": "r2", "name": "Never reached", "conditions": [], "memberIds": ["third"], "active": True, "order": 3},
        ]
        repo.fallback_sales = []
        actor = main.User(id="admin", name="Admin", email="admin@example.test", role="Admin")
        with caplog.at_level("WARNING", logger="call-center.assignment"):
            result = main.run_assignment(main.AssignmentRunRequest(scope="all-open", submissionId="ordered"), actor)
        assert (result["candidates"], result["assigned"], result["skipped"]) == (2, 2, 0)
        # r0 matched (empty conditions) but its only member is no longer active Sales: skipped with a warning, falling through to r1.
        assert any("Empty-eligible" in message and "no eligible active-Sales members" in message for message in caplog.messages)
        # r1 (fewer open-owned wins) reassigns 000001 to "second" (0 open < first's 1); the in-run
        # count bump then shifts 000002's pick to "first" via the ascending-id tie-break. r2's member
        # "third" is never picked -- r1 (earlier position) always had eligible members.
        assert repo.customers["000001"]["ownerId"] == "second" and repo.customers["000001"]["version"] == 5
        assert repo.customers["000002"]["ownerId"] == "first" and repo.customers["000002"]["version"] == 1
    finally:
        repo.reset()

def test_assignment_fallback_used_and_no_eligible_anywhere_leaves_unchanged(monkeypatch, caplog) -> None:
    """No rule matches (none configured): falls back to the global fallback list. An already-correct
    owner is a no-op. An empty fallback (no rule and no fallback member) leaves ownership unchanged
    and reports a 'No eligible salesperson' skip reason."""
    from .. import main
    monkeypatch.setattr(main, "assignment_db", None)
    monkeypatch.setattr(main, "customer_db", None)
    repo.reset()
    try:
        repo.users["second"] = {"id": "second", "name": "second", "role": "Sales", "active": True}
        repo.customers = {"000010": {"bcn": "000010", "ownerId": None, "status": "Open", "version": 0}}
        repo.rules = []
        repo.fallback_sales = ["second"]
        actor = main.User(id="admin", name="Admin", email="admin@example.test", role="Admin")
        result = main.run_assignment(main.AssignmentRunRequest(scope="unassigned", submissionId="fb1"), actor)
        assert (result["assigned"], result["skipped"]) == (1, 0)
        assert repo.customers["000010"]["ownerId"] == "second" and repo.customers["000010"]["version"] == 1
        result = main.run_assignment(main.AssignmentRunRequest(scope="all-open", submissionId="fb2"), actor)
        assert (result["assigned"], result["skipped"]) == (0, 1)
        assert repo.customers["000010"]["version"] == 1
        repo.fallback_sales = []
        with caplog.at_level("WARNING", logger="call-center.assignment"):
            result = main.run_assignment(main.AssignmentRunRequest(scope="all-open", submissionId="fb3"), actor)
        assert (result["assigned"], result["skipped"]) == (0, 1)
        assert repo.customers["000010"]["ownerId"] == "second" and repo.customers["000010"]["version"] == 1
        assert any("No eligible salesperson" in message for message in caplog.messages)
    finally:
        repo.reset()

def test_condition_null_semantics_reject_wrong_operators_and_unknown_fields() -> None:
    from .. import assignment_rules
    assert assignment_rules.evaluate_condition(None, {"field": "propensity_tier", "operator": "=", "value": "Gold"}) is False
    assert assignment_rules.evaluate_condition(None, {"field": "propensity_tier", "operator": "!=", "value": "Gold"}) is False  # != never matches null
    assert assignment_rules.evaluate_condition(None, {"field": "propensity_tier", "operator": "is-null", "value": None}) is True
    assert assignment_rules.evaluate_condition("Gold", {"field": "propensity_tier", "operator": "is-not-null", "value": None}) is True
    assert assignment_rules.evaluate_condition(None, {"field": "propensity_tier", "operator": "is-not-null", "value": None}) is False
    try:
        assignment_rules.validate_condition("propensity_score", "contains", "5")
        assert False
    except assignment_rules.ConditionError:
        pass
    for forbidden in ("status", "owner_id", "bogus_field"):
        try:
            assignment_rules.validate_condition(forbidden, "=", "x")
            assert False
        except assignment_rules.ConditionError:
            pass

def test_condition_between_is_inclusive_and_rejects_invalid_range_at_save_time() -> None:
    from .. import assignment_rules
    condition = assignment_rules.validate_condition("propensity_score", "between", [10, 20])
    assert condition == {"field": "propensity_score", "operator": "between", "value": ["10", "20"]}
    assert assignment_rules.evaluate_condition(10, condition) is True
    assert assignment_rules.evaluate_condition(20, condition) is True
    assert assignment_rules.evaluate_condition(9, condition) is False
    assert assignment_rules.evaluate_condition(21, condition) is False
    for bad_range in ([20, 10], [None, 20], [10, None]):
        try:
            assignment_rules.validate_condition("propensity_score", "between", bad_range)
            assert False
        except assignment_rules.ConditionError:
            pass

def test_pick_candidate_tie_break_ascending_id_and_mid_run_shift() -> None:
    from .. import assignment_rules
    counts = {"a": 0, "b": 0}
    assert assignment_rules.pick_candidate(["b", "a"], counts) == "a"
    assignment_rules.record_pick(counts, "a")
    assert counts == {"a": 1, "b": 0}
    assert assignment_rules.pick_candidate(["a", "b"], counts) == "b"

def test_assignment_fallback_rejects_invalid_members_without_partial_write() -> None:
    repo.reset(); admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    repo.users["sales-river"] = {"id": "sales-river", "name": "River Sales", "email": "river@example.test", "role": "Sales", "active": True, "password": "unused"}
    client = TestClient(app, base_url="http://localhost"); client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"})
    assert client.put("/admin/assignment-fallback", json=["sales-river"], headers={"origin": "http://localhost:3000"}).json() == ["sales-river"]
    rejected = client.put("/admin/assignment-fallback", json=["sales-river", "ghost-user"], headers={"origin": "http://localhost:3000"})
    assert rejected.status_code == 422
    assert client.get("/admin/assignment-fallback", headers={"origin": "http://localhost:3000"}).json() == ["sales-river"]

def test_audit_endpoint_returns_only_whitelisted_detail_keys(tmp_path, monkeypatch) -> None:
    from .. import main
    database = AssignmentDatabase(f"sqlite+pysqlite:///{tmp_path / 'audit.db'}")
    monkeypatch.setattr(main, "assignment_db", database)
    try:
        database.append_audit(actor_id="admin", action="Customer assigned", target="000123", details={"oldOwner": None, "newOwner": "sales", "secret": "should-not-appear", "password": "nope"})
        result = main.admin_audit(page=1, page_size=25)
        item = result["items"][0]
        assert set(item) == {"id", "actor", "actorId", "action", "target", "timestamp", "details"}
        assert item["details"] == {"oldOwner": None, "newOwner": "sales"}
    finally:
        database.engine.dispose()

def test_activity_idempotency_replays_and_rejects_payload_reuse() -> None:
    database = ActivityDatabase("sqlite+pysqlite:///:memory:")
    assert database.save_idempotent(actor_id="u1", operation="note", submission_id="s1", payload="hello", result={"id": "n1"}) == {"id": "n1"}
    assert database.save_idempotent(actor_id="u1", operation="note", submission_id="s1", payload="hello", result={"id": "ignored"}) == {"id": "n1"}
    try: database.save_idempotent(actor_id="u1", operation="note", submission_id="s1", payload="different", result={})
    except ValueError as exc: assert "already used" in str(exc)
    else: assert False
    database.save_activity(record_id="a1", bcn="000123", actor_id="u1", kind="Interaction", outcome="Attempt", text="x")
    assert database.soft_delete("a1", "u1") is True
    assert database.history("000123")[0][0].deleted_by == "u1"
    assert database.history_map(["000123"])["000123"][0]["deletedBy"] == "u1"

def test_followup_completion_rejects_cross_customer_and_stale_links() -> None:
    database = ActivityDatabase("sqlite+pysqlite:///:memory:")
    followup = {"id": "f1", "bcn": "000123", "actorId": "u1", "type": "Reminder", "due": None, "status": "Open", "note": None}
    database.save_followup(followup)
    try: database.complete_followup(followup, {"id": "a1", "bcn": "000124", "actorId": "u1"})
    except ValueError as exc: assert "another customer" in str(exc)
    else: assert False
    database.complete_followup(followup, {"id": "a1", "bcn": "000123", "actorId": "u1"})
    try: database.complete_followup(followup, {"id": "a2", "bcn": "000123", "actorId": "u1"})
    except ValueError as exc: assert "no longer open" in str(exc)
    else: assert False

def test_close_lifecycle_commits_cancellations_and_retry_record() -> None:
    from datetime import datetime, timezone
    database = ActivityDatabase("sqlite+pysqlite:///:memory:")
    database.save_followup({"id": "f-close", "bcn": "000123", "actorId": "u1", "type": "Reminder", "due": None, "status": "Open", "note": "next"})
    event = {"id": "closure-000123-1", "reason": "Completed", "timestamp": datetime.now(timezone.utc).isoformat()}
    database.close_lifecycle("000123", event, actor_id="u1", submission_id="close-1", payload="000123|close|r1", result={"status": "Closed"})
    assert database.followups("000123")[0].status == "Cancelled"
    assert database.get_idempotent(actor_id="u1", operation="close", submission_id="close-1", payload="000123|close|r1") == {"status": "Closed"}

def test_close_lifecycle_rolls_back_on_activity_conflict() -> None:
    from datetime import datetime, timezone
    from sqlalchemy.exc import IntegrityError
    from ..db_activity import ActivityRow
    database = ActivityDatabase("sqlite+pysqlite:///:memory:")
    database.save_followup({"id": "f-rollback", "bcn": "000123", "actorId": "u1", "type": "Reminder", "due": None, "status": "Open", "note": "next"})
    with database.engine.begin() as connection:
        connection.execute(ActivityRow.__table__.insert().values(id="activity-f-rollback-cancel", bcn="000123", actor_id="u1", kind="Existing", created_at=datetime.now(timezone.utc)))
    try: database.close_lifecycle("000123", {"id": "closure-rollback", "reason": "Done", "timestamp": datetime.now(timezone.utc).isoformat()}, actor_id="u1", submission_id="close-rollback", payload="p", result={})
    except IntegrityError: pass
    else: assert False
    assert database.followups("000123")[0].status == "Open" and database.get_idempotent(actor_id="u1", operation="close", submission_id="close-rollback", payload="p") is None

def test_reopen_customer_commits_owner_status_and_idempotency_atomically(tmp_path) -> None:
    """Regression for the #25 QA bug: reopen used to be three independently committed writes
    (activity event, customer operational row, idempotency record). reopen_customer must commit
    all three together, and an injected failure must leave none of them applied. Also covers #30:
    when ownership is released on reopen, the assignment_history/audit_events rows written in the
    same call must commit -- and roll back -- together with everything else too."""
    from datetime import datetime, timezone
    from sqlalchemy import func, select
    from sqlalchemy.exc import IntegrityError
    from ..db_activity import ActivityRow
    from ..db_assignment import AssignmentDatabase, AssignmentHistoryRow, AuditRow
    from ..db_customers import CustomerDatabase

    url = f"sqlite+pysqlite:///{tmp_path / 'reopen.db'}"
    customers = CustomerDatabase(url)
    activities = ActivityDatabase(url)
    assignments = AssignmentDatabase(url)
    try:
        customers.upsert_source(bcn="000123", name="Synthetic", source={})
        customers.save_operational(bcn="000123", owner_id="sales", status="Closed", version=2)

        activities.reopen_customer("000123", {"id": "reopen-000123-3", "timestamp": datetime.now(timezone.utc).isoformat()}, owner_id="sales", status="Open", version=3, actor_id="sales", submission_id="reopen-1", payload="000123|reopen", result={"status": "Open"})
        assert (customers.get("000123").status, customers.get("000123").version, customers.get("000123").owner_id) == ("Open", 3, "sales")
        assert activities.get_idempotent(actor_id="sales", operation="reopen", submission_id="reopen-1", payload="000123|reopen") == {"status": "Open"}
        with Session(activities.engine) as session:
            assert session.get(ActivityRow, "reopen-000123-3").kind == "Reopen"

        # A conflicting activity id (simulating an injected failure partway through) rolls back the
        # whole write: owner/status/version and the idempotency record stay unchanged together.
        with activities.engine.begin() as connection:
            connection.execute(ActivityRow.__table__.insert().values(id="reopen-000123-4", bcn="000123", actor_id="sales", kind="Existing", created_at=datetime.now(timezone.utc)))
        try:
            activities.reopen_customer("000123", {"id": "reopen-000123-4", "timestamp": datetime.now(timezone.utc).isoformat()}, owner_id=None, status="Open", version=4, actor_id="sales", submission_id="reopen-2", payload="p2", result={})
        except IntegrityError: pass
        else: assert False
        assert (customers.get("000123").status, customers.get("000123").version, customers.get("000123").owner_id) == ("Open", 3, "sales")
        assert activities.get_idempotent(actor_id="sales", operation="reopen", submission_id="reopen-2", payload="p2") is None

        # #30 release path: a successful release writes exactly one assignment_history row and one
        # audit_events row alongside owner/status/version and the idempotency record.
        customers.save_operational(bcn="000123", owner_id="sales", status="Closed", version=4)
        activities.reopen_customer("000123", {"id": "reopen-000123-5", "timestamp": datetime.now(timezone.utc).isoformat()}, owner_id=None, status="Open", version=5, actor_id="sales", submission_id="reopen-3", payload="000123|reopen-release", result={"status": "Open", "ownerId": None}, released_owner_id="sales")
        assert (customers.get("000123").status, customers.get("000123").version, customers.get("000123").owner_id) == ("Open", 5, None)
        with Session(assignments.engine) as session:
            history_row = session.execute(select(AssignmentHistoryRow).where(AssignmentHistoryRow.bcn == "000123")).scalar_one()
            assert history_row.old_owner_id == "sales" and history_row.new_owner_id is None and history_row.reason == "Owner no longer active Sales"
            audit_row = session.execute(select(AuditRow).where(AuditRow.target == "000123", AuditRow.action == "Customer ownership released")).scalar_one()
            assert audit_row.actor_id == "sales" and audit_row.details == {"oldOwner": "sales", "newOwner": None, "reason": "Owner no longer active Sales"}

        # An injected failure on a release write rolls back the assignment_history/audit_events rows
        # together with owner/status/version and the idempotency record -- no partial release survives.
        with activities.engine.begin() as connection:
            connection.execute(ActivityRow.__table__.insert().values(id="reopen-000123-6", bcn="000123", actor_id="sales", kind="Existing", created_at=datetime.now(timezone.utc)))
        try:
            activities.reopen_customer("000123", {"id": "reopen-000123-6", "timestamp": datetime.now(timezone.utc).isoformat()}, owner_id=None, status="Open", version=6, actor_id="sales", submission_id="reopen-4", payload="p4", result={}, released_owner_id="sales")
        except IntegrityError: pass
        else: assert False
        assert (customers.get("000123").status, customers.get("000123").version, customers.get("000123").owner_id) == ("Open", 5, None)
        assert activities.get_idempotent(actor_id="sales", operation="reopen", submission_id="reopen-4", payload="p4") is None
        with Session(assignments.engine) as session:
            assert session.scalar(select(func.count()).select_from(AssignmentHistoryRow).where(AssignmentHistoryRow.bcn == "000123")) == 1
            assert session.scalar(select(func.count()).select_from(AuditRow).where(AuditRow.target == "000123", AuditRow.action == "Customer ownership released")) == 1
    finally:
        customers.engine.dispose(); activities.engine.dispose(); assignments.engine.dispose()

def test_reopen_releases_ownership_when_owner_not_active_sales_memory_mode() -> None:
    """#30, memory mode: reopen must apply the same active-Sales rule as PostgreSQL mode through the
    real HTTP endpoint -- covering the prior owner still being active Sales (retained, unchanged),
    role changed away from Sales while still active, an owner id no longer present in the user store
    at all, and a customer that was already unassigned at close time (no spurious history entry)."""
    repo.reset()
    admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    repo.users["sales-active"] = {"id": "sales-active", "name": "Active Sales", "email": "active@example.test", "role": "Sales", "active": True, "password": ""}
    repo.users["sales-demoted"] = {"id": "sales-demoted", "name": "Demoted", "email": "demoted@example.test", "role": "Admin", "active": True, "password": ""}
    repo.customers["900001"] = {"bcn": "900001", "name": "Retained Co", "ownerId": "sales-active", "ownerName": "Active Sales", "status": "Closed", "phones": [], "source": {}, "version": 0, "histories": []}
    repo.customers["900002"] = {"bcn": "900002", "name": "Demoted Co", "ownerId": "sales-demoted", "ownerName": "Demoted", "status": "Closed", "phones": [], "source": {}, "version": 0, "histories": []}
    repo.customers["900003"] = {"bcn": "900003", "name": "Ghost Owner Co", "ownerId": "sales-ghost", "ownerName": "Ghost", "status": "Closed", "phones": [], "source": {}, "version": 0, "histories": []}
    repo.customers["900004"] = {"bcn": "900004", "name": "Already Unassigned Co", "ownerId": None, "ownerName": None, "status": "Closed", "phones": [], "source": {}, "version": 0, "histories": []}
    client = TestClient(app, base_url="http://localhost"); origin = {"origin": "http://localhost:3000"}
    assert client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"}).status_code == 200

    retained = client.post("/customers/900001/reopen", json={"submissionId": "retained"}, headers=origin)
    assert retained.status_code == 200
    retained_body = retained.json()
    assert retained_body["ownerId"] == "sales-active" and retained_body["ownerName"] == "Active Sales"
    assert not any(item["kind"] == "Assignment" for item in retained_body["histories"])

    demoted = client.post("/customers/900002/reopen", json={"submissionId": "demoted"}, headers=origin)
    assert demoted.status_code == 200
    demoted_body = demoted.json()
    assert demoted_body["ownerId"] is None and demoted_body["ownerName"] is None and demoted_body["status"] == "Open"
    demoted_entries = [item for item in demoted_body["histories"] if item["kind"] == "Assignment"]
    assert len(demoted_entries) == 1 and demoted_entries[0]["oldOwner"] == "sales-demoted" and demoted_entries[0]["newOwner"] is None and demoted_entries[0]["reason"] == "Owner no longer active Sales"
    replay = client.post("/customers/900002/reopen", json={"submissionId": "demoted"}, headers=origin)
    assert replay.status_code == 200 and replay.json() == demoted_body

    ghost = client.post("/customers/900003/reopen", json={"submissionId": "ghost"}, headers=origin)
    assert ghost.status_code == 200
    ghost_body = ghost.json()
    assert ghost_body["ownerId"] is None and ghost_body["ownerName"] is None
    ghost_entries = [item for item in ghost_body["histories"] if item["kind"] == "Assignment"]
    assert len(ghost_entries) == 1 and ghost_entries[0]["oldOwner"] == "sales-ghost" and ghost_entries[0]["newOwner"] is None

    unassigned = client.post("/customers/900004/reopen", json={"submissionId": "unassigned"}, headers=origin)
    assert unassigned.status_code == 200
    unassigned_body = unassigned.json()
    assert unassigned_body["ownerId"] is None and unassigned_body["status"] == "Open"
    assert not any(item["kind"] == "Assignment" for item in unassigned_body["histories"])

def test_real_http_admin_sales_journey() -> None:
    repo.reset()
    admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    sales = provision_user(type("P", (), {"name": "River", "email": "river@example.test", "role": "Sales", "password": "correct horse battery staple"})())
    repo.customers["000123"]["ownerId"] = sales.id; repo.customers["000123"]["ownerName"] = sales.name
    client = TestClient(app, base_url="http://localhost"); origin = {"origin": "http://localhost:3000"}
    assert client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"}).status_code == 200
    workbook = Workbook(); workbook.active.append(["bcn", "customer_name"]); workbook.active.append(["009990", "Journey Co"]); payload = BytesIO(); workbook.save(payload)
    imported = client.post("/admin/imports?submission_id=journey-import", files={"file": ("journey.xlsx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=origin)
    assert imported.status_code == 201 and imported.json()["created"] == 1
    assigned = client.post("/admin/assignments/manual/009990", json={"ownerId": sales.id, "submissionId": "journey-assign"}, headers=origin)
    assert assigned.status_code == 200
    client.post("/session/logout", headers=origin)
    assert client.post("/session/login", json={"email": sales.email, "password": "correct horse battery staple"}).status_code == 200
    assert client.get("/customers", params={"mine": "true"}).json()["total"] == 2
    interaction = client.post("/customers/009990/interactions", json={"outcome": "Contact", "note": "Journey", "submissionId": "journey-contact"}, headers=origin)
    assert interaction.status_code == 200, interaction.text
    followup = client.post("/customers/009990/follow-ups", json={"type": "Reminder", "due": "2026-09-20", "note": "Next", "submissionId": "journey-followup"}, headers=origin)
    assert followup.status_code == 200
    complete_url = f"/customers/009990/follow-ups/{followup.json()['id']}/complete"
    assert client.post(complete_url, json={"outcome": "Attempt", "submissionId": "journey-complete"}, headers=origin).status_code == 200
    assert client.post(complete_url, json={"outcome": "Contact", "submissionId": "journey-complete"}, headers=origin).status_code == 409
    assert client.post("/customers/000125/interactions", json={"outcome": "Attempt", "submissionId": "foreign"}, headers=origin).status_code == 403

def test_import_preserves_source_columns_and_operational_phone_history() -> None:
    repo.reset()
    admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    repo.customers["000123"]["ownerId"] = "sales-owner"
    repo.customers["000123"]["phones"] = ["old-primary", "independent"]
    client = TestClient(app, base_url="http://localhost")
    assert client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"}).status_code == 200
    workbook = Workbook(); workbook.active.append(["bcn", "customer_name", "propensity_score", "LAST_PURCHASE_DATE"]); workbook.active.append(["000123", "Updated", 0.875, datetime(2025, 1, 2)])
    payload = BytesIO(); workbook.save(payload)
    response = client.post("/admin/imports?submission_id=source-columns", files={"file": ("source.xlsx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers={"origin": "http://localhost:3000"})
    assert response.status_code == 201
    row = repo.customers["000123"]
    assert row["ownerId"] == "sales-owner" and row["phones"] == ["independent"]
    assert row["source"]["propensity_score"] == 0.875 and row["source"]["LAST_PURCHASE_DATE"].startswith("2025-01-02")

def test_import_retry_key_is_scoped_to_actor() -> None:
    database = CustomerDatabase("sqlite+pysqlite:///:memory:")
    result = {"jobId": "job-1", "submissionId": "same", "filename": "a.xlsx", "processed": 0, "created": 0, "updated": 0, "errorRows": 0, "status": "Completed", "errors": [], "actorId": "u1"}
    database.save_import_job(result)
    other = {**result, "jobId": "job-2", "actorId": "u2"}
    database.save_import_job(other)
    assert database.import_job("u1", "same")["jobId"] == "job-1"
    assert database.import_job("u2", "same")["jobId"] == "job-2"

def test_import_fingerprint_round_trips_with_durable_job() -> None:
    database = CustomerDatabase("sqlite+pysqlite:///:memory:")
    database.save_import_job({"jobId": "job-fingerprint", "submissionId": "fingerprint", "filename": "x.xlsx", "processed": 0, "created": 0, "updated": 0, "errorRows": 0, "status": "Completed", "errors": [], "actorId": "u1", "fingerprint": "a" * 64})
    assert database.import_job("u1", "fingerprint")["fingerprint"] == "a" * 64

def test_import_rejects_missing_required_customer_name() -> None:
    repo.reset()
    admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    client = TestClient(app, base_url="http://localhost")
    client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"})
    workbook = Workbook(); workbook.active.append(["bcn", "customer_name"]); workbook.active.append(["123456", ""])
    payload = BytesIO(); workbook.save(payload)
    response = client.post("/admin/imports?submission_id=missing-name", files={"file": ("source.xlsx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers={"origin": "http://localhost:3000"})
    assert response.status_code == 201 and response.json()["created"] == 0 and response.json()["errorRows"] == 1

def test_import_errors_are_paginated() -> None:
    repo.reset(); admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    repo.imports[(admin.id, "errors")] = {"errors": [{"row": n, "field": "bcn", "reason": "Invalid bcn"} for n in range(2, 5)]}
    client = TestClient(app, base_url="http://localhost"); client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"})
    response = client.get("/admin/imports/errors/errors?page=2&page_size=2")
    assert response.status_code == 200 and response.json()["total"] == 3 and len(response.json()["items"]) == 1

def test_database_import_errors_are_paginated_in_query() -> None:
    database = CustomerDatabase("sqlite+pysqlite:///:memory:")
    database.save_import_job({"jobId": "job-errors", "submissionId": "errors", "filename": "x.xlsx", "processed": 2, "created": 0, "updated": 0, "errorRows": 2, "status": "Partial", "errors": [{"row": 2, "field": "bcn", "reason": "bad"}, {"row": 3, "field": "bcn", "reason": "bad"}], "actorId": "admin"})
    items, total = database.import_errors("admin", "errors", 2, 1)
    assert total == 2 and items == [{"row": 3, "field": "bcn", "reason": "bad"}]

def test_import_storage_failure_restores_in_memory_staging(monkeypatch) -> None:
    from .. import main
    repo.reset(); admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    client = TestClient(app, base_url="http://localhost"); client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"})
    workbook = Workbook(); workbook.active.append(["bcn", "customer_name"]); workbook.active.append(["999999", "Transient"])
    payload = BytesIO(); workbook.save(payload)
    class FailingCustomerDB:
        failed = None
        def ingest_sources(self, *_): raise RuntimeError("database unavailable")
        def import_job(self, *_): return None
        def save_import_job(self, result): self.failed = result
    failing = FailingCustomerDB(); monkeypatch.setattr(main, "customer_db", failing)
    response = client.post("/admin/imports?submission_id=rollback", files={"file": ("source.xlsx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers={"origin": "http://localhost:3000"})
    assert response.status_code == 422 and "999999" not in repo.customers and failing.failed["status"] == "Failed"

def test_import_accepts_header_only_workbook_with_zero_totals() -> None:
    repo.reset()
    admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    client = TestClient(app, base_url="http://localhost"); client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"})
    workbook = Workbook(); workbook.active.append(["bcn", "customer_name"])
    payload = BytesIO(); workbook.save(payload)
    response = client.post("/admin/imports?submission_id=empty-book", files={"file": ("empty.xlsx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers={"origin": "http://localhost:3000"})
    body = response.json()
    assert response.status_code == 201
    assert (body["processed"], body["created"], body["updated"], body["errorRows"], body["status"]) == (0, 0, 0, 0, "Completed")

def test_import_rejects_wrong_extension_without_persisting_a_job() -> None:
    repo.reset()
    admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    client = TestClient(app, base_url="http://localhost"); client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"})
    response = client.post("/admin/imports?submission_id=wrong-ext", files={"file": ("customers.csv", b"bcn,customer_name\n000123,Renamed\n", "text/csv")}, headers={"origin": "http://localhost:3000"})
    assert response.status_code == 422
    assert client.get("/admin/imports/wrong-ext").status_code == 404

def test_import_rejects_corrupt_workbook_archive_without_persisting_a_job() -> None:
    repo.reset()
    admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    client = TestClient(app, base_url="http://localhost"); client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"})
    response = client.post("/admin/imports?submission_id=corrupt-zip", files={"file": ("corrupt.xlsx", b"this is not a real zip archive", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers={"origin": "http://localhost:3000"})
    assert response.status_code == 422
    assert client.get("/admin/imports/corrupt-zip").status_code == 404

def test_memory_assignment_retry_is_scoped_to_actor() -> None:
    repo.reset()
    admin_one = provision_user(type("P", (), {"name": "One", "email": "one@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    admin_two = provision_user(type("P", (), {"name": "Two", "email": "two@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    repo.assignment_runs[(admin_one.id, "same-run")] = {"actorId": admin_one.id}
    repo.assignment_runs[(admin_two.id, "same-run")] = {"actorId": admin_two.id}
    assert repo.assignment_runs[(admin_one.id, "same-run")] != repo.assignment_runs[(admin_two.id, "same-run")]

def test_memory_followup_retry_rejects_payload_reuse() -> None:
    repo.reset()
    sales = provision_user(type("P", (), {"name": "Sales", "email": "sales@example.test", "role": "Sales", "password": "correct horse battery staple"})())
    repo.customers["000123"]["ownerId"] = sales.id
    client = TestClient(app, base_url="http://localhost")
    client.post("/session/login", json={"email": sales.email, "password": "correct horse battery staple"})
    origin = {"origin": "http://localhost:3000"}
    body = {"type": "Reminder", "due": "2026-09-20", "note": "Call", "submissionId": "same-followup"}
    assert client.post("/customers/000123/follow-ups", json=body, headers=origin).status_code == 200
    changed = {**body, "note": "Different"}
    assert client.post("/customers/000123/follow-ups", json=changed, headers=origin).status_code == 409

def test_memory_lifecycle_retries_are_idempotent() -> None:
    repo.reset()
    admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    client = TestClient(app, base_url="http://localhost"); client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"})
    origin = {"origin": "http://localhost:3000"}
    body = {"reasonId": "closure-1", "submissionId": "close-once"}
    assert client.post("/customers/000123/close", json=body, headers=origin).status_code == 200
    replay = client.post("/customers/000123/close", json=body, headers=origin)
    assert replay.status_code == 200 and repo.customers["000123"]["version"] == 1

def test_memory_import_retry_rejects_changed_workbook() -> None:
    repo.reset()
    admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    client = TestClient(app, base_url="http://localhost"); client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"})
    origin = {"origin": "http://localhost:3000"}
    def workbook(name: str) -> bytes:
        book = Workbook(); book.active.append(["bcn", "customer_name"]); book.active.append(["991001", name]); output = BytesIO(); book.save(output); return output.getvalue()
    first = client.post("/admin/imports?submission_id=same-import", files={"file": ("source.xlsx", workbook("One"), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=origin)
    second = client.post("/admin/imports?submission_id=same-import", files={"file": ("source.xlsx", workbook("Two"), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=origin)
    assert first.status_code == 201 and second.status_code == 409

def test_import_result_can_be_read_by_owner() -> None:
    repo.reset()
    admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    client = TestClient(app, base_url="http://localhost"); client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"})
    book = Workbook(); book.active.append(["bcn", "customer_name"]); book.active.append(["991002", "Readback"]); payload = BytesIO(); book.save(payload)
    origin = {"origin": "http://localhost:3000"}
    assert client.post("/admin/imports?submission_id=readback", files={"file": ("source.xlsx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=origin).status_code == 201
    result = client.get("/admin/imports/readback")
    assert result.status_code == 200 and result.json()["submissionId"] == "readback"
    page = client.get("/admin/imports?page=1&page_size=1")
    assert page.status_code == 200 and page.json()["total"] == 1 and len(page.json()["items"]) == 1

def test_memory_import_submission_is_scoped_to_actor() -> None:
    repo.reset()
    one = provision_user(type("P", (), {"name": "One", "email": "one@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    two = provision_user(type("P", (), {"name": "Two", "email": "two@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    repo.imports[(one.id, "shared")] = {"actorId": one.id, "submissionId": "shared"}
    repo.imports[(two.id, "shared")] = {"actorId": two.id, "submissionId": "shared"}
    assert repo.imports[(one.id, "shared")]["actorId"] != repo.imports[(two.id, "shared")]["actorId"]

def test_mutation_routes_bind_json_bodies() -> None:
    paths = {route.path: {field.name for field in route.dependant.body_params} for route in app.routes if getattr(route, "dependant", None)}
    assert paths["/customers/{bcn}/interactions"] == {"body"}
    assert paths["/admin/users"] == {"data"}
    assert paths["/admin/assignment-fallback"] == {"ids"}


def test_expired_session_is_rejected() -> None:
    repo.reset()
    user = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    repo.sessions["expired"] = (user.id, datetime.now(timezone.utc) - timedelta(seconds=1))
    client = TestClient(app); client.cookies.set("call_center_session", "expired")
    assert client.get("/session/me").status_code == 401


def test_foreign_origin_is_rejected_before_mutation() -> None:
    repo.reset(); client = TestClient(app)
    assert client.post("/session/logout", headers={"origin": "https://foreign.example"}).status_code == 403


def test_memory_unit_of_work_rolls_back_auth_state() -> None:
    repo.reset()
    try:
        with repo.transaction():
            provision_user(type("P", (), {"name": "Transient", "email": "transient@example.test", "role": "Admin", "password": "correct horse battery staple"})())
            raise RuntimeError("rollback")
    except RuntimeError:
        pass
    assert repo.users == {}


def test_operator_provision_and_recovery_revoke_session() -> None:
    repo.reset(); client = TestClient(app, base_url="http://localhost"); headers = {"origin": "http://localhost:3000"}
    provision = client.post("/operator/provision", json={"name": "Initial Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"}, headers=headers)
    assert provision.status_code == 200 and "password" not in provision.json()
    login = client.post("/session/login", json={"email": "admin@example.test", "password": "correct horse battery staple"}); assert login.status_code == 200
    recovered = client.post(f"/operator/reset-password/{provision.json()['id']}", json={"password": "new correct horse battery staple"}, headers=headers)
    assert recovered.status_code == 200 and "password" not in recovered.json()
    assert client.get("/session/me").status_code == 401
    assert client.post("/operator/provision", json={"name": "Second", "email": "second@example.test", "role": "Admin", "password": "correct horse battery staple"}, headers=headers).status_code == 409


def test_password_bounds_and_inactive_users_have_safe_failures() -> None:
    repo.reset(); provision_user(type("P", (), {"name": "Inactive", "email": "inactive@example.test", "role": "Sales", "password": "correct horse battery staple"})()); repo.users["user-1"]["active"] = False
    client = TestClient(app, base_url="http://localhost")
    assert client.post("/session/login", json={"email": "inactive@example.test", "password": "correct horse battery staple"}).status_code == 401
    short = client.post("/session/login", json={"email": "inactive@example.test", "password": "short"})
    assert short.status_code == 422 and "short" not in short.text


def test_customer_reads_and_admin_assignment() -> None:
    repo.reset(); admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})()); repo.users["sales-river"] = {"id": "sales-river", "name": "River Sales", "email": "river@example.test", "role": "Sales", "active": True, "password": "unused"}
    client = TestClient(app, base_url="http://localhost"); client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"})
    assert client.get("/customers/000123").status_code == 200
    changed = client.post("/admin/assignments/manual/000125", json={"ownerId": "sales-river", "submissionId": "assign-1"}, headers={"origin": "http://localhost:3000"})
    assert changed.status_code == 200 and changed.json()["ownerId"] == "sales-river"
    assert client.post("/admin/assignments/manual/000125", json={"ownerId": "sales-river", "submissionId": "assign-1"}, headers={"origin": "http://localhost:3000"}).status_code == 200
    assert client.post("/admin/assignments/manual/000125", json={"ownerId": None, "submissionId": "assign-1"}, headers={"origin": "http://localhost:3000"}).status_code == 409
    assert client.post("/admin/assignments/manual/000125", json={"ownerId": None, "submissionId": "assign-2", "expectedVersion": 0}, headers={"origin": "http://localhost:3000"}).status_code == 409
    assert client.post("/admin/assignments/manual/000125", json={"ownerId": "missing", "submissionId": "assign-3"}, headers={"origin": "http://localhost:3000"}).status_code == 422
    assert client.get("/customers/missing").status_code == 404


def test_sales_my_scope_cannot_be_widened_and_reads_are_paginated() -> None:
    repo.reset(); provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})()); repo.users["sales-river"] = {"id": "sales-river", "name": "River Sales", "email": "river@example.test", "role": "Sales", "active": True, "password": password_hash.hash("correct horse battery staple")}
    client = TestClient(app, base_url="http://localhost"); client.post("/session/login", json={"email": "river@example.test", "password": "correct horse battery staple"})
    scoped = client.get("/customers", params={"mine": "true", "page_size": 1}); assert scoped.status_code == 200 and scoped.json()["total"] == 1 and scoped.json()["items"][0]["bcn"] == "000123"
    assert client.get("/customers", params={"mine": "true", "owner": "sales-sky"}).json()["total"] == 0
    assert client.get("/customers", params={"page": 0}).status_code == 422

def test_admin_xlsx_import_preserves_assignment_and_is_idempotent() -> None:
    repo.reset(); admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    client = TestClient(app, base_url="http://localhost"); client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"})
    workbook = Workbook(); sheet = workbook.active; sheet.append(["bcn", "customer_name", "phone"]); sheet.append(["000123", "Renamed", "555-0001"]); sheet.append(["009999", "New Co", None]); sheet.append(["009999", "Duplicate", None]); payload = BytesIO(); workbook.save(payload); payload.seek(0)
    response = client.post("/admin/imports?submission_id=job-1", files={"file": ("customers.xlsx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers={"origin": "http://localhost:3000"})
    assert response.status_code == 201 and response.json()["processed"] == 3 and response.json()["created"] == 1 and len(response.json()["errors"]) == 1
    assert client.get("/customers/000123").json()["ownerId"] == "sales-river"
    assert client.post("/admin/imports?submission_id=job-1", files={"file": ("customers.xlsx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers={"origin": "http://localhost:3000"}).json()["jobId"] == response.json()["jobId"]
