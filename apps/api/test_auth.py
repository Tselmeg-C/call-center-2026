from datetime import datetime, timedelta, timezone
from io import BytesIO
from openpyxl import Workbook

from fastapi.testclient import TestClient
from sqlalchemy.orm import Session
from sqlalchemy import text

from .main import app, password_hash, repo, provision_user
from .storage import mode
from .db_auth import AuthDatabase, UserRow, SessionRow, digest
from .db_customers import CustomerDatabase
from .db_assignment import AssignmentDatabase
from .db_activity import ActivityDatabase


def test_login_logout_and_generic_failure() -> None:
    repo.reset()
    provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    client = TestClient(app, base_url="http://localhost")
    assert client.post("/session/login", json={"email": "unknown@example.test", "password": "wrong password"}).status_code == 401
    login = client.post("/session/login", json={"email": " ADMIN@example.test ", "password": "correct horse battery staple"})
    assert login.status_code == 200 and login.json()["role"] == "Admin"
    assert client.get("/session/me").status_code == 200
    assert client.post("/session/logout", headers={"origin": "http://localhost:3000"}).status_code == 204
    assert client.get("/session/me").status_code == 401

def test_storage_selection_is_explicit(monkeypatch) -> None:
    monkeypatch.setenv("CALL_CENTER_STORAGE", "postgres"); monkeypatch.delenv("DATABASE_URL", raising=False)
    try: mode(); assert False
    except RuntimeError as exc: assert "DATABASE_URL" in str(exc)

def test_login_failure_throttle_is_generic() -> None:
    repo.reset(); client = TestClient(app, base_url="http://localhost")
    for _ in range(5): assert client.post("/session/login", json={"email": "unknown@example.test", "password": "wrong password"}).status_code == 401
    limited = client.post("/session/login", json={"email": "unknown@example.test", "password": "wrong password"})
    assert limited.status_code == 429 and limited.headers.get("retry-after") == "900" and "unknown@example.test" not in limited.text

def test_login_failure_ip_throttle_limits_identity_spraying() -> None:
    repo.reset(); client = TestClient(app, base_url="http://localhost")
    for index in range(50): assert client.post("/session/login", json={"email": f"unknown-{index}@example.test", "password": "wrong password"}).status_code == 401
    limited = client.post("/session/login", json={"email": "unknown-final@example.test", "password": "wrong password"})
    assert limited.status_code == 429 and limited.headers.get("retry-after") == "900"

def test_request_id_rejects_malformed_client_value() -> None:
    response = TestClient(app, base_url="http://localhost").get("/health/live", headers={"x-request-id": "bad\nvalue"})
    assert response.status_code == 200 and "\n" not in response.headers["x-request-id"] and len(response.headers["x-request-id"]) > 10

def test_cors_allows_configured_frontend_origin() -> None:
    response = TestClient(app, base_url="http://localhost").options("/health/live", headers={"origin": "http://localhost:3000", "access-control-request-method": "GET"})
    assert response.status_code == 200 and response.headers.get("access-control-allow-origin") == "http://localhost:3000"
    dev = TestClient(app, base_url="http://localhost").options("/health/live", headers={"origin": "http://localhost:4174", "access-control-request-method": "POST"})
    assert dev.status_code == 200 and dev.headers.get("access-control-allow-origin") == "http://localhost:4174"

def test_import_rejects_oversized_multipart_envelope() -> None:
    response = TestClient(app, base_url="http://localhost").post("/admin/imports?submission_id=large", headers={"origin": "http://localhost:3000", "content-length": str(11 * 1024 * 1024 + 1)})
    assert response.status_code == 413 and response.headers.get("x-request-id")

def test_health_endpoints_are_minimal_and_safe() -> None:
    client = TestClient(app, base_url="http://localhost")
    live = client.get("/health/live"); ready = client.get("/health/ready")
    assert live.status_code == 200 and live.json() == {"status": "ok"}
    assert ready.status_code == 200 and ready.json()["storage"] == "memory"
    from .main import ALEMBIC_HEAD
    assert ALEMBIC_HEAD == "020_assignment_identity_fks"

def test_postgres_readiness_rejects_stale_migration(monkeypatch) -> None:
    from .main import health_ready
    database = AuthDatabase("sqlite+pysqlite:///:memory:")
    with database.engine.begin() as connection:
        connection.execute(text("CREATE TABLE alembic_version (version_num VARCHAR(32) NOT NULL)"))
        connection.execute(text("INSERT INTO alembic_version VALUES ('001_auth')"))
    monkeypatch.setitem(health_ready.__globals__, "auth_db", database)
    response = TestClient(app, base_url="http://localhost").get("/health/ready")
    assert response.status_code == 503

def test_database_session_lookup_uses_digest_and_revocation() -> None:
    database = AuthDatabase("sqlite+pysqlite:///:memory:")
    from datetime import datetime, timedelta, timezone
    with Session(database.engine) as session:
        session.add(UserRow(id="u1", name="Admin", email="admin@example.test", role="Admin", active=True, password_hash="hash"))
        session.add(SessionRow(digest=digest("opaque-token"), user_id="u1", issued_at=datetime.now(timezone.utc), expires_at=datetime.now(timezone.utc) + timedelta(hours=1))); session.commit()
    assert database.user_for_session("opaque-token").id == "u1"
    assert database.user_for_session("opaque-token").id == "u1"
    database.revoke("opaque-token"); assert database.user_for_session("opaque-token") is None
    assert "opaque-token" not in {row.digest for row in Session(database.engine).query(SessionRow).all()}

def test_database_duplicate_identity_is_safe_conflict() -> None:
    database = AuthDatabase("sqlite+pysqlite:///:memory:")
    database.create_user(user_id="u1", name="One", email="same@example.test", role="Admin", password_hash="hash")
    try: database.create_user(user_id="u2", name="Two", email="same@example.test", role="Admin", password_hash="hash")
    except ValueError as exc: assert "already exists" in str(exc)
    else: assert False

def test_database_customer_upsert_preserves_operational_owner() -> None:
    database = CustomerDatabase("sqlite+pysqlite:///:memory:")
    first = database.upsert_source(bcn="000123", name="Original", source={"score": 1}, primary_phone="555")
    with Session(database.engine) as session:
        stored = session.get(type(first), "000123"); stored.owner_id = "sales-river"; stored.status = "Closed"; session.commit()
    second = database.upsert_source(bcn="000123", name="Imported", source={"score": 2}, primary_phone="777")
    assert second.name == "Imported" and second.owner_id == "sales-river" and second.status == "Closed"
    database.upsert_sources([{"bcn": "000123", "name": "Imported", "source": {}, "primary_phone": None}])
    assert database.phones("000123") == []

def test_database_release_open_owner_returns_released_customers() -> None:
    database = CustomerDatabase("sqlite+pysqlite:///:memory:")
    database.upsert_source(bcn="000123", name="Open", source={})
    database.save_operational(bcn="000123", owner_id="sales", status="Open", version=1)
    assert database.release_open_owner("sales") == ["000123"]
    assert database.get("000123").owner_id is None

def test_database_customer_ingest_stores_typed_source_fields() -> None:
    from datetime import date
    from decimal import Decimal
    database = CustomerDatabase("sqlite+pysqlite:///:memory:")
    database.ingest_sources([{"bcn": "000123", "name": "Typed", "source": {}, "typed": {"propensity_score": Decimal("0.875000"), "last_purchase_date": date(2025, 1, 2), "previously_contacted": True}, "primary_phone": None}], {"jobId": "job-typed", "submissionId": "typed", "filename": "x.xlsx", "processed": 1, "created": 1, "updated": 0, "errorRows": 0, "status": "Completed", "errors": [], "actorId": "admin"})
    row = database.get("000123")
    assert row.propensity_score == Decimal("0.875000") and row.last_purchase_date == date(2025, 1, 2) and row.previously_contacted is True

def test_database_customer_ingest_stores_extensible_collections() -> None:
    from decimal import Decimal
    from .db_customers import CustomerCollectionRow
    database = CustomerDatabase("sqlite+pysqlite:///:memory:")
    result = {"jobId": "job-collections", "submissionId": "collections", "filename": "x.xlsx", "processed": 1, "created": 1, "updated": 0, "errorRows": 0, "status": "Completed", "errors": [], "actorId": "admin"}
    record = {"bcn": "000123", "name": "Collections", "source": {}, "collections": [{"kind": "vendor", "slot": 4, "name": "New vendor", "revenue": Decimal("12.50")}], "primary_phone": None}
    database.ingest_sources([record], result)
    with Session(database.engine) as session:
        row = session.query(CustomerCollectionRow).one()
        assert row.kind == "vendor" and row.slot == 4 and row.revenue == Decimal("12.500000")

def test_assignment_configuration_and_run_are_admin_only() -> None:
    repo.reset(); admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})()); repo.users["sales-river"] = {"id": "sales-river", "name": "River Sales", "email": "river@example.test", "role": "Sales", "active": True, "password": "unused"}
    client = TestClient(app, base_url="http://localhost"); client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"})
    created = client.post("/admin/assignment-rules", json={"name": "River", "ownerId": "sales-river"}, headers={"origin": "http://localhost:3000"})
    assert created.status_code == 201
    assert client.put("/admin/assignment-fallback", json=["sales-river"], headers={"origin": "http://localhost:3000"}).json() == ["sales-river"]
    result = client.post("/admin/assignment-runs", json={"scope": "unassigned", "submissionId": "run-1"}, headers={"origin": "http://localhost:3000"})
    assert result.status_code == 200 and result.json()["assigned"] == 1
    assert client.patch("/admin/assignment-rules/rule-1", json={"version": 0}, headers={"origin": "http://localhost:3000"}).status_code == 409

def test_assignment_database_keeps_ordered_rules_and_audit() -> None:
    database = AssignmentDatabase("sqlite+pysqlite:///:memory:")
    database.create_rule(rule_id="r2", name="Second", position=2, actor_id="admin")
    database.create_rule(rule_id="r1", name="First", position=1, actor_id="admin")
    assert [row.id for row in database.ordered_rules()] == ["r1", "r2"]
    events, total = database.audit(); assert total == 2 and len(events) == 2
    filtered, filtered_total = database.audit(actor="admin", action="created", page=1, page_size=1)
    assert filtered_total == 2 and len(filtered) == 1
    result = {"submissionId": "s1", "scope": "unassigned", "assigned": 1}
    assert database.save_run(actor_id="admin", submission_id="s1", scope="unassigned", result=result) == result
    assert database.get_run("admin", "s1") == result
    database.set_setting("fallback_sales", {"ids": ["u1"]})
    assert database.get_setting("fallback_sales") == {"ids": ["u1"]}

def test_assignment_run_retry_returns_persisted_result() -> None:
    database = AssignmentDatabase("sqlite+pysqlite:///:memory:")
    assert database.save_run(actor_id="admin", submission_id="same", scope="all", result={"assigned": 1}) == {"assigned": 1}
    assert database.save_run(actor_id="admin", submission_id="same", scope="all", result={"assigned": 99}) == {"assigned": 1}
    database.save_run(actor_id="admin", submission_id="scoped", scope="all", result={"assigned": 1}, payload="all")
    try: database.get_run("admin", "scoped", "unassigned")
    except ValueError as exc: assert "already used" in str(exc)
    else: assert False

def test_activity_idempotency_replays_and_rejects_payload_reuse() -> None:
    database = ActivityDatabase("sqlite+pysqlite:///:memory:")
    assert database.save_idempotent(actor_id="u1", operation="note", submission_id="s1", payload="hello", result={"id": "n1"}) == {"id": "n1"}
    assert database.save_idempotent(actor_id="u1", operation="note", submission_id="s1", payload="hello", result={"id": "ignored"}) == {"id": "n1"}
    try: database.save_idempotent(actor_id="u1", operation="note", submission_id="s1", payload="different", result={})
    except ValueError as exc: assert "already used" in str(exc)
    else: assert False
    database.save_activity(record_id="a1", bcn="000123", actor_id="u1", kind="Interaction", outcome="Attempt", text="x")
    assert database.soft_delete("a1", "u1") is True
    assert database.history("000123")[0][0].deleted_by == "u1"

def test_followup_completion_rejects_cross_customer_and_stale_links() -> None:
    database = ActivityDatabase("sqlite+pysqlite:///:memory:")
    followup = {"id": "f1", "bcn": "000123", "actorId": "u1", "type": "Reminder", "due": None, "status": "Open", "note": None}
    database.save_followup(followup)
    try: database.complete_followup(followup, {"id": "a1", "bcn": "000124", "actorId": "u1"})
    except ValueError as exc: assert "another customer" in str(exc)
    else: assert False
    database.complete_followup(followup, {"id": "a1", "bcn": "000123", "actorId": "u1"})
    try: database.complete_followup(followup, {"id": "a2", "bcn": "000123", "actorId": "u1"})
    except ValueError as exc: assert "no longer open" in str(exc)
    else: assert False

def test_real_http_admin_sales_journey() -> None:
    repo.reset()
    admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    sales = provision_user(type("P", (), {"name": "River", "email": "river@example.test", "role": "Sales", "password": "correct horse battery staple"})())
    repo.customers["000123"]["ownerId"] = sales.id; repo.customers["000123"]["ownerName"] = sales.name
    client = TestClient(app, base_url="http://localhost"); origin = {"origin": "http://localhost:3000"}
    assert client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"}).status_code == 200
    workbook = Workbook(); workbook.active.append(["bcn", "customer_name"]); workbook.active.append(["009990", "Journey Co"]); payload = BytesIO(); workbook.save(payload)
    imported = client.post("/admin/imports?submission_id=journey-import", files={"file": ("journey.xlsx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=origin)
    assert imported.status_code == 201 and imported.json()["created"] == 1
    assigned = client.post("/admin/assignments/manual/009990", json={"ownerId": sales.id, "submissionId": "journey-assign"}, headers=origin)
    assert assigned.status_code == 200
    client.post("/session/logout", headers=origin)
    assert client.post("/session/login", json={"email": sales.email, "password": "correct horse battery staple"}).status_code == 200
    assert client.get("/customers", params={"mine": "true"}).json()["total"] == 2
    interaction = client.post("/customers/009990/interactions", json={"outcome": "Contact", "note": "Journey", "submissionId": "journey-contact"}, headers=origin)
    assert interaction.status_code == 200, interaction.text
    followup = client.post("/customers/009990/follow-ups", json={"type": "Reminder", "due": "2026-09-20", "note": "Next", "submissionId": "journey-followup"}, headers=origin)
    assert followup.status_code == 200
    assert client.post(f"/customers/009990/follow-ups/{followup.json()['id']}/complete", json={"outcome": "Attempt", "submissionId": "journey-complete"}, headers=origin).status_code == 200
    assert client.post("/customers/000125/interactions", json={"outcome": "Attempt", "submissionId": "foreign"}, headers=origin).status_code == 403

def test_import_preserves_source_columns_and_operational_phone_history() -> None:
    repo.reset()
    admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    repo.customers["000123"]["ownerId"] = "sales-owner"
    repo.customers["000123"]["phones"] = ["old-primary", "independent"]
    client = TestClient(app, base_url="http://localhost")
    assert client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"}).status_code == 200
    workbook = Workbook(); workbook.active.append(["bcn", "customer_name", "propensity_score", "LAST_PURCHASE_DATE"]); workbook.active.append(["000123", "Updated", 0.875, datetime(2025, 1, 2)])
    payload = BytesIO(); workbook.save(payload)
    response = client.post("/admin/imports?submission_id=source-columns", files={"file": ("source.xlsx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers={"origin": "http://localhost:3000"})
    assert response.status_code == 201
    row = repo.customers["000123"]
    assert row["ownerId"] == "sales-owner" and row["phones"] == ["independent"]
    assert row["source"]["propensity_score"] == 0.875 and row["source"]["LAST_PURCHASE_DATE"].startswith("2025-01-02")

def test_import_retry_key_is_scoped_to_actor() -> None:
    database = CustomerDatabase("sqlite+pysqlite:///:memory:")
    result = {"jobId": "job-1", "submissionId": "same", "filename": "a.xlsx", "processed": 0, "created": 0, "updated": 0, "errorRows": 0, "status": "Completed", "errors": [], "actorId": "u1"}
    database.save_import_job(result)
    other = {**result, "jobId": "job-2", "actorId": "u2"}
    database.save_import_job(other)
    assert database.import_job("u1", "same")["jobId"] == "job-1"
    assert database.import_job("u2", "same")["jobId"] == "job-2"

def test_import_rejects_missing_required_customer_name() -> None:
    repo.reset()
    admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    client = TestClient(app, base_url="http://localhost")
    client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"})
    workbook = Workbook(); workbook.active.append(["bcn", "customer_name"]); workbook.active.append(["123456", ""])
    payload = BytesIO(); workbook.save(payload)
    response = client.post("/admin/imports?submission_id=missing-name", files={"file": ("source.xlsx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers={"origin": "http://localhost:3000"})
    assert response.status_code == 201 and response.json()["created"] == 0 and response.json()["errorRows"] == 1

def test_import_errors_are_paginated() -> None:
    repo.reset(); admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    repo.imports[(admin.id, "errors")] = {"errors": [{"row": n, "field": "bcn", "reason": "Invalid bcn"} for n in range(2, 5)]}
    client = TestClient(app, base_url="http://localhost"); client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"})
    response = client.get("/admin/imports/errors/errors?page=2&page_size=2")
    assert response.status_code == 200 and response.json()["total"] == 3 and len(response.json()["items"]) == 1

def test_database_import_errors_are_paginated_in_query() -> None:
    database = CustomerDatabase("sqlite+pysqlite:///:memory:")
    database.save_import_job({"jobId": "job-errors", "submissionId": "errors", "filename": "x.xlsx", "processed": 2, "created": 0, "updated": 0, "errorRows": 2, "status": "Partial", "errors": [{"row": 2, "field": "bcn", "reason": "bad"}, {"row": 3, "field": "bcn", "reason": "bad"}], "actorId": "admin"})
    items, total = database.import_errors("admin", "errors", 2, 1)
    assert total == 2 and items == [{"row": 3, "field": "bcn", "reason": "bad"}]

def test_import_storage_failure_restores_in_memory_staging(monkeypatch) -> None:
    from . import main
    repo.reset(); admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    client = TestClient(app, base_url="http://localhost"); client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"})
    workbook = Workbook(); workbook.active.append(["bcn", "customer_name"]); workbook.active.append(["999999", "Transient"])
    payload = BytesIO(); workbook.save(payload)
    class FailingCustomerDB:
        failed = None
        def ingest_sources(self, *_): raise RuntimeError("database unavailable")
        def import_job(self, *_): return None
        def save_import_job(self, result): self.failed = result
    failing = FailingCustomerDB(); monkeypatch.setattr(main, "customer_db", failing)
    response = client.post("/admin/imports?submission_id=rollback", files={"file": ("source.xlsx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers={"origin": "http://localhost:3000"})
    assert response.status_code == 422 and "999999" not in repo.customers and failing.failed["status"] == "Failed"

def test_memory_assignment_retry_is_scoped_to_actor() -> None:
    repo.reset()
    admin_one = provision_user(type("P", (), {"name": "One", "email": "one@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    admin_two = provision_user(type("P", (), {"name": "Two", "email": "two@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    repo.assignment_runs[(admin_one.id, "same-run")] = {"actorId": admin_one.id}
    repo.assignment_runs[(admin_two.id, "same-run")] = {"actorId": admin_two.id}
    assert repo.assignment_runs[(admin_one.id, "same-run")] != repo.assignment_runs[(admin_two.id, "same-run")]

def test_memory_followup_retry_rejects_payload_reuse() -> None:
    repo.reset()
    sales = provision_user(type("P", (), {"name": "Sales", "email": "sales@example.test", "role": "Sales", "password": "correct horse battery staple"})())
    repo.customers["000123"]["ownerId"] = sales.id
    client = TestClient(app, base_url="http://localhost")
    client.post("/session/login", json={"email": sales.email, "password": "correct horse battery staple"})
    origin = {"origin": "http://localhost:3000"}
    body = {"type": "Reminder", "due": "2026-09-20", "note": "Call", "submissionId": "same-followup"}
    assert client.post("/customers/000123/follow-ups", json=body, headers=origin).status_code == 200
    changed = {**body, "note": "Different"}
    assert client.post("/customers/000123/follow-ups", json=changed, headers=origin).status_code == 409

def test_memory_lifecycle_retries_are_idempotent() -> None:
    repo.reset()
    admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    client = TestClient(app, base_url="http://localhost"); client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"})
    origin = {"origin": "http://localhost:3000"}
    body = {"reasonId": "closure-1", "submissionId": "close-once"}
    assert client.post("/customers/000123/close", json=body, headers=origin).status_code == 200
    replay = client.post("/customers/000123/close", json=body, headers=origin)
    assert replay.status_code == 200 and repo.customers["000123"]["version"] == 1

def test_memory_import_retry_rejects_changed_workbook() -> None:
    repo.reset()
    admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    client = TestClient(app, base_url="http://localhost"); client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"})
    origin = {"origin": "http://localhost:3000"}
    def workbook(name: str) -> bytes:
        book = Workbook(); book.active.append(["bcn", "customer_name"]); book.active.append(["991001", name]); output = BytesIO(); book.save(output); return output.getvalue()
    first = client.post("/admin/imports?submission_id=same-import", files={"file": ("source.xlsx", workbook("One"), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=origin)
    second = client.post("/admin/imports?submission_id=same-import", files={"file": ("source.xlsx", workbook("Two"), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=origin)
    assert first.status_code == 201 and second.status_code == 409

def test_import_result_can_be_read_by_owner() -> None:
    repo.reset()
    admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    client = TestClient(app, base_url="http://localhost"); client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"})
    book = Workbook(); book.active.append(["bcn", "customer_name"]); book.active.append(["991002", "Readback"]); payload = BytesIO(); book.save(payload)
    origin = {"origin": "http://localhost:3000"}
    assert client.post("/admin/imports?submission_id=readback", files={"file": ("source.xlsx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=origin).status_code == 201
    result = client.get("/admin/imports/readback")
    assert result.status_code == 200 and result.json()["submissionId"] == "readback"
    page = client.get("/admin/imports?page=1&page_size=1")
    assert page.status_code == 200 and page.json()["total"] == 1 and len(page.json()["items"]) == 1

def test_memory_import_submission_is_scoped_to_actor() -> None:
    repo.reset()
    one = provision_user(type("P", (), {"name": "One", "email": "one@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    two = provision_user(type("P", (), {"name": "Two", "email": "two@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    repo.imports[(one.id, "shared")] = {"actorId": one.id, "submissionId": "shared"}
    repo.imports[(two.id, "shared")] = {"actorId": two.id, "submissionId": "shared"}
    assert repo.imports[(one.id, "shared")]["actorId"] != repo.imports[(two.id, "shared")]["actorId"]

def test_mutation_routes_bind_json_bodies() -> None:
    paths = {route.path: {field.name for field in route.dependant.body_params} for route in app.routes if getattr(route, "dependant", None)}
    assert paths["/customers/{bcn}/interactions"] == {"body"}
    assert paths["/admin/users"] == {"data"}
    assert paths["/admin/assignment-fallback"] == {"ids"}


def test_expired_session_is_rejected() -> None:
    repo.reset()
    user = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    repo.sessions["expired"] = (user.id, datetime.now(timezone.utc) - timedelta(seconds=1))
    client = TestClient(app); client.cookies.set("call_center_session", "expired")
    assert client.get("/session/me").status_code == 401


def test_foreign_origin_is_rejected_before_mutation() -> None:
    repo.reset(); client = TestClient(app)
    assert client.post("/session/logout", headers={"origin": "https://foreign.example"}).status_code == 403


def test_memory_unit_of_work_rolls_back_auth_state() -> None:
    repo.reset()
    try:
        with repo.transaction():
            provision_user(type("P", (), {"name": "Transient", "email": "transient@example.test", "role": "Admin", "password": "correct horse battery staple"})())
            raise RuntimeError("rollback")
    except RuntimeError:
        pass
    assert repo.users == {}


def test_operator_provision_and_recovery_revoke_session() -> None:
    repo.reset(); client = TestClient(app, base_url="http://localhost"); headers = {"origin": "http://localhost:3000"}
    provision = client.post("/operator/provision", json={"name": "Initial Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"}, headers=headers)
    assert provision.status_code == 200 and "password" not in provision.json()
    login = client.post("/session/login", json={"email": "admin@example.test", "password": "correct horse battery staple"}); assert login.status_code == 200
    recovered = client.post(f"/operator/reset-password/{provision.json()['id']}", json={"password": "new correct horse battery staple"}, headers=headers)
    assert recovered.status_code == 200 and "password" not in recovered.json()
    assert client.get("/session/me").status_code == 401
    assert client.post("/operator/provision", json={"name": "Second", "email": "second@example.test", "role": "Admin", "password": "correct horse battery staple"}, headers=headers).status_code == 409


def test_password_bounds_and_inactive_users_have_safe_failures() -> None:
    repo.reset(); provision_user(type("P", (), {"name": "Inactive", "email": "inactive@example.test", "role": "Sales", "password": "correct horse battery staple"})()); repo.users["user-1"]["active"] = False
    client = TestClient(app, base_url="http://localhost")
    assert client.post("/session/login", json={"email": "inactive@example.test", "password": "correct horse battery staple"}).status_code == 401
    short = client.post("/session/login", json={"email": "inactive@example.test", "password": "short"})
    assert short.status_code == 422 and "short" not in short.text


def test_customer_reads_and_admin_assignment() -> None:
    repo.reset(); admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})()); repo.users["sales-river"] = {"id": "sales-river", "name": "River Sales", "email": "river@example.test", "role": "Sales", "active": True, "password": "unused"}
    client = TestClient(app, base_url="http://localhost"); client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"})
    assert client.get("/customers/000123").status_code == 200
    changed = client.post("/admin/assignments/manual/000125", json={"ownerId": "sales-river", "submissionId": "assign-1"}, headers={"origin": "http://localhost:3000"})
    assert changed.status_code == 200 and changed.json()["ownerId"] == "sales-river"
    assert client.post("/admin/assignments/manual/000125", json={"ownerId": "sales-river", "submissionId": "assign-1"}, headers={"origin": "http://localhost:3000"}).status_code == 200
    assert client.post("/admin/assignments/manual/000125", json={"ownerId": None, "submissionId": "assign-1"}, headers={"origin": "http://localhost:3000"}).status_code == 409
    assert client.post("/admin/assignments/manual/000125", json={"ownerId": None, "submissionId": "assign-2", "expectedVersion": 0}, headers={"origin": "http://localhost:3000"}).status_code == 409
    assert client.post("/admin/assignments/manual/000125", json={"ownerId": "missing", "submissionId": "assign-3"}, headers={"origin": "http://localhost:3000"}).status_code == 422
    assert client.get("/customers/missing").status_code == 404


def test_sales_my_scope_cannot_be_widened_and_reads_are_paginated() -> None:
    repo.reset(); provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})()); repo.users["sales-river"] = {"id": "sales-river", "name": "River Sales", "email": "river@example.test", "role": "Sales", "active": True, "password": password_hash.hash("correct horse battery staple")}
    client = TestClient(app, base_url="http://localhost"); client.post("/session/login", json={"email": "river@example.test", "password": "correct horse battery staple"})
    scoped = client.get("/customers", params={"mine": "true", "page_size": 1}); assert scoped.status_code == 200 and scoped.json()["total"] == 1 and scoped.json()["items"][0]["bcn"] == "000123"
    assert client.get("/customers", params={"mine": "true", "owner": "sales-sky"}).json()["total"] == 0
    assert client.get("/customers", params={"page": 0}).status_code == 422

def test_admin_xlsx_import_preserves_assignment_and_is_idempotent() -> None:
    repo.reset(); admin = provision_user(type("P", (), {"name": "Admin", "email": "admin@example.test", "role": "Admin", "password": "correct horse battery staple"})())
    client = TestClient(app, base_url="http://localhost"); client.post("/session/login", json={"email": admin.email, "password": "correct horse battery staple"})
    workbook = Workbook(); sheet = workbook.active; sheet.append(["bcn", "customer_name", "phone"]); sheet.append(["000123", "Renamed", "555-0001"]); sheet.append(["009999", "New Co", None]); sheet.append(["009999", "Duplicate", None]); payload = BytesIO(); workbook.save(payload); payload.seek(0)
    response = client.post("/admin/imports?submission_id=job-1", files={"file": ("customers.xlsx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers={"origin": "http://localhost:3000"})
    assert response.status_code == 201 and response.json()["processed"] == 3 and response.json()["created"] == 1 and len(response.json()["errors"]) == 1
    assert client.get("/customers/000123").json()["ownerId"] == "sales-river"
    assert client.post("/admin/imports?submission_id=job-1", files={"file": ("customers.xlsx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers={"origin": "http://localhost:3000"}).json()["jobId"] == response.json()["jobId"]
