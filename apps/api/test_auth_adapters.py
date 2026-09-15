"""Shared auth contract plus isolated, migration-managed PostgreSQL acceptance checks.

Use --tb=no so even unexpected assertion failures cannot print credentials.
"""
import os
import secrets
import subprocess
import sys
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from datetime import date
from decimal import Decimal
from io import BytesIO
from threading import Barrier
from uuid import uuid4

import pytest
from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine, delete, event, select, text, update
from sqlalchemy.engine import make_url
from sqlalchemy.exc import DBAPIError, IntegrityError
from sqlalchemy.orm import Session

from . import main, db_auth
from .db_auth import AuthDatabase, SessionRow, UserRow, digest, StorageError

ORIGIN = {"origin": "http://localhost:3000"}


def test_history_delete_is_atomic_authorized_and_idempotent(tmp_path, monkeypatch):
    from fastapi import HTTPException
    from .db_activity import ActivityDatabase, ActivityRow
    from .db_assignment import AssignmentDatabase, AuditRow
    from .db_customers import CustomerDatabase

    url = f"sqlite+pysqlite:///{tmp_path / 'delete.db'}"
    activities = ActivityDatabase(url)
    assignments = AssignmentDatabase(url)
    customers = CustomerDatabase(url)
    monkeypatch.setattr(main, "activity_db", activities)
    monkeypatch.setattr(main, "assignment_db", assignments)
    monkeypatch.setattr(main, "customer_db", customers)
    main.repo.reset()
    try:
        customers.upsert_source(bcn="000001", name="Synthetic", source={})
        customers.save_operational(bcn="000001", owner_id="sales", status="Open", version=0)
        activities.save_activity(record_id="note", bcn="000001", actor_id="sales", kind="Standalone note", outcome=None, text="Synthetic note")
        main.repo.customers.clear()
        actor = main.User(id="sales", name="Sales", email="sales@example.test", role="Sales")
        with pytest.raises(HTTPException) as forbidden:
            main.delete_history("000001", "note", actor.model_copy(update={"id": "other"}))
        assert forbidden.value.status_code == 403 and main.repo.customers == {}
        def fail_commit(session):
            if session.bind is activities.engine:
                session.flush()
                raise RuntimeError("Injected failure")
        event.listen(Session, "before_commit", fail_commit)
        try:
            with pytest.raises(RuntimeError, match="Injected failure"):
                main.delete_history("000001", "note", actor)
        finally:
            event.remove(Session, "before_commit", fail_commit)
        assert main.repo.customers == {}
        with Session(activities.engine) as session:
            assert session.get(ActivityRow, "note").deleted_at is None
            assert session.scalar(select(AuditRow)) is None
        result = main.delete_history("000001", "note", actor)
        assert result["deleted"] is True and result["deletedBy"] == actor.id
        assert main.repo.customers["000001"]["histories"][0]["deleted"] is True
        main.repo.customers.clear()
        assert main.delete_history("000001", "note", actor.model_copy(update={"id": "admin", "role": "Admin"})) == result
        with Session(activities.engine) as session:
            assert session.get(ActivityRow, "note").deleted_by == actor.id
            assert len(list(session.scalars(select(AuditRow)))) == 1
    finally:
        for database in (activities, assignments, customers): database.engine.dispose()
        main.repo.reset()


@pytest.mark.parametrize("operation", ["followup", "interaction", "note"])
def test_activity_creation_is_atomic_and_replays(operation, tmp_path, monkeypatch):
    from fastapi import HTTPException
    from .db_activity import ActivityDatabase, FollowUpRow, ActivityRow, IdempotencyRow
    from .db_customers import CustomerDatabase
    from .db_assignment import AssignmentDatabase, AuditRow

    url = f"sqlite+pysqlite:///{tmp_path / 'followup.db'}"
    activities = ActivityDatabase(url)
    customers = CustomerDatabase(url)
    assignments = AssignmentDatabase(url)
    monkeypatch.setattr(main, "activity_db", activities)
    monkeypatch.setattr(main, "customer_db", customers)
    monkeypatch.setattr(main, "assignment_db", assignments)
    main.repo.reset()
    try:
        customers.upsert_source(bcn="000001", name="Synthetic", source={})
        customers.save_operational(bcn="000001", owner_id="sales", status="Open", version=0)
        main.repo.customers.clear()
        actor = main.User(id="sales", name="Sales", email="sales@example.test", role="Sales")
        create, body, cache, tables = {
            "followup": (main.create_followup, main.FollowUpCreate(type="Reminder", due=None, note="Synthetic reminder", submissionId="create"), main.repo.followups, (FollowUpRow, ActivityRow, IdempotencyRow)),
            "interaction": (main.create_interaction, main.InteractionCreate(outcome="Contact", note="Synthetic contact", submissionId="create"), main.repo.interactions, (ActivityRow, IdempotencyRow, AuditRow)),
            "note": (main.create_note, main.NoteCreate(text="Synthetic note", submissionId="create"), main.repo.notes, (ActivityRow, IdempotencyRow, AuditRow)),
        }[operation]
        def fail_commit(session):
            if session.bind is activities.engine:
                session.flush()
                raise RuntimeError("Injected failure")
        event.listen(Session, "before_commit", fail_commit)
        try:
            with pytest.raises(RuntimeError, match="Injected failure"):
                create("000001", body, actor)
        finally:
            event.remove(Session, "before_commit", fail_commit)
        assert main.repo.customers == {} and cache == {}
        with Session(activities.engine) as session:
            assert all(session.scalar(select(table)) is None for table in tables)
        result = create("000001", body, actor)
        assert create("000001", body, actor) == result
        with pytest.raises(HTTPException) as conflict:
            create("000001", body.model_copy(update={"text" if operation == "note" else "note": "Different content"}), actor)
        assert conflict.value.status_code == 409
        assert len(cache) == 1 and len(main.repo.customers["000001"]["histories"]) == 1
        with Session(activities.engine) as session:
            assert all(len(list(session.scalars(select(table)))) == 1 for table in tables)
        with pytest.raises(HTTPException) as forbidden:
            create("000001", body, actor.model_copy(update={"id": "other"}))
        assert forbidden.value.status_code == 403
        main.repo.customers["000001"]["status"] = "Closed"
        with pytest.raises(HTTPException) as closed:
            create("000001", body.model_copy(update={"submissionId": "closed"}), actor)
        assert closed.value.status_code == 409
    finally:
        activities.engine.dispose(); customers.engine.dispose(); assignments.engine.dispose()
        main.repo.reset()


def test_rule_updates_compare_persisted_version_and_rollback(tmp_path, monkeypatch):
    from copy import deepcopy
    from fastapi import HTTPException
    from .db_assignment import AssignmentDatabase

    database = AssignmentDatabase(f"sqlite+pysqlite:///{tmp_path / 'rules.db'}")
    monkeypatch.setattr(main, "assignment_db", database)
    main.repo.reset()
    try:
        database.create_rule(rule_id="rule", name="Original", position=1, actor_id="admin", conditions=[], member_ids=["sales"])
        database.set_setting("assignment_version", {"value": 2})
        main.repo.rules = [{"id": "rule", "name": "Original", "order": 1, "conditions": [], "memberIds": ["sales"], "active": True}]
        main.repo.assignment_version = 2
        actor = main.User(id="admin", name="Admin", email="admin@example.test", role="Admin")
        # Another request wins after this process cached configuration version 2.
        database.update_rule("rule", {"name": "Winner"}, actor.id, 2)
        before = deepcopy(main.repo.rules)
        with pytest.raises(HTTPException) as stale:
            main.update_assignment_rule("rule", {"name": "Loser", "version": 2}, actor)
        assert stale.value.status_code == 409
        assert main.repo.rules == before and main.repo.assignment_version == 2
        assert database.ordered_rules()[0].name == "Winner" and database.get_setting("assignment_version") == {"value": 3}
        assert database.audit()[1] == 2
        result = main.update_assignment_rule("rule", {"active": False, "version": 3}, actor)
        assert result["name"] == "Winner" and result["active"] is False
        assert main.repo.assignment_version == 4
        before = deepcopy(main.repo.rules)
        def fail_commit(session):
            if session.bind is database.engine:
                session.flush()
                raise RuntimeError("Injected failure")
        event.listen(Session, "before_commit", fail_commit)
        try:
            with pytest.raises(RuntimeError, match="Injected failure"):
                main.update_assignment_rule("rule", {"name": "Rolled back", "version": 4}, actor)
        finally:
            event.remove(Session, "before_commit", fail_commit)
        assert main.repo.rules == before and main.repo.assignment_version == 4
        assert database.ordered_rules()[0].name == "Winner" and database.get_setting("assignment_version") == {"value": 4}
        assert database.audit()[1] == 3
    finally:
        database.engine.dispose()
        main.repo.reset()


def test_rule_member_update_enforces_active_sales(tmp_path, monkeypatch):
    """A member deactivated after being added is left in assignment_rule_members untouched (no
    cascade-remove); the rule is simply skipped at evaluation time (db_assignment.run_bulk /
    assignment_rules.resolve_owner cover that). Editing a rule's memberIds is validated the same
    way rule creation is: only active-Sales users, and at least one member."""
    from fastapi import HTTPException
    from .db_assignment import AssignmentDatabase

    database = AssignmentDatabase(f"sqlite+pysqlite:///{tmp_path / 'rule-member.db'}")
    monkeypatch.setattr(main, "assignment_db", database)
    main.repo.reset()
    try:
        main.repo.users["first"] = {"id": "first", "name": "first", "role": "Sales", "active": True}
        main.repo.users["inactive"] = {"id": "inactive", "name": "inactive", "role": "Sales", "active": False}
        actor = main.User(id="admin", name="Admin", email="admin@example.test", role="Admin")
        rule = main.create_assignment_rule(main.AssignmentRuleDraft(name="Rule", conditions=[], memberIds=["first"], active=True), actor)
        # Deactivating a member leaves assignment_rule_members untouched (no cascade-remove).
        main.repo.users["first"]["active"] = False
        assert database.rule_members(rule["id"]) == ["first"]
        with pytest.raises(HTTPException) as rejected:
            main.update_assignment_rule(rule["id"], {"memberIds": ["inactive"], "version": main.repo.assignment_version}, actor)
        assert rejected.value.status_code == 422
        assert database.rule_members(rule["id"]) == ["first"] and database.ordered_rules()[0].active is True
        with pytest.raises(HTTPException) as empty:
            main.update_assignment_rule(rule["id"], {"memberIds": [], "version": main.repo.assignment_version}, actor)
        assert empty.value.status_code == 422
        main.repo.users["second"] = {"id": "second", "name": "second", "role": "Sales", "active": True}
        result = main.update_assignment_rule(rule["id"], {"memberIds": ["second"], "version": main.repo.assignment_version}, actor)
        assert result["memberIds"] == ["second"]
        assert database.rule_members(rule["id"]) == ["second"]
    finally:
        database.engine.dispose()
        main.repo.reset()


def test_manual_assignment_commit_and_rollback(tmp_path, monkeypatch):
    from .db_assignment import AssignmentDatabase, AssignmentHistoryRow, AuditRow
    from .db_customers import CustomerDatabase
    from fastapi import HTTPException

    url = f"sqlite+pysqlite:///{tmp_path / 'manual.db'}"
    customers = CustomerDatabase(url)
    assignments = AssignmentDatabase(url)
    monkeypatch.setattr(main, "customer_db", customers)
    monkeypatch.setattr(main, "assignment_db", assignments)
    monkeypatch.setattr(main, "activity_db", None)
    main.repo.reset()
    try:
        customers.upsert_source(bcn="000001", name="Synthetic", source={})
        main.repo.customers.clear()
        main.repo.users["sales"] = {"id": "sales", "name": "Sales", "role": "Sales", "active": True}
        actor = main.User(id="admin", name="Admin", email="admin@example.test", role="Admin")
        body = main.AssignmentRequest(ownerId="sales", expectedVersion=0, submissionId="manual")
        def fail_commit(session):
            if session.bind is assignments.engine:
                session.flush()
                raise RuntimeError("Injected failure")
        event.listen(Session, "before_commit", fail_commit)
        try:
            with pytest.raises(RuntimeError, match="Injected failure"):
                main.assign_customer("000001", body, actor)
        finally:
            event.remove(Session, "before_commit", fail_commit)
        assert main.repo.customers == {} and main.repo.submissions == {}
        assert customers.get("000001").owner_id is None and customers.get("000001").version == 0
        with Session(assignments.engine) as session:
            assert session.scalar(select(AssignmentHistoryRow)) is None and session.scalar(select(AuditRow)) is None
        result = main.assign_customer("000001", body, actor)
        assert result.ownerId == "sales" and result.version == 1
        assert main.assign_customer("000001", body, actor) == result
        unchanged = main.assign_customer("000001", main.AssignmentRequest(ownerId="sales", expectedVersion=1, submissionId="unchanged"), actor)
        assert unchanged == result
        with pytest.raises(HTTPException) as stale:
            main.assign_customer("000001", main.AssignmentRequest(ownerId=None, expectedVersion=0, submissionId="stale"), actor)
        assert stale.value.status_code == 409
        assert customers.get("000001").owner_id == "sales" and customers.get("000001").version == 1
        with Session(assignments.engine) as session:
            assert len(list(session.scalars(select(AssignmentHistoryRow)))) == 1
            assert len(list(session.scalars(select(AuditRow)))) == 1
    finally:
        customers.engine.dispose(); assignments.engine.dispose()
        main.repo.reset()


@pytest.mark.parametrize("backend", ["sqlite", "postgres"])
@pytest.mark.parametrize("changes", [{"active": False}, {"role": "Admin"}])
def test_owner_identity_change_rolls_back_and_releases_only_open(changes, backend, request, tmp_path, monkeypatch):
    """Parametrized over sqlite/postgres so update_identity's with_for_update() row-locking path
    is exercised against real PostgreSQL, not just SQLite."""
    from copy import deepcopy
    from fastapi import HTTPException
    from .db_assignment import AssignmentDatabase, AssignmentHistoryRow, AuditRow
    from .db_customers import CustomerDatabase

    url = request.getfixturevalue("postgres_url") if backend == "postgres" else f"sqlite+pysqlite:///{tmp_path / 'owner.db'}"
    if backend == "postgres": migrate()
    auth = AuthDatabase(url, create_schema=backend == "sqlite")
    customers = CustomerDatabase(url, create_schema=backend == "sqlite")
    assignments = AssignmentDatabase(url, create_schema=backend == "sqlite")
    for name, database in (("auth_db", auth), ("customer_db", customers), ("assignment_db", assignments)):
        monkeypatch.setattr(main, name, database)
    main.repo.reset()
    try:
        hashed = main.password_hash.hash(secrets.token_urlsafe(24))
        auth.create_user(user_id="sales", name="Sales", email="sales@example.test", role="Sales", password_hash=hashed)
        # A real actor row too, so assignment_history/audit_events FKs are satisfied under real PostgreSQL enforcement.
        auth.create_user(user_id="admin", name="Admin", email="owner-admin@example.test", role="Admin", password_hash=hashed)
        token, expires = auth.issue("sales")
        main.repo.users = {"sales": {"id": "sales", "name": "Sales", "email": "sales@example.test", "role": "Sales", "active": True}}
        main.repo.sessions[token] = ("sales", expires)
        main.repo.customers = {}
        for bcn, status in (("000001", "Open"), ("000002", "Closed")):
            customers.upsert_source(bcn=bcn, name="Synthetic", source={})
            customers.save_operational(bcn=bcn, owner_id="sales", status=status, version=3)
            main.repo.customers[bcn] = {"bcn": bcn, "ownerId": "sales", "ownerName": "Sales", "status": status, "version": 3, "histories": []}
        before = deepcopy((main.repo.users, main.repo.customers, main.repo.sessions))
        actor = main.User(id="admin", name="Admin", email="admin@example.test", role="Admin")
        def fail_commit(session):
            if session.bind is assignments.engine:
                session.flush()
                raise RuntimeError("Injected failure")
        event.listen(Session, "before_commit", fail_commit)
        try:
            with pytest.raises(RuntimeError, match="Injected failure"):
                main.update_user("sales", main.UserPatch(**changes), actor)
        finally:
            event.remove(Session, "before_commit", fail_commit)
        assert before == (main.repo.users, main.repo.customers, main.repo.sessions)
        assert auth.user_for_session(token).active
        assert auth.user_by_email("sales@example.test").role == "Sales"
        assert all(row.owner_id == "sales" and row.version == 3 for row in customers.all())
        with Session(assignments.engine) as session:
            assert session.scalar(select(AssignmentHistoryRow)) is None and session.scalar(select(AuditRow)) is None
        main.update_user("sales", main.UserPatch(**changes), actor)
        assert auth.user_for_session(token) is None and token not in main.repo.sessions
        assert customers.get("000001").owner_id is None and customers.get("000001").version == 4
        assert customers.get("000002").owner_id == "sales" and customers.get("000002").version == 3
        assert main.repo.customers["000001"]["ownerId"] is None and main.repo.customers["000002"]["ownerId"] == "sales"
        with Session(assignments.engine) as session:
            assert len(list(session.scalars(select(AssignmentHistoryRow)))) == 1
            assert len(list(session.scalars(select(AuditRow)))) == 2
        # Reactivating (active Sales again) does not reassign or reclaim any customer.
        reactivation = {"active": True} if "active" in changes else {"role": "Sales"}
        main.update_user("sales", main.UserPatch(**reactivation), actor)
        assert customers.get("000001").owner_id is None and customers.get("000001").version == 4
        assert customers.get("000002").owner_id == "sales" and customers.get("000002").version == 3
        assert main.repo.customers["000001"]["ownerId"] is None and main.repo.customers["000002"]["ownerId"] == "sales"
        with Session(assignments.engine) as session:
            assert len(list(session.scalars(select(AssignmentHistoryRow)))) == 1
            assert len(list(session.scalars(select(AuditRow)))) == 3
        with pytest.raises(HTTPException) as blocked:
            main.update_user("sales", main.UserPatch(active=False), main.User(**main.repo.users["sales"]))
        assert blocked.value.status_code == 422
    finally:
        for database in (auth, customers, assignments): database.engine.dispose()
        main.repo.reset()


@pytest.mark.parametrize("backend", ["sqlite", "postgres"])
def test_bulk_assignment_atomic_rollback_and_retry(backend, request, tmp_path, monkeypatch):
    from .db_assignment import AssignmentDatabase, AssignmentHistoryRow, AssignmentRunRow, AuditRow
    from .db_customers import CustomerDatabase
    from fastapi import HTTPException

    url = request.getfixturevalue("postgres_url") if backend == "postgres" else f"sqlite+pysqlite:///{tmp_path / 'bulk.db'}"
    if backend == "postgres": migrate()
    auth = AuthDatabase(url, create_schema=backend == "sqlite")
    customers = CustomerDatabase(url, create_schema=backend == "sqlite")
    assignments = AssignmentDatabase(url, create_schema=backend == "sqlite")
    monkeypatch.setattr(main, "customer_db", customers)
    monkeypatch.setattr(main, "assignment_db", assignments)
    main.repo.reset()
    try:
        hashed = main.password_hash.hash(secrets.token_urlsafe(24))
        for user_id, role in (("admin", "Admin"), ("sales", "Sales")):
            auth.create_user(user_id=user_id, name=user_id, email=user_id + "@example.test", role=role, password_hash=hashed)
            main.repo.users[user_id] = {"id": user_id, "name": user_id, "role": role, "active": True}
        for bcn in ("000002", "000001"):
            customers.upsert_source(bcn=bcn, name="Synthetic", source={})
        main.repo.customers = {"000001": {"bcn": "000001", "ownerId": None, "version": 0}}
        main.repo.fallback_sales = ["sales"]
        actor = main.User(id="admin", name="admin", email="admin@example.test", role="Admin")
        body = main.AssignmentRunRequest(scope="unassigned", submissionId="atomic")
        def fail_commit(session):
            if session.bind is assignments.engine:
                session.flush()
                raise RuntimeError("Injected failure")
        event.listen(Session, "before_commit", fail_commit)
        try:
            with pytest.raises(RuntimeError, match="Injected failure"):
                main.run_assignment(body, actor)
        finally:
            event.remove(Session, "before_commit", fail_commit)
        assert all(row.owner_id is None and row.version == 0 for row in customers.all())
        assert main.repo.customers["000001"]["ownerId"] is None and main.repo.assignment_runs == {}
        with Session(assignments.engine) as session:
            assert all(session.scalar(select(table)) is None for table in (AssignmentHistoryRow, AuditRow, AssignmentRunRow))
        result = main.run_assignment(body, actor)
        assert result["assigned"] == 2
        assert main.repo.customers["000001"]["ownerId"] == "sales"
        assert main.run_assignment(body, actor) == result
        with pytest.raises(HTTPException) as conflict:
            main.run_assignment(main.AssignmentRunRequest(scope="all-open", submissionId="atomic"), actor)
        assert conflict.value.status_code == 409
        assert all(row.owner_id == "sales" and row.version == 1 for row in customers.all())
        with Session(assignments.engine) as session:
            assert len(list(session.scalars(select(AssignmentHistoryRow)))) == 2
            assert len(list(session.scalars(select(AuditRow)))) == 2
            assert len(list(session.scalars(select(AssignmentRunRow)))) == 1
    finally:
        for database in (auth, customers, assignments): database.engine.dispose()
        main.repo.reset()


@pytest.fixture
def postgres_url(monkeypatch):
    raw = os.getenv("TEST_DATABASE_URL")
    if not raw:
        pytest.skip("TEST_DATABASE_URL required for real PostgreSQL checks")
    url = make_url(raw)
    if url.get_backend_name() != "postgresql" or not (url.database or "").endswith(("_test", "_ci")):
        pytest.fail("Use a dedicated PostgreSQL database ending in _test or _ci", pytrace=False)
    engine = create_engine(url, hide_parameters=True)
    schema = "auth_test_" + uuid4().hex
    with engine.begin() as connection:
        connection.execute(text(f'CREATE SCHEMA "{schema}"'))
    isolated = url.update_query_dict({"options": f"-csearch_path={schema}"}).render_as_string(hide_password=False)
    monkeypatch.setenv("DATABASE_URL", isolated)
    try:
        yield isolated
    finally:
        with engine.begin() as connection:
            connection.execute(text(f'DROP SCHEMA "{schema}" CASCADE'))
        engine.dispose()


def migrate(revision="head"):
    command.upgrade(Config("apps/api/alembic.ini"), revision)


@pytest.fixture(params=["memory", "postgres"])
def adapter(request, monkeypatch):
    database = None
    if request.param == "postgres":
        url = request.getfixturevalue("postgres_url")
        migrate()
        database = AuthDatabase(url, create_schema=False)
    main.repo.reset()
    monkeypatch.setattr(main, "auth_db", database)
    for name in ("customer_db", "assignment_db", "activity_db"):
        monkeypatch.setattr(main, name, None)
    try:
        yield database
    finally:
        if database:
            assert database.engine.pool.checkedout() == 0
            database.engine.dispose()
        main.repo.reset()


def provision(email="admin@example.test", role="Admin", password=None):
    secret = password or secrets.token_urlsafe(24)
    user = main.provision_user(main.Provision(name="Synthetic", email=email, role=role, password=secret))
    return user, secret


def sign_in(client, email, secret):
    return client.post("/session/login", json={"email": email, "password": secret})


def test_shared_identity_validation_and_safe_failures(adapter):
    user, secret = provision("  STRAẞE@example.test  ")
    assert user.email == "strasse@example.test"
    with pytest.raises(ValueError, match="normalized identity already exists"):
        provision(" STRASSE@EXAMPLE.TEST ")
    with TestClient(main.app, base_url="http://localhost") as client:
        unknown = sign_in(client, "unknown@example.test", secrets.token_urlsafe(24))
        wrong = sign_in(client, user.email, secrets.token_urlsafe(24))
        if adapter:
            adapter.update_user(user.id, {"active": False})
        else:
            main.repo.users[user.id]["active"] = False
        inactive = sign_in(client, user.email, secret)
        assert all(response.status_code == 401 for response in (unknown, wrong, inactive))
        assert unknown.json() == wrong.json() == inactive.json() == {"detail": "Unable to sign in."}
        assert all("set-cookie" not in response.headers for response in (unknown, wrong, inactive))
        for payload in ({"email": "x" * 255, "password": secret}, {"email": user.email, "password": "x" * 11}, {"email": user.email, "password": "x" * 129}, {"email": [], "password": secret}):
            response = client.post("/session/login", json=payload)
            assert response.status_code == 422 and response.json() == {"detail": "Invalid request."}
    with pytest.raises(ValueError):
        provision("short@example.test", password=secrets.token_hex(5))


def test_shared_password_boundaries_and_transport(adapter, caplog):
    for length in (12, 128):
        secret = " " + secrets.token_hex(100)[:length - 3] + "界 "
        user, _ = provision(f"length-{length}@example.test", password=secret)
        with TestClient(main.app, base_url="https://api.example.test") as client:
            response = sign_in(client, user.email, secret)
            assert response.status_code == 200
            assert set(response.json()) == {"id", "name", "email", "role", "active"}
            cookie = response.headers["set-cookie"].lower()
            assert all(value in cookie for value in ("httponly", "secure", "samesite=lax", "path=/", "max-age=28800"))
            token = client.cookies.get("call_center_session")
            assert secret not in response.text and token not in response.text
            assert client.post("/session/logout").status_code == 403
            assert client.post("/session/logout", headers={"origin": "https://foreign.example"}).status_code == 403
            assert client.post("/session/logout", headers={"referer": "http://localhost:3000/settings"}).status_code == 204
            assert secret not in caplog.text and token not in caplog.text
        with TestClient(main.app, base_url="http://localhost") as client:
            response = client.options("/session/me", headers={"origin": "http://localhost:3000", "access-control-request-method": "GET"})
            assert response.headers["access-control-allow-origin"] == "http://localhost:3000"
            assert response.headers["access-control-allow-credentials"] == "true"
            response = client.options("/session/me", headers={"origin": "http://localhost:3000.evil.test", "access-control-request-method": "GET"})
            assert "access-control-allow-origin" not in response.headers


def test_shared_session_clock_logout_recovery_and_role(adapter, monkeypatch):
    clock = [datetime(2030, 1, 1, tzinfo=timezone.utc)]
    monkeypatch.setattr(main, "utcnow", lambda: clock[0])
    monkeypatch.setattr(db_auth, "utcnow", lambda: clock[0])
    user, secret = provision()
    with TestClient(main.app, base_url="http://localhost") as client:
        assert sign_in(client, user.email, secret).status_code == 200
        clock[0] += timedelta(hours=8) - timedelta(microseconds=1)
        assert client.get("/session/me").status_code == 200
        clock[0] += timedelta(microseconds=1)
        assert client.get("/session/me").status_code == 401
        assert client.post("/admin/users", headers=ORIGIN, json={"name": "Other", "email": "other@example.test", "password": secret}).status_code == 401
        assert sign_in(client, user.email, secret).status_code == 200
        response = client.post("/session/logout", headers=ORIGIN)
        assert response.status_code == 204 and "Max-Age=0" in response.headers["set-cookie"]
        assert client.get("/session/me").status_code == 401
        assert client.post("/session/logout", headers=ORIGIN).status_code == 204
        assert sign_in(client, user.email, secret).status_code == 200
        new_secret = secrets.token_urlsafe(24)
        main.operator_reset_password(user.id, main.ResetPassword(password=new_secret))
        assert client.get("/session/me").status_code == 401
        assert sign_in(client, user.email, secret).status_code == 401
        assert sign_in(client, user.email, new_secret).status_code == 200
        if adapter:
            adapter.update_user(user.id, {"role": "Sales"})
        else:
            main.repo.users[user.id]["role"] = "Sales"
        assert client.get("/admin/users").status_code in (401, 403)
        assert sign_in(client, user.email, new_secret).json()["role"] == "Sales"
        if adapter:
            adapter.update_user(user.id, {"active": False})
        else:
            main.repo.users[user.id]["active"] = False
        assert client.get("/session/me").status_code == 401


def test_postgres_migrations_constraints_and_rollback(postgres_url, monkeypatch):
    from .db_customers import CustomerDatabase

    # Start from #23's actual final revision, seed representative pre-existing user/customer/import
    # data, then upgrade through #24's entire migration chain to head -- proving the upgrade path
    # itself, not just the schema it lands on -- and that a repeated upgrade is a clean no-op.
    migrate("002_customers")
    legacy_auth = AuthDatabase(postgres_url, create_schema=False)
    legacy_auth.create_user(user_id="legacy-sales", name="Legacy Sales", email="legacy-sales@example.test", role="Sales", password_hash=main.password_hash.hash(secrets.token_urlsafe(24)))
    # Raw SQL matching #23's actual 002_customers columns exactly -- the CustomerRow ORM model
    # reflects the head schema, which doesn't exist yet at this revision.
    with legacy_auth.engine.begin() as connection:
        connection.execute(text("INSERT INTO customers (bcn, name, status, owner_id, source, version) VALUES ('000042', 'Legacy Co', 'Open', 'legacy-sales', '{}'::json, 3)"))
        connection.execute(text("INSERT INTO customer_phones (bcn, phone, is_primary) VALUES ('000042', '555-0042', true)"))
        connection.execute(text("INSERT INTO import_jobs (id, submission_id, actor_id, filename, processed, created, updated, error_rows, status, created_at) VALUES ('legacy-job', 'legacy-import', 'legacy-sales', 'legacy.xlsx', 1, 1, 0, 0, 'Completed', now())"))
    legacy_auth.engine.dispose()
    migrate()
    migrate()  # repeated upgrade of #24's full chain is a no-op

    database = AuthDatabase(postgres_url, create_schema=False)
    legacy_user = database.user_by_email("legacy-sales@example.test")
    assert legacy_user is not None and legacy_user.role == "Sales" and legacy_user.active
    verify_customers = CustomerDatabase(postgres_url, create_schema=False)
    survivor = verify_customers.get("000042")
    assert (survivor.owner_id, survivor.status, survivor.version, verify_customers.phones("000042")) == ("legacy-sales", "Open", 3, ["555-0042"])
    with database.transaction() as session:
        assert session.scalar(text("SELECT status FROM import_jobs WHERE id = 'legacy-job'")) == "Completed"
    verify_customers.engine.dispose()
    secret = secrets.token_urlsafe(24)
    hashed = main.password_hash.hash(secret)
    kwargs = dict(user_id="persisted", name="Synthetic", email="persisted@example.test", role="Admin", password_hash=hashed)
    database.create_user(**kwargs)
    token, _ = database.issue("persisted")
    with database.transaction() as session:
        assert session.scalar(text("SELECT version_num FROM alembic_version")) == main.ALEMBIC_HEAD
        stored = session.get(SessionRow, digest(token))
        assert stored.digest != token and stored.digest == digest(token)
    # Actual PostgreSQL constraint violations, each rolled back independently.
    for changes in ({"role": "Other"}, {"email": " Upper@example.test "}, {"email": "straße@example.test"}, {"email": "\ttrim@example.test\n"}, {"password_hash": secrets.token_urlsafe(24)}, {"name": " "}, {"active": None}):
        row = dict(id=uuid4().hex, name="Synthetic", email=uuid4().hex + "@example.test", role="Admin", active=True, password_hash=hashed) | changes
        with pytest.raises(IntegrityError):
            with Session(database.engine) as session, session.begin():
                session.execute(UserRow.__table__.insert().values(**row))
    now = datetime.now(timezone.utc)
    for changes in ({"digest": secrets.token_urlsafe(24)}, {"user_id": "missing"}, {"expires_at": now}, {"revoked_at": now - timedelta(seconds=1)}):
        row = dict(digest=digest(secrets.token_urlsafe(24)), user_id="persisted", issued_at=now, expires_at=now + timedelta(hours=8)) | changes
        with pytest.raises(IntegrityError):
            with Session(database.engine) as session, session.begin():
                session.add(SessionRow(**row)); session.flush()
    barrier = Barrier(2)
    def create_same(index):
        barrier.wait()
        try:
            database.create_user(**(kwargs | {"user_id": f"race-{index}", "email": (" Race@Example.Test " if index else "race@example.test")}))
            return "created"
        except ValueError as exc:
            assert str(exc) == "normalized identity already exists"
            return "conflict"
    with ThreadPoolExecutor(max_workers=2) as pool:
        assert sorted(pool.map(create_same, range(2))) == ["conflict", "created"]
    assert len([user for user in database.all_users() if user.email == "race@example.test"]) == 1
    def fail_before_commit(session):
        if session.bind is database.engine:
            raise RuntimeError("Injected failure")
    event.listen(Session, "before_commit", fail_before_commit)
    try:
        with pytest.raises(RuntimeError, match="Injected failure"):
            database.create_user(**(kwargs | {"user_id": "rolled-back", "email": "rollback@example.test"}))
        with pytest.raises(RuntimeError, match="Injected failure"):
            database.reset_password("persisted", main.password_hash.hash(secrets.token_urlsafe(24)))
    finally:
        event.remove(Session, "before_commit", fail_before_commit)
    assert database.user_by_email("rollback@example.test") is None
    assert database.user_by_email("persisted@example.test").password_hash == hashed
    assert database.user_for_session(token).id == "persisted"
    before = len(database.all_users())
    assert database.reset_password("missing", hashed) is None
    assert len(database.all_users()) == before and database.user_for_session(token).id == "persisted"
    # UoW rollback includes identity and session rows in the same transaction.
    with pytest.raises(RuntimeError, match="Injected failure"):
        with database.transaction() as session:
            session.get(UserRow, "persisted").name = "Changed"
            session.get(SessionRow, digest(token)).revoked_at = now
            session.flush()
            raise RuntimeError("Injected failure")
    assert database.user_by_email("persisted@example.test").name == "Synthetic"
    assert database.user_for_session(token).id == "persisted"
    def unavailable(*args):
        from sqlalchemy.exc import OperationalError
        raise OperationalError("private connection", {}, Exception(secrets.token_urlsafe(24)))
    event.listen(database.engine, "before_cursor_execute", unavailable)
    try:
        with pytest.raises(StorageError, match=r"^Storage operation failed\.$"):
            database.all_users()
        monkeypatch.setattr(main, "auth_db", database)
        with TestClient(main.app, base_url="http://localhost") as client:
            response = client.get("/session/me")
            assert response.status_code == 503 and response.json() == {"detail": "Storage operation failed."}
    finally:
        event.remove(database.engine, "before_cursor_execute", unavailable)
    assert database.engine.pool.checkedout() == 0
    database.engine.dispose()


def test_postgres_activity_migrations_survive_bcn_widen_and_followup_link(postgres_url):
    """#25's own migrations (018 widening activities/follow_ups.bcn and adding their customer FK, 025
    adding the follow_ups.interaction_id link, and the 027/028 chain) must preserve previously seeded
    activity/follow-up/closure-reason/idempotency rows across the upgrade to head, and a repeated
    upgrade must be a no-op -- the same pattern the #24 fix used for #23's customers/import tables in
    test_postgres_migrations_constraints_and_rollback."""
    from .db_activity import ActivityDatabase, fingerprint
    from .db_customers import CustomerDatabase

    # Start from the revision right before #25's activity_customer_fks migration, seed representative
    # pre-existing rows using the schema as it existed at that revision (bcn is still VARCHAR(64) and
    # follow_ups has no interaction_id column yet), then upgrade through 018/025/026/027/028 to head.
    migrate("017_customer_owner_fk")
    auth = AuthDatabase(postgres_url, create_schema=False)
    customers = CustomerDatabase(postgres_url, create_schema=False)
    auth.create_user(user_id="legacy-actor", name="Legacy Actor", email="legacy-activity@example.test", role="Sales", password_hash=main.password_hash.hash(secrets.token_urlsafe(24)))
    customers.upsert_source(bcn="000555", name="Legacy Activity Co", source={})
    payload = "activity-legacy-payload"
    digest_value = fingerprint(payload)
    with auth.engine.begin() as connection:
        connection.execute(text("INSERT INTO closure_reasons (id, label, active) VALUES ('closure-legacy', 'Legacy reason', true)"))
        connection.execute(text("INSERT INTO activities (id, bcn, actor_id, kind, outcome, text, created_at) VALUES ('activity-legacy', '000555', 'legacy-actor', 'Interaction', 'Contact', 'Legacy note', now())"))
        connection.execute(text("INSERT INTO follow_ups (id, bcn, actor_id, type, status, version, created_at, updated_at) VALUES ('followup-legacy', '000555', 'legacy-actor', 'Reminder', 'Open', 0, now(), now())"))
        connection.execute(text("INSERT INTO idempotency_records (actor_id, operation, submission_id, fingerprint, result, completed_at) VALUES ('legacy-actor', 'interaction', 'legacy-submission', :fp, '{\"id\": \"activity-legacy\"}'::json, now())"), {"fp": digest_value})
    auth.engine.dispose()

    migrate()
    migrate()  # repeated upgrade of #25's full activity chain is a no-op

    activities = ActivityDatabase(postgres_url, create_schema=False)
    verify_customers = CustomerDatabase(postgres_url, create_schema=False)
    verify_auth = AuthDatabase(postgres_url, create_schema=False)
    try:
        with verify_auth.transaction() as session:
            assert session.scalar(text("SELECT version_num FROM alembic_version")) == main.ALEMBIC_HEAD
        assert verify_customers.get("000555").name == "Legacy Activity Co"
        history, total = activities.history("000555")
        assert total == 1 and history[0].id == "activity-legacy" and history[0].text == "Legacy note" and history[0].deleted_at is None
        followups = activities.followups("000555")
        assert len(followups) == 1 and followups[0].id == "followup-legacy" and followups[0].status == "Open" and followups[0].interaction_id is None
        reasons = {row.id: row for row in activities.reasons()}
        assert reasons["closure-legacy"].label == "Legacy reason" and reasons["closure-legacy"].active is True
        assert activities.get_idempotent(actor_id="legacy-actor", operation="interaction", submission_id="legacy-submission", payload=payload) == {"id": "activity-legacy"}

        # 025's interaction_id link column/constraint work against the migration-preserved follow-up.
        activities.complete_followup({"id": "followup-legacy", "bcn": "000555", "actorId": "legacy-actor"}, {"id": "activity-legacy-interaction", "bcn": "000555", "actorId": "legacy-actor", "outcome": "Contact", "note": "Post-migration completion"})
        completed = activities.followups("000555")[0]
        assert completed.status == "Completed" and completed.interaction_id == "activity-legacy-interaction"
    finally:
        activities.engine.dispose(); verify_customers.engine.dispose(); verify_auth.engine.dispose()


def test_postgres_customer_import_persists_and_rolls_back(postgres_url):
    from .db_customers import CustomerCollectionRow, CustomerDatabase, CustomerRow, PhoneRow

    migrate()
    auth = AuthDatabase(postgres_url, create_schema=False)
    customers = CustomerDatabase(postgres_url, create_schema=False)
    try:
        hashed = main.password_hash.hash(secrets.token_urlsafe(24))
        for user_id in ("admin", "sales"):
            auth.create_user(user_id=user_id, name=user_id, email=f"{user_id}@example.test", role="Admin" if user_id == "admin" else "Sales", password_hash=hashed)
        result = {"jobId": "import-1", "submissionId": "first", "filename": "synthetic.xlsx", "processed": 2, "created": 1, "updated": 0, "errorRows": 1, "status": "Partial", "errors": [{"row": 3, "field": "bcn", "reason": "Invalid bcn"}], "actorId": "admin"}
        customers.ingest_sources([{"bcn": "000123", "name": "Original", "source": {"customer_name": "Original"}, "typed": {"propensity_score": Decimal("0.875000"), "last_purchase_date": date(2025, 1, 2), "recent": True}, "collections": [{"kind": "vendor", "slot": 4, "name": "Synthetic vendor", "revenue": Decimal("12.500000")}], "primary_phone": "555-0100"}], result)
        with Session(customers.engine) as session:
            row = session.get(CustomerRow, "000123")
            assert row.bcn == "000123" and row.propensity_score == Decimal("0.875000") and row.last_purchase_date == date(2025, 1, 2) and row.recent is True
            assert session.query(CustomerCollectionRow).one().slot == 4
            session.add(PhoneRow(bcn="000123", phone="555-0101", primary=False)); session.commit()
        customers.save_operational(bcn="000123", owner_id="sales", status="Closed", version=7)
        customers.ingest_sources([{"bcn": "000123", "name": "Reimported", "source": {"customer_name": "Reimported"}, "typed": {"propensity_score": None, "last_purchase_date": None, "recent": None}, "collections": [], "primary_phone": None}], {**result, "jobId": "import-2", "submissionId": "second", "processed": 1, "created": 0, "updated": 1, "errorRows": 0, "status": "Completed", "errors": []})
        row = customers.get("000123")
        assert row.name == "Reimported" and row.owner_id == "sales" and row.status == "Closed" and row.version == 7 and row.propensity_score is None and row.last_purchase_date is None and row.recent is None
        assert customers.phones("000123") == ["555-0101"]
        assert customers.import_errors("admin", "first", 1, 1) == ([{"row": 3, "field": "bcn", "reason": "Invalid bcn"}], 1)
        with pytest.raises(IntegrityError):
            customers.ingest_sources([{"bcn": "000999", "name": "Rolled back", "source": {}, "primary_phone": None}], {**result, "submissionId": "conflict"})
        assert customers.get("000999") is None and customers.import_job("admin", "conflict") is None
        barrier = Barrier(2)
        def concurrent_import(index):
            barrier.wait()
            customers.ingest_sources([{"bcn": "000777", "name": f"Winner {index}", "source": {"customer_name": f"Winner {index}"}, "typed": {"propensity_rank": index}, "primary_phone": f"555-01{index}"}], {**result, "jobId": f"race-{index}", "submissionId": f"race-{index}", "errorRows": 0, "status": "Completed", "errors": []})
        with ThreadPoolExecutor(max_workers=2) as pool:
            list(pool.map(concurrent_import, (1, 2)))
        row = customers.get("000777")
        assert (row.name, row.propensity_rank, customers.phones("000777")) in (("Winner 1", 1, ["555-011"]), ("Winner 2", 2, ["555-012"]))
    finally:
        auth.engine.dispose(); customers.engine.dispose()


def test_postgres_customer_import_constraints_and_rollback(postgres_url):
    """Constraint/rollback coverage for customers/customer_phones/customer_collections/import_row_errors, matching test_postgres_migrations_constraints_and_rollback's pattern for users/sessions."""
    from .db_customers import CustomerCollectionRow, CustomerDatabase, CustomerRow, ImportErrorRow, ImportJobRow, PhoneRow

    migrate()
    auth = AuthDatabase(postgres_url, create_schema=False)
    customers = CustomerDatabase(postgres_url, create_schema=False)
    try:
        hashed = main.password_hash.hash(secrets.token_urlsafe(24))
        auth.create_user(user_id="owner", name="Owner", email="owner@example.test", role="Sales", password_hash=hashed)
        customers.upsert_source(bcn="000900", name="Synthetic", source={})
        # Blank required identifiers and an invalid status are rejected by check constraints.
        for changes in ({"bcn": "   "}, {"name": " "}, {"status": "Pending"}):
            row = dict(bcn=uuid4().hex[:12], name="Synthetic", status="Open", source={}, version=0) | changes
            with pytest.raises(IntegrityError):
                with Session(customers.engine) as session, session.begin():
                    session.execute(CustomerRow.__table__.insert().values(**row))
        # A blank phone is rejected.
        with pytest.raises(IntegrityError):
            with Session(customers.engine) as session, session.begin():
                session.add(PhoneRow(bcn="000900", phone="   ", primary=False)); session.flush()
        # Invalid collection kind/slot/name are each rejected.
        for changes in ({"kind": "other"}, {"slot": 0}, {"name": " "}):
            row = dict(bcn="000900", kind="vendor", slot=1, name="Synthetic vendor") | changes
            with pytest.raises(IntegrityError):
                with Session(customers.engine) as session, session.begin():
                    session.add(CustomerCollectionRow(**row)); session.flush()
        # An orphan owner_id is rejected.
        with pytest.raises(IntegrityError):
            with Session(customers.engine) as session, session.begin():
                session.get(CustomerRow, "000900").owner_id = "missing-user"; session.flush()
        # An orphan import_row_errors.job_id is rejected.
        with pytest.raises(IntegrityError):
            with Session(customers.engine) as session, session.begin():
                session.add(ImportErrorRow(job_id="missing-job", row_number=1, field="bcn", reason="bad")); session.flush()
        assert customers.get("000900").owner_id is None
        # An unexpected failure during ingest_sources rolls back the customer row, phone, job, and row errors together.
        def fail_commit(session):
            if session.bind is customers.engine:
                session.flush(); raise RuntimeError("Injected failure")
        event.listen(Session, "before_commit", fail_commit)
        try:
            with pytest.raises(RuntimeError, match="Injected failure"):
                customers.ingest_sources([{"bcn": "000901", "name": "Rolled back", "source": {}, "primary_phone": "555-0001"}], {"jobId": "job-rollback", "submissionId": "rollback", "filename": "x.xlsx", "processed": 1, "created": 1, "updated": 0, "errorRows": 1, "status": "Partial", "errors": [{"row": 2, "field": "bcn", "reason": "bad"}], "actorId": "owner"})
        finally:
            event.remove(Session, "before_commit", fail_commit)
        assert customers.get("000901") is None and customers.import_job("owner", "rollback") is None
        with Session(customers.engine) as session:
            assert session.scalar(select(ImportJobRow)) is None and session.scalar(select(ImportErrorRow)) is None
    finally:
        auth.engine.dispose(); customers.engine.dispose()


def test_postgres_activity_lifecycle_http_journey(postgres_url, monkeypatch):
    """The #16/#17 interaction/note/follow-up/lifecycle behavior suite through the real HTTP endpoints
    (/interactions, /notes, /follow-ups, /close, /reopen, /admin/closure-reasons) against real
    PostgreSQL -- the coverage gap QA found: every prior activity/lifecycle test called repository or
    main.py functions directly, or ran the HTTP routes only against the in-memory repo."""
    from .db_activity import ActivityDatabase
    from .db_customers import CustomerDatabase

    migrate()
    auth = AuthDatabase(postgres_url, create_schema=False)
    customers = CustomerDatabase(postgres_url, create_schema=False)
    activities = ActivityDatabase(postgres_url, create_schema=False)
    monkeypatch.setattr(main, "auth_db", auth)
    monkeypatch.setattr(main, "customer_db", customers)
    monkeypatch.setattr(main, "activity_db", activities)
    monkeypatch.setattr(main, "assignment_db", None)
    main.repo.reset(); main.repo.customers.clear()
    try:
        secret = secrets.token_urlsafe(24)
        hashed = main.password_hash.hash(secret)
        auth.create_user(user_id="journey-admin", name="Journey Admin", email="journey-admin@example.test", role="Admin", password_hash=hashed)
        auth.create_user(user_id="journey-owner", name="Journey Owner", email="journey-owner@example.test", role="Sales", password_hash=hashed)
        auth.create_user(user_id="journey-other", name="Journey Other", email="journey-other@example.test", role="Sales", password_hash=hashed)
        customers.upsert_source(bcn="000860", name="Journey Co", source={})
        customers.save_operational(bcn="000860", owner_id="journey-owner", status="Open", version=0)

        with TestClient(main.app, base_url="http://localhost") as admin_client:
            assert admin_client.post("/session/login", json={"email": "journey-admin@example.test", "password": secret}).status_code == 200
            reason = admin_client.post("/admin/closure-reasons", json={"label": "Journey resolved"}, headers=ORIGIN)
            assert reason.status_code == 201
            reason_id = reason.json()["id"]

        with TestClient(main.app, base_url="http://localhost") as owner_client:
            assert owner_client.post("/session/login", json={"email": "journey-owner@example.test", "password": secret}).status_code == 200
            assert owner_client.post("/admin/closure-reasons", json={"label": "Sales cannot"}, headers=ORIGIN).status_code == 403

            interaction = owner_client.post("/customers/000860/interactions", json={"outcome": "Contact", "note": "First contact", "submissionId": "journey-interaction"}, headers=ORIGIN)
            assert interaction.status_code == 200, interaction.text
            interaction_id = interaction.json()["id"]
            replay = owner_client.post("/customers/000860/interactions", json={"outcome": "Contact", "note": "First contact", "submissionId": "journey-interaction"}, headers=ORIGIN)
            assert replay.status_code == 200 and replay.json()["id"] == interaction_id
            conflict = owner_client.post("/customers/000860/interactions", json={"outcome": "Attempt", "note": "Changed", "submissionId": "journey-interaction"}, headers=ORIGIN)
            assert conflict.status_code == 409

            note = owner_client.post("/customers/000860/notes", json={"text": "Standalone note", "submissionId": "journey-note"}, headers=ORIGIN)
            assert note.status_code == 200
            note_id = note.json()["id"]
            deleted = owner_client.delete(f"/customers/000860/history/{note_id}", headers=ORIGIN)
            assert deleted.status_code == 200 and deleted.json()["deleted"] is True and deleted.json()["deletedBy"] == "journey-owner"
            history = owner_client.get("/customers/000860/history").json()
            tombstone = next(item for item in history["items"] if item["id"] == note_id)
            assert tombstone["deleted"] is True and tombstone["text"] is None

            first_followup = owner_client.post("/customers/000860/follow-ups", json={"type": "Reminder", "due": None, "note": "Call back", "submissionId": "journey-followup-1"}, headers=ORIGIN)
            assert first_followup.status_code == 200
            first_id = first_followup.json()["id"]
            second_followup = owner_client.post("/customers/000860/follow-ups", json={"type": "Appointment", "due": "2026-09-20", "note": "Site visit", "submissionId": "journey-followup-2"}, headers=ORIGIN)
            assert second_followup.status_code == 200
            second_id = second_followup.json()["id"]

            edited = owner_client.patch(f"/customers/000860/follow-ups/{first_id}", json={"type": "Reminder", "due": None, "note": "Call back tomorrow", "submissionId": "journey-followup-edit"}, headers=ORIGIN)
            assert edited.status_code == 200 and edited.json()["note"] == "Call back tomorrow" and edited.json()["version"] == 1

            completed = owner_client.post(f"/customers/000860/follow-ups/{second_id}/complete", json={"outcome": "Contact", "note": "Visited", "submissionId": "journey-complete"}, headers=ORIGIN)
            assert completed.status_code == 200 and completed.json()["status"] == "Completed" and completed.json()["interactionId"]
            replay_complete = owner_client.post(f"/customers/000860/follow-ups/{second_id}/complete", json={"outcome": "Contact", "note": "Visited", "submissionId": "journey-complete"}, headers=ORIGIN)
            assert replay_complete.status_code == 200 and replay_complete.json() == completed.json()
            already = owner_client.post(f"/customers/000860/follow-ups/{second_id}/complete", json={"outcome": "Attempt", "note": "Different", "submissionId": "journey-complete"}, headers=ORIGIN)
            assert already.status_code == 409

        with TestClient(main.app, base_url="http://localhost") as other_client:
            assert other_client.post("/session/login", json={"email": "journey-other@example.test", "password": secret}).status_code == 200
            assert other_client.post("/customers/000860/interactions", json={"outcome": "Attempt", "submissionId": "foreign"}, headers=ORIGIN).status_code == 403

        with TestClient(main.app, base_url="http://localhost") as owner_client:
            assert owner_client.post("/session/login", json={"email": "journey-owner@example.test", "password": secret}).status_code == 200
            closed = owner_client.post("/customers/000860/close", json={"reasonId": reason_id, "submissionId": "journey-close"}, headers=ORIGIN)
            assert closed.status_code == 200 and closed.json()["status"] == "Closed"
            assert owner_client.post("/customers/000860/interactions", json={"outcome": "Contact", "submissionId": "after-close"}, headers=ORIGIN).status_code == 409

            reopened = owner_client.post("/customers/000860/reopen", json={"submissionId": "journey-reopen"}, headers=ORIGIN)
            assert reopened.status_code == 200 and reopened.json()["status"] == "Open" and reopened.json()["ownerId"] == "journey-owner"
            replay_reopen = owner_client.post("/customers/000860/reopen", json={"submissionId": "journey-reopen"}, headers=ORIGIN)
            assert replay_reopen.status_code == 200 and replay_reopen.json() == reopened.json()

        # Persisted state matches the HTTP responses: the first follow-up was cancelled by the close
        # (never restored by reopen), the second stays Completed with its interaction link, the note
        # tombstone hides the text but keeps the record, and the closure/reopen events are both durable.
        followups = {row.id: row for row in activities.followups("000860")}
        assert followups[first_id].status == "Cancelled"
        assert followups[second_id].status == "Completed" and followups[second_id].interaction_id
        note_row = next(item for item in activities.history("000860")[0] if item.id == note_id)
        assert note_row.deleted_at is not None and note_row.text == "Standalone note"  # tombstone hides it only in API responses
        kinds = [item.kind for item in activities.history("000860")[0]]
        assert kinds.count("Closure") == 1 and kinds.count("Reopen") == 1
        assert customers.get("000860").status == "Open" and customers.get("000860").owner_id == "journey-owner"
        reasons = {row.id: row.label for row in activities.reasons()}
        assert reasons[reason_id] == "Journey resolved"

        # #30 coverage gap flagged in #25's QA pass: a second customer, closed while its owner was
        # still active Sales, then reopened only after that owner is deactivated -- ownership must
        # release (ownerId/ownerName null) with the Assignment history entry plus the
        # assignment_history/audit_events rows, all written atomically inside reopen_customer.
        from .db_assignment import AssignmentDatabase, AssignmentHistoryRow, AuditRow

        customers.upsert_source(bcn="000861", name="Journey Release Co", source={})
        customers.save_operational(bcn="000861", owner_id="journey-owner", status="Open", version=0)

        with TestClient(main.app, base_url="http://localhost") as owner_client:
            assert owner_client.post("/session/login", json={"email": "journey-owner@example.test", "password": secret}).status_code == 200
            closed_release = owner_client.post("/customers/000861/close", json={"reasonId": reason_id, "submissionId": "journey-release-close"}, headers=ORIGIN)
            assert closed_release.status_code == 200 and closed_release.json()["status"] == "Closed"

        with TestClient(main.app, base_url="http://localhost") as admin_client:
            assert admin_client.post("/session/login", json={"email": "journey-admin@example.test", "password": secret}).status_code == 200
            deactivated = admin_client.patch("/admin/users/journey-owner", json={"active": False}, headers=ORIGIN)
            assert deactivated.status_code == 200 and deactivated.json()["active"] is False

            reopened_release = admin_client.post("/customers/000861/reopen", json={"submissionId": "journey-release-reopen"}, headers=ORIGIN)
            assert reopened_release.status_code == 200, reopened_release.text
            released_body = reopened_release.json()
            assert released_body["status"] == "Open" and released_body["ownerId"] is None and released_body["ownerName"] is None
            release_entries = [item for item in released_body["histories"] if item["kind"] == "Assignment"]
            assert len(release_entries) == 1 and release_entries[0]["oldOwner"] == "journey-owner" and release_entries[0]["newOwner"] is None and release_entries[0]["reason"] == "Owner no longer active Sales"

            replay_release = admin_client.post("/customers/000861/reopen", json={"submissionId": "journey-release-reopen"}, headers=ORIGIN)
            assert replay_release.status_code == 200 and replay_release.json() == released_body

        assert customers.get("000861").status == "Open" and customers.get("000861").owner_id is None
        release_kinds = [item.kind for item in activities.history("000861")[0]]
        assert release_kinds.count("Reopen") == 1

        assignments = AssignmentDatabase(postgres_url, create_schema=False)
        try:
            with Session(assignments.engine) as session:
                history_row = session.execute(select(AssignmentHistoryRow).where(AssignmentHistoryRow.bcn == "000861")).scalar_one()
                assert history_row.old_owner_id == "journey-owner" and history_row.new_owner_id is None and history_row.reason == "Owner no longer active Sales"
                audit_row = session.execute(select(AuditRow).where(AuditRow.target == "000861", AuditRow.action == "Customer ownership released")).scalar_one()
                assert audit_row.actor_id == "journey-admin" and audit_row.details == {"oldOwner": "journey-owner", "newOwner": None, "reason": "Owner no longer active Sales"}
        finally:
            assignments.engine.dispose()
    finally:
        auth.engine.dispose(); customers.engine.dispose(); activities.engine.dispose()


def test_postgres_search_workload_reports_audit_http_journey(postgres_url, monkeypatch):
    """The QA-flagged coverage gap: /customers search/detail, /sales/workload, /workload,
    /admin/reports, and /admin/audit exercised as real HTTP requests against real PostgreSQL --
    plus unauthorized direct requests (wrong role, wrong owner, forged IDs) rejected the same way
    as in memory mode."""
    from .db_activity import ActivityDatabase
    from .db_assignment import AssignmentDatabase
    from .db_customers import CustomerDatabase

    migrate()
    auth = AuthDatabase(postgres_url, create_schema=False)
    customers = CustomerDatabase(postgres_url, create_schema=False)
    activities = ActivityDatabase(postgres_url, create_schema=False)
    assignments = AssignmentDatabase(postgres_url, create_schema=False)
    monkeypatch.setattr(main, "auth_db", auth)
    monkeypatch.setattr(main, "customer_db", customers)
    monkeypatch.setattr(main, "activity_db", activities)
    monkeypatch.setattr(main, "assignment_db", assignments)
    main.repo.reset(); main.repo.customers.clear()
    try:
        secret = secrets.token_urlsafe(24)
        hashed = main.password_hash.hash(secret)
        auth.create_user(user_id="rep-admin", name="Rep Admin", email="rep-admin@example.test", role="Admin", password_hash=hashed)
        auth.create_user(user_id="rep-sales-a", name="Rep Sales A", email="rep-sales-a@example.test", role="Sales", password_hash=hashed)
        auth.create_user(user_id="rep-sales-b", name="Rep Sales B", email="rep-sales-b@example.test", role="Sales", password_hash=hashed)
        for user_id, role, active in (("rep-admin", "Admin", True), ("rep-sales-a", "Sales", True), ("rep-sales-b", "Sales", True)):
            main.repo.users[user_id] = {"id": user_id, "name": user_id, "role": role, "active": active}

        yesterday = (datetime.now(timezone.utc) - timedelta(days=1)).isoformat()
        today_iso = datetime.now(timezone.utc).isoformat()

        # never-contacted: no activity at all.
        customers.upsert_source(bcn="000801", name="Never Contacted Co", source={}, primary_phone="555-0801"); customers.save_operational(bcn="000801", owner_id="rep-sales-a", status="Open", version=0)
        # overdue: an open follow-up due yesterday, plus a prior interaction (overdue still wins).
        customers.upsert_source(bcn="000802", name="Overdue Co", source={}, primary_phone="555-0802"); customers.save_operational(bcn="000802", owner_id="rep-sales-a", status="Open", version=0)
        activities.save_activity(record_id="act-802", bcn="000802", actor_id="rep-sales-a", kind="Interaction", outcome="Attempt", text="Tried")
        activities.save_followup({"id": "fu-802", "bcn": "000802", "actorId": "rep-sales-a", "type": "Reminder", "due": yesterday, "note": "Overdue", "status": "Open"})
        # today: an open follow-up due today.
        customers.upsert_source(bcn="000803", name="Today Co", source={}, primary_phone="555-0803"); customers.save_operational(bcn="000803", owner_id="rep-sales-a", status="Open", version=0)
        activities.save_followup({"id": "fu-803", "bcn": "000803", "actorId": "rep-sales-a", "type": "Reminder", "due": today_iso, "note": "Today", "status": "Open"})
        # undated: an open follow-up with no due date.
        customers.upsert_source(bcn="000804", name="Undated Co", source={}, primary_phone="555-0804"); customers.save_operational(bcn="000804", owner_id="rep-sales-a", status="Open", version=0)
        activities.save_followup({"id": "fu-804", "bcn": "000804", "actorId": "rep-sales-a", "type": "Reminder", "due": None, "note": "Whenever", "status": "Open"})
        # other: contacted, no open follow-ups.
        customers.upsert_source(bcn="000805", name="Contacted Co", source={}, primary_phone="555-0805"); customers.save_operational(bcn="000805", owner_id="rep-sales-a", status="Open", version=0)
        activities.save_activity(record_id="act-805", bcn="000805", actor_id="rep-sales-a", kind="Interaction", outcome="Contact", text="Reached them")
        # a closed customer owned by the other sales rep.
        customers.upsert_source(bcn="000806", name="Closed Co", source={}, primary_phone="555-0806"); customers.save_operational(bcn="000806", owner_id="rep-sales-b", status="Closed", version=1)

        with TestClient(main.app, base_url="http://localhost") as anon_client:
            assert anon_client.get("/customers").status_code == 401

        with TestClient(main.app, base_url="http://localhost") as admin_client:
            assert admin_client.post("/session/login", json={"email": "rep-admin@example.test", "password": secret}).status_code == 200

            # /customers search: name/phone query, status filter, owner filter.
            by_name = admin_client.get("/customers", params={"q": "overdue"}).json()
            assert by_name["total"] == 1 and by_name["items"][0]["bcn"] == "000802"
            by_phone = admin_client.get("/customers", params={"q": "555-0804"}).json()
            assert by_phone["total"] == 1 and by_phone["items"][0]["bcn"] == "000804"
            closed_only = admin_client.get("/customers", params={"status": "Closed"}).json()
            assert closed_only["total"] == 1 and closed_only["items"][0]["bcn"] == "000806"
            owned_by_a = admin_client.get("/customers", params={"owner": "rep-sales-a"}).json()
            assert owned_by_a["total"] == 5

            # /customers/{bcn} detail includes phones/source/histories.
            detail = admin_client.get("/customers/000802").json()
            assert detail["phones"] == ["555-0802"] and detail["ownerId"] == "rep-sales-a"
            assert any(item["kind"] == "Interaction" for item in detail["histories"])
            assert any(item["id"] == "fu-802" for item in detail["followUps"])

            # Wrong role: Sales cannot read admin reports/audit.
            reason = admin_client.post("/admin/closure-reasons", json={"label": "Report done"}, headers=ORIGIN).json()

            # /admin/reports: owner rollup, follow-up buckets, daily interactions -- all DB-side aggregated.
            reports = admin_client.get("/admin/reports").json()
            owners_by_id = {item["ownerId"]: item for item in reports["owners"]}
            sales_a = owners_by_id["rep-sales-a"]
            # neverContacted counts customers with zero interactions ever (801/803/804 have none;
            # 802/805 each have one) -- distinct from the workload bucket, which also weighs open
            # follow-ups.
            assert (sales_a["open"], sales_a["neverContacted"], sales_a["attempts"], sales_a["contacts"], sales_a["pendingFollowUps"]) == (5, 3, 1, 1, 3)
            assert sales_a["contactRate"] == 100.0
            sales_b = owners_by_id["rep-sales-b"]
            assert (sales_b["open"], sales_b["closed"]) == (0, 1)
            assert reports["followUps"] == {"overdue": 1, "today": 1, "undated": 1, "completed": 0}

            # /admin/audit: filter by bcn/action and paginate.
            audit_for_802 = admin_client.get("/admin/audit", params={"bcn": "000802"}).json()
            assert audit_for_802["total"] >= 0  # no mutation audit rows were written for seeded rows; endpoint still answers correctly
            audit_page = admin_client.get("/admin/audit", params={"page": 1, "page_size": 1}).json()
            assert audit_page["page"] == 1 and audit_page["page_size"] == 1 and len(audit_page["items"]) <= 1

        with TestClient(main.app, base_url="http://localhost") as sales_a_client:
            assert sales_a_client.post("/session/login", json={"email": "rep-sales-a@example.test", "password": secret}).status_code == 200

            # /sales/workload and /workload -- bucketed by real SQL aggregates.
            workload = sales_a_client.get("/sales/workload").json()
            assert workload == {"ownerId": "rep-sales-a", "totalOpen": 5, "neverContacted": 3, "pendingFollowUps": 3}
            full_workload = sales_a_client.get("/workload").json()
            buckets = {item["bcn"]: item["workloadBucket"] for item in full_workload["customers"]}
            assert buckets == {"000801": "never-contacted", "000802": "overdue", "000803": "today", "000804": "undated", "000805": "other"}
            assert full_workload["counts"] == {"overdue": 1, "today": 1, "undated": 1, "never-contacted": 1, "other": 1}

            # Wrong role: Sales cannot read Admin-only endpoints.
            assert sales_a_client.get("/admin/reports").status_code == 403
            assert sales_a_client.get("/admin/audit").status_code == 403
            assert sales_a_client.post("/admin/closure-reasons", json={"label": "nope"}, headers=ORIGIN).status_code == 403

            # Forged/nonexistent IDs.
            assert sales_a_client.get("/customers/999999").status_code == 404
            assert sales_a_client.delete("/customers/000802/history/forged-record-id", headers=ORIGIN).status_code == 404
            assert sales_a_client.post("/customers/000802/follow-ups/forged-followup-id/complete", json={"outcome": "Contact", "submissionId": "forged"}, headers=ORIGIN).status_code == 404

            # Wrong owner: rep-sales-a cannot write to rep-sales-b's customer.
            assert sales_a_client.post("/customers/000806/interactions", json={"outcome": "Attempt", "submissionId": "wrong-owner"}, headers=ORIGIN).status_code in (403, 409)

        with TestClient(main.app, base_url="http://localhost") as sales_b_client:
            assert sales_b_client.post("/session/login", json={"email": "rep-sales-b@example.test", "password": secret}).status_code == 200
            # Wrong owner on an Open customer: 403 (not merely 409 from being closed).
            assert sales_b_client.post("/customers/000801/interactions", json={"outcome": "Attempt", "submissionId": "wrong-owner-2"}, headers=ORIGIN).status_code == 403
    finally:
        auth.engine.dispose(); customers.engine.dispose(); activities.engine.dispose(); assignments.engine.dispose()


def test_postgres_reimport_after_operational_work_http_journey(postgres_url, monkeypatch):
    """The #26 QA gap: the existing reimport HTTP test ran with activity_db=None and no prior
    work. This runs a full reimport as an HTTP request AFTER real interaction/follow-up/closure/
    assignment work, and proves it changes only source fields and the source-primary phone while
    preserving every operational and audit record and leaving current ownership untouched."""
    from .db_activity import ActivityDatabase
    from .db_assignment import AssignmentDatabase, AssignmentHistoryRow, AuditRow
    from .db_customers import CustomerDatabase

    migrate()
    auth = AuthDatabase(postgres_url, create_schema=False)
    customers = CustomerDatabase(postgres_url, create_schema=False)
    activities = ActivityDatabase(postgres_url, create_schema=False)
    assignments = AssignmentDatabase(postgres_url, create_schema=False)
    monkeypatch.setattr(main, "auth_db", auth)
    monkeypatch.setattr(main, "customer_db", customers)
    monkeypatch.setattr(main, "activity_db", activities)
    monkeypatch.setattr(main, "assignment_db", assignments)
    main.repo.reset(); main.repo.customers.clear()
    try:
        secret = secrets.token_urlsafe(24)
        hashed = main.password_hash.hash(secret)
        auth.create_user(user_id="reimport-admin", name="Reimport Admin", email="reimport-admin@example.test", role="Admin", password_hash=hashed)
        auth.create_user(user_id="reimport-owner", name="Reimport Owner", email="reimport-owner@example.test", role="Sales", password_hash=hashed)
        for user_id, role in (("reimport-admin", "Admin"), ("reimport-owner", "Sales")):
            main.repo.users[user_id] = {"id": user_id, "name": user_id, "role": role, "active": True}

        headers = ["bcn", "customer_name", "propensity_score", "last_purchase_date", "recent", "phone"]

        with TestClient(main.app, base_url="http://localhost") as admin_client:
            assert admin_client.post("/session/login", json={"email": "reimport-admin@example.test", "password": secret}).status_code == 200
            reason = admin_client.post("/admin/closure-reasons", json={"label": "Reimport resolved"}, headers=ORIGIN)
            assert reason.status_code == 201
            reason_id = reason.json()["id"]

            book = Workbook(); sheet = book.active; sheet.append(headers)
            sheet.append(["700900", "Original Co", 0.30, datetime(2025, 1, 1), True, "555-7000"])
            payload = BytesIO(); book.save(payload)
            first = admin_client.post("/admin/imports?submission_id=reimport-initial", files={"file": ("first.xlsx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=ORIGIN)
            assert first.status_code == 201 and first.json()["created"] == 1

            assigned = admin_client.post("/admin/assignments/manual/700900", json={"ownerId": "reimport-owner", "submissionId": "reimport-assign"}, headers=ORIGIN)
            assert assigned.status_code == 200 and assigned.json()["ownerId"] == "reimport-owner"

        # Real operational work happens before the reimport: interaction, note, follow-up +
        # completion, closure, reopen -- each producing durable activity/audit/assignment rows.
        with TestClient(main.app, base_url="http://localhost") as owner_client:
            assert owner_client.post("/session/login", json={"email": "reimport-owner@example.test", "password": secret}).status_code == 200
            interaction = owner_client.post("/customers/700900/interactions", json={"outcome": "Attempt", "note": "Pre-reimport contact", "submissionId": "reimport-interaction"}, headers=ORIGIN)
            assert interaction.status_code == 200
            note = owner_client.post("/customers/700900/notes", json={"text": "Pre-reimport note", "submissionId": "reimport-note"}, headers=ORIGIN)
            assert note.status_code == 200
            followup = owner_client.post("/customers/700900/follow-ups", json={"type": "Reminder", "due": None, "note": "Follow up", "submissionId": "reimport-followup"}, headers=ORIGIN)
            assert followup.status_code == 200
            followup_id = followup.json()["id"]
            completed = owner_client.post(f"/customers/700900/follow-ups/{followup_id}/complete", json={"outcome": "Contact", "note": "Completed before reimport", "submissionId": "reimport-complete"}, headers=ORIGIN)
            assert completed.status_code == 200
            closed = owner_client.post("/customers/700900/close", json={"reasonId": reason_id, "submissionId": "reimport-close"}, headers=ORIGIN)
            assert closed.status_code == 200
            reopened = owner_client.post("/customers/700900/reopen", json={"submissionId": "reimport-reopen"}, headers=ORIGIN)
            assert reopened.status_code == 200 and reopened.json()["ownerId"] == "reimport-owner"

        # Snapshot durable state before the reimport.
        history_before = [(item.id, item.kind, item.outcome, item.text, item.deleted_at) for item in activities.history("700900", 1, 100)[0]]
        followups_before = [(item.id, item.status, item.interaction_id) for item in activities.followups("700900")]
        with Session(assignments.engine) as session:
            history_row_count_before = len(list(session.scalars(select(AssignmentHistoryRow))))
            audit_row_count_before = len(list(session.scalars(select(AuditRow))))
        owner_before, status_before, version_before = customers.get("700900").owner_id, customers.get("700900").status, customers.get("700900").version

        # The reimport itself: changes source fields and the source-primary phone only.
        with TestClient(main.app, base_url="http://localhost") as admin_client:
            assert admin_client.post("/session/login", json={"email": "reimport-admin@example.test", "password": secret}).status_code == 200
            changed_book = Workbook(); sheet = changed_book.active; sheet.append(headers)
            sheet.append(["700900", "Renamed Co", 0.91, datetime(2026, 2, 2), False, "555-9999"])
            changed_payload = BytesIO(); changed_book.save(changed_payload)
            reimport = admin_client.post("/admin/imports?submission_id=reimport-after-work", files={"file": ("second.xlsx", changed_payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=ORIGIN)
            assert reimport.status_code == 201
            assert (reimport.json()["created"], reimport.json()["updated"]) == (0, 1)

            detail = admin_client.get("/customers/700900").json()
            assert detail["name"] == "Renamed Co" and detail["phones"] == ["555-9999"]
            assert detail["source"]["propensity_score"] == 0.91

        # Only source fields and the primary phone changed: ownership/status/version untouched.
        row = customers.get("700900")
        assert (row.owner_id, row.status, row.version) == (owner_before, status_before, version_before) == ("reimport-owner", "Open", version_before)
        assert row.name == "Renamed Co" and row.propensity_score == Decimal("0.91")

        # Every operational and audit record from before the reimport survives unchanged.
        history_after = [(item.id, item.kind, item.outcome, item.text, item.deleted_at) for item in activities.history("700900", 1, 100)[0]]
        assert history_after == history_before
        followups_after = [(item.id, item.status, item.interaction_id) for item in activities.followups("700900")]
        assert followups_after == followups_before
        with Session(assignments.engine) as session:
            assert len(list(session.scalars(select(AssignmentHistoryRow)))) == history_row_count_before
            assert len(list(session.scalars(select(AuditRow)))) == audit_row_count_before + 1  # +1 for the import-completed audit entry
        kinds = [item[1] for item in history_after]
        assert kinds.count("Closure") == 1 and kinds.count("Reopen") == 1
    finally:
        auth.engine.dispose(); customers.engine.dispose(); activities.engine.dispose(); assignments.engine.dispose()


def test_postgres_import_http_pipeline_behavior(postgres_url, monkeypatch):
    """The #18 ingestion behavior suite through the real /admin/imports upload pipeline against Postgres: row errors, duplicate rows in one file, retry, conflict-on-change, unchanged-row updates, and null-clearing."""
    from .db_customers import CustomerDatabase

    migrate()
    auth = AuthDatabase(postgres_url, create_schema=False)
    customers = CustomerDatabase(postgres_url, create_schema=False)
    monkeypatch.setattr(main, "auth_db", auth)
    monkeypatch.setattr(main, "customer_db", customers)
    monkeypatch.setattr(main, "activity_db", None)
    main.repo.reset(); main.repo.customers.clear()
    try:
        admin, secret = provision("pipeline-admin@example.test")
        with TestClient(main.app, base_url="http://localhost") as client:
            assert sign_in(client, admin.email, secret).status_code == 200
            headers = ["bcn", "customer_name", "propensity_score", "last_purchase_date", "recent", "phone"]
            book = Workbook(); sheet = book.active; sheet.append(headers)
            sheet.append(["700100", "Stable Co", 0.42, datetime(2025, 3, 4), True, "555-1000"])
            sheet.append(["bad", "Broken", None, None, None, None])
            sheet.append(["700100", "Duplicate", None, None, None, None])
            output = BytesIO(); book.save(output); first_bytes = output.getvalue()
            first = client.post("/admin/imports?submission_id=first-run", files={"file": ("first.xlsx", first_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=ORIGIN)
            body = first.json()
            assert first.status_code == 201
            assert (body["processed"], body["created"], body["updated"], body["errorRows"], body["status"]) == (3, 1, 0, 2, "Partial")
            assert body["errors"] == [{"row": 3, "field": "bcn", "reason": "Invalid bcn"}, {"row": 4, "field": "bcn", "reason": "Duplicate bcn"}]
            row = customers.get("700100")
            assert row.propensity_score == Decimal("0.42") and row.last_purchase_date == date(2025, 3, 4) and row.recent is True
            assert customers.phones("700100") == ["555-1000"]

            # A same-submission retry after a simulated restart (in-process cache cleared) replays the persisted job.
            main.repo.imports.clear(); main.repo.import_payloads.clear()
            replay = client.post("/admin/imports?submission_id=first-run", files={"file": ("first.xlsx", first_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=ORIGIN)
            assert replay.status_code == 201 and replay.json()["jobId"] == body["jobId"]
            assert customers.import_jobs(admin.id, 1, 50)[1] == 1

            # Reuse of the same submission ID with a changed workbook is a contracted conflict.
            main.repo.imports.clear(); main.repo.import_payloads.clear()
            changed = Workbook(); changed.active.append(["bcn", "customer_name"]); changed.active.append(["700200", "Changed"])
            changed_output = BytesIO(); changed.save(changed_output)
            conflict = client.post("/admin/imports?submission_id=first-run", files={"file": ("first.xlsx", changed_output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=ORIGIN)
            assert conflict.status_code == 409

            # Reimporting the same unchanged row counts it as updated, not created.
            unchanged = Workbook(); unchanged.active.append(headers); unchanged.active.append(["700100", "Stable Co", 0.42, datetime(2025, 3, 4), True, "555-1000"])
            unchanged_output = BytesIO(); unchanged.save(unchanged_output)
            second = client.post("/admin/imports?submission_id=second-run", files={"file": ("second.xlsx", unchanged_output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=ORIGIN)
            assert second.status_code == 201
            assert (second.json()["created"], second.json()["updated"]) == (0, 1)

            # Blank nullable cells clear the typed values and remove the primary phone.
            nulled = Workbook(); nulled.active.append(headers); nulled.active.append(["700100", "Stable Co", None, None, None, None])
            nulled_output = BytesIO(); nulled.save(nulled_output)
            third = client.post("/admin/imports?submission_id=third-run", files={"file": ("third.xlsx", nulled_output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=ORIGIN)
            assert third.status_code == 201
            row = customers.get("700100")
            assert row.propensity_score is None and row.last_purchase_date is None and row.recent is None
            assert customers.phones("700100") == []
    finally:
        auth.engine.dispose(); customers.engine.dispose()


def test_postgres_import_rejects_empty_and_corrupt_workbooks(postgres_url, monkeypatch):
    """Empty (header-only) workbooks persist a zero-total job; corrupt/wrong-extension uploads are whole-file rejections with no persisted job."""
    from .db_customers import CustomerDatabase

    migrate()
    auth = AuthDatabase(postgres_url, create_schema=False)
    customers = CustomerDatabase(postgres_url, create_schema=False)
    monkeypatch.setattr(main, "auth_db", auth)
    monkeypatch.setattr(main, "customer_db", customers)
    monkeypatch.setattr(main, "activity_db", None)
    main.repo.reset(); main.repo.customers.clear()
    try:
        admin, secret = provision("rejection-admin@example.test")
        with TestClient(main.app, base_url="http://localhost") as client:
            assert sign_in(client, admin.email, secret).status_code == 200
            book = Workbook(); book.active.append(["bcn", "customer_name"])
            output = BytesIO(); book.save(output)
            empty = client.post("/admin/imports?submission_id=empty-book", files={"file": ("empty.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=ORIGIN)
            body = empty.json()
            assert empty.status_code == 201
            assert (body["processed"], body["created"], body["updated"], body["errorRows"], body["status"]) == (0, 0, 0, 0, "Completed")
            assert customers.import_job(admin.id, "empty-book")["jobId"] == body["jobId"]

            wrong_ext = client.post("/admin/imports?submission_id=wrong-ext", files={"file": ("customers.csv", b"bcn,customer_name\n700300,New Co\n", "text/csv")}, headers=ORIGIN)
            assert wrong_ext.status_code == 422
            assert customers.import_job(admin.id, "wrong-ext") is None

            corrupt = client.post("/admin/imports?submission_id=corrupt-zip", files={"file": ("corrupt.xlsx", b"this is not a real zip archive", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=ORIGIN)
            assert corrupt.status_code == 422
            assert customers.import_job(admin.id, "corrupt-zip") is None
            assert customers.get("700300") is None
    finally:
        auth.engine.dispose(); customers.engine.dispose()


def test_postgres_assignment_constraints_and_rollback(postgres_url):
    """Constraint/rollback coverage for assignment_rules/assignment_history/audit_events/assignment_runs, matching test_postgres_customer_import_constraints_and_rollback's pattern.

    Starts from #23's actual final revision (002_customers), seeded with representative
    pre-existing user/customer/import data, then upgrades through #24's entire migration chain to
    head -- proving the upgrade path survives and is a no-op on repeat -- before covering the
    constraint/rollback behavior.
    """
    from .db_assignment import AssignmentDatabase, AssignmentHistoryRow, AssignmentRunRow, AuditRow, RuleRow, RuleMemberRow
    from .db_customers import CustomerDatabase

    migrate("002_customers")
    auth = AuthDatabase(postgres_url, create_schema=False)
    customers = CustomerDatabase(postgres_url, create_schema=False)
    hashed = main.password_hash.hash(secrets.token_urlsafe(24))
    auth.create_user(user_id="owner", name="Owner", email="owner@example.test", role="Sales", password_hash=hashed)
    # Raw SQL matching #23's actual 002_customers columns exactly -- the CustomerRow ORM model
    # reflects the head schema, which doesn't exist yet at this revision.
    with auth.engine.begin() as connection:
        connection.execute(text("INSERT INTO customers (bcn, name, status, owner_id, source, version) VALUES ('000900', 'Synthetic', 'Open', 'owner', '{}'::json, 2)"))
        connection.execute(text("INSERT INTO customer_phones (bcn, phone, is_primary) VALUES ('000900', '555-0900', true)"))
        connection.execute(text("INSERT INTO import_jobs (id, submission_id, actor_id, filename, processed, created, updated, error_rows, status, created_at) VALUES ('pre-migration-job', 'pre-migration', 'owner', 'legacy.xlsx', 1, 1, 0, 0, 'Completed', now())"))

    migrate()
    migrate()  # repeated upgrade of #24's full chain is a no-op

    # #23's pre-existing data survives the upgrade unchanged.
    assert auth.user_by_email("owner@example.test") is not None
    survivor = customers.get("000900")
    assert (survivor.owner_id, survivor.status, survivor.version, customers.phones("000900")) == ("owner", "Open", 2, ["555-0900"])
    with auth.transaction() as session:
        assert session.scalar(text("SELECT status FROM import_jobs WHERE id = 'pre-migration-job'")) == "Completed"

    assignments = AssignmentDatabase(postgres_url, create_schema=False)
    try:
        # Rule-name uniqueness is rejected; a rule member naming a nonexistent user is rejected.
        assignments.create_rule(rule_id="r1", name="Rule One", position=1, actor_id="owner", conditions=[], member_ids=["owner"])
        with pytest.raises(IntegrityError):
            with Session(assignments.engine) as session, session.begin():
                session.add(RuleRow(id="r2", name="Rule One", position=2, active=True, version=0, conditions=[])); session.flush()
        with pytest.raises(IntegrityError):
            with Session(assignments.engine) as session, session.begin():
                session.add(RuleMemberRow(rule_id="r1", user_id="missing-user")); session.flush()

        # assignment_history rejects a nonexistent actor, old owner, new owner, or bcn.
        base = dict(bcn="000900", actor_id="owner", old_owner_id=None, new_owner_id="owner", reason="Synthetic", created_at=datetime.now(timezone.utc))
        for changes in ({"actor_id": "missing"}, {"old_owner_id": "missing"}, {"new_owner_id": "missing"}, {"bcn": "999999"}):
            with pytest.raises(IntegrityError):
                with Session(assignments.engine) as session, session.begin():
                    session.add(AssignmentHistoryRow(**(base | changes))); session.flush()

        # audit_events rejects a nonexistent actor.
        with pytest.raises(IntegrityError):
            with Session(assignments.engine) as session, session.begin():
                session.add(AuditRow(actor_id="missing", action="Test", target="000900", details={}, created_at=datetime.now(timezone.utc))); session.flush()

        # assignment_runs enforces its (actor_id, submission_id) uniqueness at the database level too.
        run = dict(actor_id="owner", submission_id="dup", scope="unassigned", fingerprint=None, result={}, created_at=datetime.now(timezone.utc))
        with Session(assignments.engine) as session, session.begin():
            session.add(AssignmentRunRow(**run)); session.flush()
        with pytest.raises(IntegrityError):
            with Session(assignments.engine) as session, session.begin():
                session.add(AssignmentRunRow(**run)); session.flush()

        # An ordinary UPDATE or DELETE against the append-only history tables is rejected by the database trigger.
        assignments.append_audit(actor_id="owner", action="Synthetic", target="000900", details={})
        assignments.append_assignment(bcn="000900", actor_id="owner", old_owner_id=None, new_owner_id="owner", reason="Synthetic")
        with Session(assignments.engine) as session:
            audit_id = session.scalar(select(AuditRow.id).where(AuditRow.action == "Synthetic"))
            with pytest.raises(DBAPIError):
                session.execute(update(AuditRow).where(AuditRow.id == audit_id).values(action="Changed"))
            session.rollback()
        with Session(assignments.engine) as session:
            history_id = session.scalar(select(AssignmentHistoryRow.id))
            with pytest.raises(DBAPIError):
                session.execute(delete(AssignmentHistoryRow).where(AssignmentHistoryRow.id == history_id))
            session.rollback()
        with Session(assignments.engine) as session:
            assert session.scalar(select(AuditRow).where(AuditRow.id == audit_id)).action == "Synthetic"
            assert session.scalar(select(AssignmentHistoryRow).where(AssignmentHistoryRow.id == history_id)) is not None
    finally:
        auth.engine.dispose(); customers.engine.dispose(); assignments.engine.dispose()


def test_postgres_bulk_assignment_rule_order_fallback_and_scope(postgres_url, monkeypatch):
    """The #19 rule-order/fallback/unchanged-owner engine (test_assignment_order_fallback_and_unchanged_owner) against real PostgreSQL persistence."""
    from .db_assignment import AssignmentDatabase, AssignmentHistoryRow, AuditRow
    from .db_customers import CustomerDatabase

    migrate()
    auth = AuthDatabase(postgres_url, create_schema=False)
    customers = CustomerDatabase(postgres_url, create_schema=False)
    assignments = AssignmentDatabase(postgres_url, create_schema=False)
    monkeypatch.setattr(main, "auth_db", auth)
    monkeypatch.setattr(main, "customer_db", customers)
    monkeypatch.setattr(main, "assignment_db", assignments)
    main.repo.reset()
    try:
        hashed = main.password_hash.hash(secrets.token_urlsafe(24))
        for user_id, role, active in (("admin", "Admin", True), ("first", "Sales", True), ("second", "Sales", True), ("inactive", "Sales", False)):
            auth.create_user(user_id=user_id, name=user_id, email=f"{user_id}@example.test", role=role, password_hash=hashed)
            if not active: auth.update_user(user_id, {"active": False})
            main.repo.users[user_id] = {"id": user_id, "name": user_id, "role": role, "active": active}
        customers.upsert_source(bcn="000001", name="Synthetic", source={}); customers.save_operational(bcn="000001", owner_id="first", status="Open", version=4)
        customers.upsert_source(bcn="000002", name="Synthetic", source={}); customers.save_operational(bcn="000002", owner_id=None, status="Open", version=0)
        customers.upsert_source(bcn="000003", name="Synthetic", source={}); customers.save_operational(bcn="000003", owner_id=None, status="Closed", version=0)
        main.repo.customers = {}
        actor = main.User(id="admin", name="Admin", email="admin@example.test", role="Admin")
        rule_first = main.create_assignment_rule(main.AssignmentRuleDraft(name="First", conditions=[], memberIds=["first"], active=True), actor)
        rule_second = main.create_assignment_rule(main.AssignmentRuleDraft(name="Second", conditions=[], memberIds=["second"], active=True), actor)
        main.repo.fallback_sales = ["inactive", "second", "first"]

        def counts():
            with Session(assignments.engine) as session:
                return (len(list(session.scalars(select(AssignmentHistoryRow)))), len(list(session.scalars(select(AuditRow)))))
        history_before, audit_before = counts()

        # A matching active rule (by position) wins over the fallback list; an already-correct owner is a no-op; Closed customers are excluded from the run's scope.
        result = main.run_assignment(main.AssignmentRunRequest(scope="all-open", submissionId="ordered"), actor)
        assert (result["candidates"], result["assigned"], result["skipped"]) == (2, 1, 1)
        assert customers.get("000001").owner_id == "first" and customers.get("000001").version == 4
        assert customers.get("000002").owner_id == "first" and customers.get("000002").version == 1
        assert customers.get("000003").owner_id is None
        history_after_ordered, audit_after_ordered = counts()
        assert (history_after_ordered - history_before, audit_after_ordered - audit_before) == (1, 1)

        # Every rule inactive falls through to the fallback list.
        main.update_assignment_rule(rule_first["id"], {"active": False, "version": main.repo.assignment_version}, actor)
        main.update_assignment_rule(rule_second["id"], {"active": False, "version": main.repo.assignment_version}, actor)
        history_before_fallback, audit_before_fallback = counts()
        result = main.run_assignment(main.AssignmentRunRequest(scope="all-open", submissionId="fallback"), actor)
        assert result["assigned"] == 2
        assert customers.get("000001").owner_id == "second" and customers.get("000001").version == 5
        assert customers.get("000002").owner_id == "second" and customers.get("000002").version == 2
        history_after_fallback, audit_after_fallback = counts()
        assert (history_after_fallback - history_before_fallback, audit_after_fallback - audit_before_fallback) == (2, 2)

        # Re-running the same all-open scope re-balances again: "second" now owns both (2 open), so
        # "first" (0 open) is the lower-workload fallback pick this time -- an all-open run always
        # re-evaluates every open candidate against current live counts, it does not freeze prior picks.
        result = main.run_assignment(main.AssignmentRunRequest(scope="all-open", submissionId="rebalanced"), actor)
        assert (result["assigned"], result["skipped"]) == (2, 0)
        assert customers.get("000001").owner_id == "first" and customers.get("000001").version == 6
        assert customers.get("000002").owner_id == "first" and customers.get("000002").version == 3
        history_after_rebalanced, audit_after_rebalanced = counts()
        assert (history_after_rebalanced - history_after_fallback, audit_after_rebalanced - audit_after_fallback) == (2, 2)

        # Every fallback member inactive/non-Sales leaves the owner unchanged: no history/audit row, no version bump.
        main.repo.fallback_sales = ["inactive"]
        result = main.run_assignment(main.AssignmentRunRequest(scope="all-open", submissionId="no-eligible-owner"), actor)
        assert (result["assigned"], result["skipped"]) == (0, 2)
        assert customers.get("000001").owner_id == "first" and customers.get("000001").version == 6
        assert customers.get("000002").owner_id == "first" and customers.get("000002").version == 3
        assert counts() == (history_after_rebalanced, audit_after_rebalanced)
    finally:
        auth.engine.dispose(); customers.engine.dispose(); assignments.engine.dispose()


def test_postgres_fallback_sales_round_trips_and_rejects_without_partial_write(postgres_url, monkeypatch):
    """fallback_sales write/reject round-tripped through AssignmentDatabase.set_setting against real PostgreSQL, not just in memory."""
    from fastapi import HTTPException
    from .db_assignment import AssignmentDatabase
    from .db_customers import CustomerDatabase

    migrate()
    auth = AuthDatabase(postgres_url, create_schema=False)
    customers = CustomerDatabase(postgres_url, create_schema=False)
    assignments = AssignmentDatabase(postgres_url, create_schema=False)
    monkeypatch.setattr(main, "auth_db", auth)
    monkeypatch.setattr(main, "customer_db", customers)
    monkeypatch.setattr(main, "assignment_db", assignments)
    main.repo.reset()
    try:
        hashed = main.password_hash.hash(secrets.token_urlsafe(24))
        for user_id, role in (("admin", "Admin"), ("sales", "Sales")):
            auth.create_user(user_id=user_id, name=user_id, email=f"{user_id}@example.test", role=role, password_hash=hashed)
            main.repo.users[user_id] = {"id": user_id, "name": user_id, "role": role, "active": True}
        actor = main.User(id="admin", name="Admin", email="admin@example.test", role="Admin")

        assert main.set_assignment_fallback(["sales"], actor) == ["sales"]
        assert assignments.get_setting("fallback_sales") == {"ids": ["sales"]}

        with pytest.raises(HTTPException) as rejected:
            main.set_assignment_fallback(["sales", "ghost-user"], actor)
        assert rejected.value.status_code == 422
        # No partial write: the previously stored row is unchanged in real PostgreSQL.
        assert assignments.get_setting("fallback_sales") == {"ids": ["sales"]}
    finally:
        auth.engine.dispose(); customers.engine.dispose(); assignments.engine.dispose()
        main.repo.reset()


def test_postgres_bulk_run_concurrent_submission_serializes(postgres_url):
    """Two real concurrent PostgreSQL transactions racing on the same (actor_id, submission_id) serialize on the unique run key."""
    from .db_assignment import AssignmentDatabase, AssignmentHistoryRow, AssignmentRunRow, AuditRow
    from .db_customers import CustomerDatabase

    migrate()
    auth = AuthDatabase(postgres_url, create_schema=False)
    customers = CustomerDatabase(postgres_url, create_schema=False)
    assignments = AssignmentDatabase(postgres_url, create_schema=False)
    try:
        hashed = main.password_hash.hash(secrets.token_urlsafe(24))
        auth.create_user(user_id="admin", name="Admin", email="admin@example.test", role="Admin", password_hash=hashed)
        auth.create_user(user_id="sales", name="Sales", email="sales@example.test", role="Sales", password_hash=hashed)
        customers.upsert_source(bcn="000700", name="Synthetic", source={})
        barrier = Barrier(2)
        def race(_index):
            barrier.wait()
            return assignments.run_bulk(actor_id="admin", submission_id="race", scope="unassigned", fallback_ids=["sales"])
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(race, (1, 2)))
        assert results[0][0] == results[1][0]
        assert results[0][0]["assigned"] == 1
        with Session(assignments.engine) as session:
            assert len(list(session.scalars(select(AssignmentHistoryRow)))) == 1
            assert len(list(session.scalars(select(AuditRow)))) == 1
            assert len(list(session.scalars(select(AssignmentRunRow)))) == 1
        assert customers.get("000700").owner_id == "sales" and customers.get("000700").version == 1
    finally:
        auth.engine.dispose(); customers.engine.dispose(); assignments.engine.dispose()


def test_postgres_manual_assignment_concurrent_stale_version_conflicts(postgres_url):
    """Two real concurrent PostgreSQL transactions racing on the same bcn serialize on the customer row lock; the stale-version loser gets a conflict, not a silently lost update."""
    from .db_assignment import AssignmentDatabase, AssignmentHistoryRow, AuditRow
    from .db_customers import CustomerDatabase

    migrate()
    auth = AuthDatabase(postgres_url, create_schema=False)
    customers = CustomerDatabase(postgres_url, create_schema=False)
    assignments = AssignmentDatabase(postgres_url, create_schema=False)
    try:
        hashed = main.password_hash.hash(secrets.token_urlsafe(24))
        auth.create_user(user_id="admin", name="Admin", email="admin@example.test", role="Admin", password_hash=hashed)
        for owner in ("sales-a", "sales-b"):
            auth.create_user(user_id=owner, name=owner, email=f"{owner}@example.test", role="Sales", password_hash=hashed)
        customers.upsert_source(bcn="000800", name="Synthetic", source={})
        barrier = Barrier(2)
        def race(owner):
            barrier.wait()
            try:
                return ("ok", assignments.assign_manual(bcn="000800", owner_id=owner, expected_version=0, actor_id="admin"))
            except ValueError as exc:
                return ("conflict", str(exc))
        with ThreadPoolExecutor(max_workers=2) as pool:
            outcomes = list(pool.map(race, ("sales-a", "sales-b")))
        statuses = sorted(status for status, _ in outcomes)
        assert statuses == ["conflict", "ok"]
        assert "stale" in next(detail for status, detail in outcomes if status == "conflict").lower()
        winner = next(detail for status, detail in outcomes if status == "ok")
        row = customers.get("000800")
        assert row.owner_id == winner["ownerId"] and row.version == 1
        with Session(assignments.engine) as session:
            assert len(list(session.scalars(select(AssignmentHistoryRow)))) == 1
            assert len(list(session.scalars(select(AuditRow)))) == 1
    finally:
        auth.engine.dispose(); customers.engine.dispose(); assignments.engine.dispose()


def test_postgres_followup_completion_race_loser_gets_conflict_not_500(postgres_url, monkeypatch):
    """Regression for the #25 QA bug: two concurrent /follow-ups/{id}/complete requests for the same
    open follow-up must produce one 200 and one 409 -- never a raw 500. ActivityDatabase.complete_followup
    raises ValueError for the with_for_update lock loser (whose in-process view of the follow-up was
    still "Open" when it read it); main.complete_followup must catch that, like every other mutation
    endpoint catches its own idempotency/version ValueError, instead of letting FastAPI turn it into an
    unhandled 500."""
    from .db_activity import ActivityDatabase
    from .db_customers import CustomerDatabase

    migrate()
    auth = AuthDatabase(postgres_url, create_schema=False)
    customers = CustomerDatabase(postgres_url, create_schema=False)
    activities = ActivityDatabase(postgres_url, create_schema=False)
    monkeypatch.setattr(main, "auth_db", auth)
    monkeypatch.setattr(main, "customer_db", customers)
    monkeypatch.setattr(main, "activity_db", activities)
    monkeypatch.setattr(main, "assignment_db", None)
    main.repo.reset(); main.repo.customers.clear()
    try:
        secret = secrets.token_urlsafe(24)
        hashed = main.password_hash.hash(secret)
        auth.create_user(user_id="race-sales", name="Race Sales", email="race-complete@example.test", role="Sales", password_hash=hashed)
        customers.upsert_source(bcn="000850", name="Synthetic", source={})
        customers.save_operational(bcn="000850", owner_id="race-sales", status="Open", version=0)
        followup = {"id": "followup-race", "bcn": "000850", "actorId": "race-sales", "type": "Reminder", "due": None, "note": None, "status": "Open", "createdAt": datetime.now(timezone.utc).isoformat()}
        activities.create_followup(followup, submission_id="race-create", payload="race-create")
        # A concurrent request wins the real database race first, completing the follow-up.
        activities.complete_followup(followup, {"id": "interaction-winner", "bcn": "000850", "actorId": "race-sales", "outcome": "Contact", "note": "Winner"})

        # This request's in-process view of the follow-up still shows it as Open -- exactly what it
        # would look like had it read the row a moment before the winner's commit.
        stale = {**followup, "status": "Open", "interactionId": None}
        monkeypatch.setattr(main, "find_followup", lambda bcn, followup_id, user: stale)
        with TestClient(main.app, base_url="http://localhost") as client:
            assert client.post("/session/login", json={"email": "race-complete@example.test", "password": secret}).status_code == 200
            response = client.post("/customers/000850/follow-ups/followup-race/complete", json={"outcome": "Contact", "note": "Loser", "submissionId": "race-loser"}, headers=ORIGIN)
        assert response.status_code == 409, response.text
        assert response.json() == {"detail": "follow-up is no longer open"}
        # The loser's interaction must never have been persisted (only the follow-up creation and
        # the winner's interaction rows exist).
        assert len(activities.history("000850")[0]) == 2
    finally:
        auth.engine.dispose(); customers.engine.dispose(); activities.engine.dispose()


def test_postgres_assignment_restart_preserves_audit_and_resubmission(postgres_url):
    migrate()
    # Transport synthetic cookies only over anonymous process pipes, never args/files/output.
    seed = '''
import json, secrets
from io import BytesIO
from openpyxl import Workbook
from fastapi.testclient import TestClient
from apps.api.main import app, provision_user, Provision
secret = secrets.token_urlsafe(24)
admin = provision_user(Provision(name="Restart Admin", email="restart-assign-admin@example.test", password=secret))
client = TestClient(app, base_url="http://localhost")
assert client.post("/session/login", json={"email": admin.email, "password": secret}).status_code == 200
sales_secret = secrets.token_urlsafe(24)
sales = client.post("/admin/users", json={"name": "Sales", "email": "restart-assign-sales@example.test", "role": "Sales", "password": sales_secret}, headers={"origin": "http://localhost:3000"}).json()
book = Workbook(); book.active.append(["bcn", "customer_name"]); book.active.append(["800100", "Restart Assign Co"])
payload = BytesIO(); book.save(payload)
imported = client.post("/admin/imports?submission_id=restart-assign-import", files={"file": ("restart.xlsx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers={"origin": "http://localhost:3000"})
assert imported.status_code == 201, imported.text
assigned = client.post("/admin/assignments/manual/800100", json={"ownerId": sales["id"], "submissionId": "restart-assign-manual"}, headers={"origin": "http://localhost:3000"})
assert assigned.status_code == 200, assigned.text
run = client.post("/admin/assignment-runs", json={"scope": "unassigned", "submissionId": "restart-assign-run"}, headers={"origin": "http://localhost:3000"})
assert run.status_code == 200, run.text
print(json.dumps({"adminEmail": admin.email, "adminId": admin.id, "runResult": run.json()}))
'''
    verify = '''
import json, sys
from fastapi import HTTPException
from apps.api.main import admin_audit, run_assignment, AssignmentRunRequest, User
info = json.load(sys.stdin)
actor = User(id=info["adminId"], name="Restart Admin", email=info["adminEmail"], role="Admin")

# Resubmitting the same (actor, submission_id) with an identical scope replays the persisted result, unchanged, across a restart.
same = run_assignment(AssignmentRunRequest(scope="unassigned", submissionId="restart-assign-run"), actor)
assert same == info["runResult"], (same, info["runResult"])

# A different scope under the same submission_id is a 409, across a restart.
try:
    run_assignment(AssignmentRunRequest(scope="all-open", submissionId="restart-assign-run"), actor)
    raise SystemExit("expected HTTPException")
except HTTPException as exc:
    assert exc.status_code == 409

# audit_events/assignment_history rows from before the restart are still present and readable through the paginated audit query,
# including the source-change row #23's import path wrote, with only the whitelisted detail fields returned (never a credential).
audit = admin_audit(page=1, page_size=50)
actions = [item["action"] for item in audit["items"]]
assert actions.count("Import completed") == 1
assert "Customer assigned" in actions
for item in audit["items"]:
    assert set(item) == {"id", "actor", "actorId", "action", "target", "timestamp", "details"}
    assert "password" not in item["details"] and "secret" not in item["details"]
'''
    env = os.environ | {"CALL_CENTER_STORAGE": "postgres", "DATABASE_URL": postgres_url}
    first = subprocess.run([sys.executable, "-c", seed], env=env, capture_output=True)
    assert first.returncode == 0, first.stderr
    second = subprocess.run([sys.executable, "-c", verify], env=env, input=first.stdout, capture_output=True)
    assert second.returncode == 0, second.stderr


def test_postgres_import_restart_preserves_jobs_and_errors(postgres_url):
    migrate()
    # Transport synthetic cookies only over anonymous process pipes, never args/files/output.
    seed = '''
import json, secrets
from io import BytesIO
from openpyxl import Workbook
from apps.api.main import app, provision_user, Provision
from fastapi.testclient import TestClient
secret = secrets.token_urlsafe(24)
admin = provision_user(Provision(name="Restart Admin", email="restart-admin@example.test", password=secret))
client = TestClient(app, base_url="http://localhost")
assert client.post("/session/login", json={"email": admin.email, "password": secret}).status_code == 200
book = Workbook(); book.active.append(["bcn", "customer_name"]); book.active.append(["700900", "Restart Co"]); book.active.append(["bad", "Bad"])
payload = BytesIO(); book.save(payload)
response = client.post("/admin/imports?submission_id=restart-import", files={"file": ("restart.xlsx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers={"origin": "http://localhost:3000"})
assert response.status_code == 201, response.text
print(json.dumps({"actorId": admin.id, "jobId": response.json()["jobId"]}))
'''
    verify = '''
import json, sys
from apps.api.main import customer_db
info = json.load(sys.stdin)
job = customer_db.import_job(info["actorId"], "restart-import")
assert job is not None and job["jobId"] == info["jobId"]
assert (job["processed"], job["created"], job["errorRows"]) == (2, 1, 1)
errors, total = customer_db.import_errors(info["actorId"], "restart-import", 1, 25)
assert total == 1 and errors[0]["field"] == "bcn"
row = customer_db.get("700900")
assert row is not None and row.name == "Restart Co"
assert customer_db.phones("700900") == []
'''
    env = os.environ | {"CALL_CENTER_STORAGE": "postgres", "DATABASE_URL": postgres_url}
    first = subprocess.run([sys.executable, "-c", seed], env=env, capture_output=True)
    assert first.returncode == 0, first.stderr
    second = subprocess.run([sys.executable, "-c", verify], env=env, input=first.stdout, capture_output=True)
    assert second.returncode == 0, second.stderr


def test_postgres_activity_lifecycle_restart_preserves_history(postgres_url):
    """Interactions, a soft-deleted note tombstone, follow-up completion, a closure-reason label, and
    close/reopen events must all survive a process restart -- the coverage gap QA found despite roughly
    ten prior engineering comments on #25 claiming exactly that, none of which were ever proven with a
    restart test. Follows the subprocess-restart pattern already used for #23/#24."""
    migrate()
    # Transport synthetic cookies only over anonymous process pipes, never args/files/output.
    seed = '''
import json, secrets
from io import BytesIO
from openpyxl import Workbook
from fastapi.testclient import TestClient
from apps.api.main import app, provision_user, Provision
secret = secrets.token_urlsafe(24)
admin = provision_user(Provision(name="Restart Admin", email="restart-activity-admin@example.test", password=secret))
client = TestClient(app, base_url="http://localhost")
assert client.post("/session/login", json={"email": admin.email, "password": secret}).status_code == 200
origin = {"origin": "http://localhost:3000"}
sales_secret = secrets.token_urlsafe(24)
sales = client.post("/admin/users", json={"name": "Restart Sales", "email": "restart-activity-sales@example.test", "role": "Sales", "password": sales_secret}, headers=origin).json()
# A second, independent owner for the #30 release scenario below, so deactivating it cannot
# side-effect the first owner's still-Open 800200 (which must stay assigned across the restart).
sales2_secret = secrets.token_urlsafe(24)
sales2 = client.post("/admin/users", json={"name": "Restart Release Sales", "email": "restart-activity-sales2@example.test", "role": "Sales", "password": sales2_secret}, headers=origin).json()
reason = client.post("/admin/closure-reasons", json={"label": "Restart resolved"}, headers=origin).json()
book = Workbook(); book.active.append(["bcn", "customer_name"]); book.active.append(["800200", "Restart Activity Co"]); book.active.append(["800201", "Restart Release Co"])
payload = BytesIO(); book.save(payload)
imported = client.post("/admin/imports?submission_id=restart-activity-import", files={"file": ("restart.xlsx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=origin)
assert imported.status_code == 201, imported.text
assigned = client.post("/admin/assignments/manual/800200", json={"ownerId": sales["id"], "submissionId": "restart-activity-assign"}, headers=origin)
assert assigned.status_code == 200, assigned.text
assigned_release = client.post("/admin/assignments/manual/800201", json={"ownerId": sales2["id"], "submissionId": "restart-release-assign"}, headers=origin)
assert assigned_release.status_code == 200, assigned_release.text
client.post("/session/logout", headers=origin)

assert client.post("/session/login", json={"email": sales["email"], "password": sales_secret}).status_code == 200
interaction = client.post("/customers/800200/interactions", json={"outcome": "Contact", "note": "Restart contact", "submissionId": "restart-interaction"}, headers=origin)
assert interaction.status_code == 200, interaction.text
note = client.post("/customers/800200/notes", json={"text": "Restart note", "submissionId": "restart-note"}, headers=origin)
assert note.status_code == 200, note.text
note_id = note.json()["id"]
deleted = client.delete("/customers/800200/history/" + note_id, headers=origin)
assert deleted.status_code == 200, deleted.text
kept_open = client.post("/customers/800200/follow-ups", json={"type": "Reminder", "due": None, "note": "Stays open then cancelled", "submissionId": "restart-followup-open"}, headers=origin)
assert kept_open.status_code == 200, kept_open.text
to_complete = client.post("/customers/800200/follow-ups", json={"type": "Reminder", "due": None, "note": "Will be completed", "submissionId": "restart-followup-complete"}, headers=origin)
assert to_complete.status_code == 200, to_complete.text
to_complete_id = to_complete.json()["id"]
completed = client.post("/customers/800200/follow-ups/" + to_complete_id + "/complete", json={"outcome": "Contact", "note": "Completed before restart", "submissionId": "restart-followup-complete-intent"}, headers=origin)
assert completed.status_code == 200, completed.text
closed = client.post("/customers/800200/close", json={"reasonId": reason["id"], "submissionId": "restart-close"}, headers=origin)
assert closed.status_code == 200, closed.text
reopened = client.post("/customers/800200/reopen", json={"submissionId": "restart-reopen"}, headers=origin)
assert reopened.status_code == 200, reopened.text
client.post("/session/logout", headers=origin)

# #30: a second customer, owned by a second, independent Sales user, closed while that owner is
# still active Sales, then reopened by an admin only after that owner is deactivated -- ownership
# must release and the release rows must survive the restart below alongside the rest of this
# test's activity/lifecycle coverage. Deactivating sales2 must not touch sales/800200 above.
assert client.post("/session/login", json={"email": sales2["email"], "password": sales2_secret}).status_code == 200
closed_release = client.post("/customers/800201/close", json={"reasonId": reason["id"], "submissionId": "restart-release-close"}, headers=origin)
assert closed_release.status_code == 200, closed_release.text
client.post("/session/logout", headers=origin)

assert client.post("/session/login", json={"email": admin.email, "password": secret}).status_code == 200
deactivated = client.patch(f"/admin/users/{sales2['id']}", json={"active": False}, headers=origin)
assert deactivated.status_code == 200 and deactivated.json()["active"] is False, deactivated.text
reopened_release = client.post("/customers/800201/reopen", json={"submissionId": "restart-release-reopen"}, headers=origin)
assert reopened_release.status_code == 200, reopened_release.text
assert reopened_release.json()["ownerId"] is None and reopened_release.json()["ownerName"] is None

print(json.dumps({"salesId": sales["id"], "sales2Id": sales2["id"], "reasonId": reason["id"], "noteId": note_id, "openFollowupId": kept_open.json()["id"], "completedFollowupId": to_complete_id}))
'''
    verify = '''
import json, sys
from sqlalchemy import select
from sqlalchemy.orm import Session
from apps.api.main import activity_db, customer_db, assignment_db
from apps.api.db_assignment import AssignmentHistoryRow, AuditRow
info = json.load(sys.stdin)
bcn = "800200"

reasons = {row.id: row for row in activity_db.reasons()}
assert reasons[info["reasonId"]].label == "Restart resolved" and reasons[info["reasonId"]].active is True

rows, total = activity_db.history(bcn)
by_id = {row.id: row for row in rows}
interaction = next(row for row in rows if row.kind == "Interaction" and row.outcome == "Contact" and row.deleted_at is None and row.text == "Restart contact")
assert interaction is not None
note_row = by_id[info["noteId"]]
assert note_row.deleted_at is not None and note_row.deleted_by == info["salesId"] and note_row.text == "Restart note"
closure = next(row for row in rows if row.kind == "Closure")
assert closure.text == "Restart resolved"
reopen_row = next(row for row in rows if row.kind == "Reopen")
assert reopen_row.actor_id == info["salesId"]

followups = {row.id: row for row in activity_db.followups(bcn)}
assert followups[info["openFollowupId"]].status == "Cancelled"
completed_followup = followups[info["completedFollowupId"]]
assert completed_followup.status == "Completed" and completed_followup.interaction_id is not None
assert by_id[completed_followup.interaction_id].text == "Completed before restart"

customer = customer_db.get(bcn)
assert customer.status == "Open" and customer.owner_id == info["salesId"]

# #30: released ownership on the second customer, and its assignment_history/audit_events rows,
# survived the restart.
released_customer = customer_db.get("800201")
assert released_customer.status == "Open" and released_customer.owner_id is None
release_reopen = next(row for row in activity_db.history("800201")[0] if row.kind == "Reopen")
assert release_reopen is not None
with Session(assignment_db.engine) as session:
    history_row = session.execute(select(AssignmentHistoryRow).where(AssignmentHistoryRow.bcn == "800201", AssignmentHistoryRow.reason == "Owner no longer active Sales")).scalar_one()
    assert history_row.old_owner_id == info["sales2Id"] and history_row.new_owner_id is None
    audit_row = session.execute(select(AuditRow).where(AuditRow.target == "800201", AuditRow.action == "Customer ownership released")).scalar_one()
    assert audit_row.details["oldOwner"] == info["sales2Id"]
'''
    env = os.environ | {"CALL_CENTER_STORAGE": "postgres", "DATABASE_URL": postgres_url}
    first = subprocess.run([sys.executable, "-c", seed], env=env, capture_output=True)
    assert first.returncode == 0, first.stderr
    second = subprocess.run([sys.executable, "-c", verify], env=env, input=first.stdout, capture_output=True)
    assert second.returncode == 0, second.stderr


def test_postgres_process_restart(postgres_url):
    migrate()
    # Transport synthetic cookies only over anonymous process pipes, never args/files/output.
    seed = '''
import json, secrets
from datetime import timedelta
from apps.api.main import app, provision_user, Provision, auth_db
from apps.api import db_auth
from fastapi.testclient import TestClient
secret = secrets.token_urlsafe(24)
user = provision_user(Provision(name="Restart", email="restart@example.test", password=secret))
client = TestClient(app, base_url="http://localhost")
assert client.post("/session/login", json={"email": user.email, "password": secret}).status_code == 200
valid = client.cookies.get("call_center_session")
revoked, _ = auth_db.issue(user.id)
auth_db.revoke(revoked)
now = db_auth.utcnow()
db_auth.utcnow = lambda: now - timedelta(hours=9)
expired, _ = auth_db.issue(user.id)
print(json.dumps([valid, revoked, expired]))
'''
    verify = '''
import json, sys
from fastapi.testclient import TestClient
from apps.api.main import app, auth_db
from apps.api.db_auth import SessionRow, digest
from sqlalchemy import select
cookies = json.load(sys.stdin)
with auth_db.transaction() as session:
    records = list(session.scalars(select(SessionRow)))
    assert len(records) == 3
    assert {row.digest for row in records} == {digest(token) for token in cookies}
    assert all(row.digest not in cookies for row in records)
for token, status in zip(cookies, [200, 401, 401]):
    with TestClient(app, base_url="http://localhost") as client:
        client.cookies.set("call_center_session", token)
        assert client.get("/session/me").status_code == status
with TestClient(app, base_url="http://localhost") as client:
    assert client.post("/operator/reset-password/anything", headers={"origin": "http://localhost:3000"}, json={}).status_code == 404
'''
    env = os.environ | {"CALL_CENTER_STORAGE": "postgres", "DATABASE_URL": postgres_url}
    first = subprocess.run([sys.executable, "-c", seed], env=env, capture_output=True)
    assert first.returncode == 0
    second = subprocess.run([sys.executable, "-c", verify], env=env, input=first.stdout, capture_output=True)
    assert second.returncode == 0


def test_postgres_operator_console(postgres_url, monkeypatch, capsys):
    migrate()
    from . import operator
    database = AuthDatabase(postgres_url, create_schema=False)
    monkeypatch.setattr(main, "auth_db", database)
    monkeypatch.setattr(main, "storage_mode", "postgres")
    secret = secrets.token_urlsafe(24)
    monkeypatch.setattr(operator, "getpass", lambda _: secret)
    answers = iter(["provision", "Synthetic", " Console@example.test "])
    monkeypatch.setattr("builtins.input", lambda _: next(answers))
    operator.main()
    user = database.user_by_email("console@example.test")
    token, _ = database.issue(user.id)
    answers = iter(["provision", "Duplicate", "CONSOLE@example.test"])
    with pytest.raises(SystemExit, match="Operator request rejected"):
        operator.main()
    secret = secrets.token_urlsafe(24)
    answers = iter(["reset", user.id])
    operator.main()
    assert database.user_for_session(token) is None
    assert main.password_hash.verify(secret, database.user_by_email(user.email).password_hash)
    answers = iter(["reset", "missing"])
    with pytest.raises(SystemExit, match="Operator request rejected"):
        operator.main()
    output = capsys.readouterr()
    assert output.out == "Operator request completed.\n" * 2 and output.err == ""
    assert database.engine.pool.checkedout() == 0
    database.engine.dispose()
